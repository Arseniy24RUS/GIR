from __future__ import annotations

import base64
import csv
import hashlib
import io
import json
import mimetypes
import re
import sqlite3
from collections import defaultdict
from functools import lru_cache
from pathlib import Path
from typing import Any, Iterable

from fastapi import APIRouter, HTTPException, Query, Request
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, RedirectResponse, StreamingResponse

try:
    from . import config as _config
except Exception:  # pragma: no cover - defensive import fallback
    _config = None

ROOT = Path(__file__).resolve().parents[1]
DB_PATH = Path(getattr(_config, "DB_PATH", ROOT / "data" / "global_index_platform.sqlite"))
STATIC_DIR = Path(getattr(_config, "STATIC_DIR", ROOT / "giip" / "static"))

router = APIRouter(tags=["data-lab"])


# ---------------------------------------------------------------------------
# Generic database helpers
# ---------------------------------------------------------------------------

def _connect() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def _dicts(cursor: sqlite3.Cursor) -> list[dict[str, Any]]:
    return [dict(row) for row in cursor.fetchall()]


def _table_exists(conn: sqlite3.Connection, table: str) -> bool:
    return conn.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name=?", (table,)
    ).fetchone() is not None


@lru_cache(maxsize=256)
def _cached_columns(db_mtime_ns: int, table: str) -> tuple[str, ...]:
    del db_mtime_ns
    with _connect() as conn:
        if not _table_exists(conn, table):
            return ()
        return tuple(row[1] for row in conn.execute(f'PRAGMA table_info("{table}")'))


def _columns(table: str) -> tuple[str, ...]:
    token = DB_PATH.stat().st_mtime_ns if DB_PATH.exists() else 0
    return _cached_columns(token, table)


def _first_existing(columns: Iterable[str], *candidates: str) -> str | None:
    available = set(columns)
    return next((name for name in candidates if name in available), None)


def _quote_identifier(value: str) -> str:
    return '"' + value.replace('"', '""') + '"'


def _pick(row: dict[str, Any] | None, *keys: str, default: Any = None) -> Any:
    if not row:
        return default
    for key in keys:
        value = row.get(key)
        if value is not None and value != "":
            return value
    return default


def _as_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    return str(value or "").strip().lower() in {
        "1", "true", "yes", "y", "allowed", "include", "included", "open", "public",
    }


def _normalise_decision(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def _decision_allows(value: Any) -> bool:
    decision = _normalise_decision(value)
    return decision in {
        "include", "included", "allow", "allowed", "public", "open", "redistributable",
        "include_raw", "include_in_release", "publish", "published",
    }


def _decision_restricts(value: Any) -> bool:
    decision = _normalise_decision(value)
    return decision in {
        "exclude", "excluded", "restricted", "do_not_distribute", "metadata_only",
        "exclude_raw_keep_derived", "no_redistribution", "private",
    }


def _safe_path(raw_path: Any) -> Path | None:
    if raw_path is None or str(raw_path).strip() == "":
        return None
    candidate = Path(str(raw_path).strip())
    probes = [candidate] if candidate.is_absolute() else [
        ROOT / candidate,
        DB_PATH.parent / candidate,
        ROOT / "data" / candidate,
    ]
    root_resolved = ROOT.resolve()
    for probe in probes:
        try:
            resolved = probe.resolve()
            resolved.relative_to(root_resolved)
        except (OSError, ValueError):
            continue
        if resolved.exists() and resolved.is_file():
            return resolved
    return None


def _sha256(path: Path | None) -> str | None:
    if not path or not path.exists():
        return None
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _human_bytes(value: Any) -> str:
    number = float(value or 0)
    units = ["B", "KB", "MB", "GB", "TB"]
    index = 0
    while number >= 1024 and index < len(units) - 1:
        number /= 1024
        index += 1
    return f"{number:.1f} {units[index]}" if index else f"{int(number)} B"


def _file_extension(path: Any, content_type: Any = None) -> str:
    suffix = Path(str(path or "")).suffix.lower().lstrip(".")
    if suffix:
        return suffix
    guessed = mimetypes.guess_extension(str(content_type or "").split(";")[0].strip()) or ""
    return guessed.lstrip(".") or "file"


def _media_type(path: Any, content_type: Any = None) -> str:
    explicit = str(content_type or "").strip()
    if explicit and "/" in explicit:
        return explicit.split(";")[0]
    return mimetypes.guess_type(str(path or ""))[0] or "application/octet-stream"


def _encode_file_id(kind: str, identifier: str) -> str:
    token = base64.urlsafe_b64encode(identifier.encode("utf-8")).decode("ascii").rstrip("=")
    return f"{kind}.{token}"


def _decode_file_id(file_id: str) -> tuple[str, str]:
    try:
        kind, token = file_id.split(".", 1)
        if kind not in {"raw", "artifact"}:
            raise ValueError
        token += "=" * (-len(token) % 4)
        identifier = base64.urlsafe_b64decode(token.encode("ascii")).decode("utf-8")
        if not identifier:
            raise ValueError
        return kind, identifier
    except Exception as exc:
        raise HTTPException(status_code=404, detail="Unknown catalog file") from exc


# ---------------------------------------------------------------------------
# Catalog assembly
# ---------------------------------------------------------------------------

def _source_rows(conn: sqlite3.Connection) -> list[dict[str, Any]]:
    if not _table_exists(conn, "source_registry"):
        return []
    return _dicts(conn.execute("SELECT * FROM source_registry ORDER BY source_id"))


def _raw_rows(conn: sqlite3.Connection) -> list[dict[str, Any]]:
    if not _table_exists(conn, "raw_snapshots"):
        return []
    return _dicts(conn.execute("SELECT * FROM raw_snapshots"))


def _artifact_rows(conn: sqlite3.Connection) -> list[dict[str, Any]]:
    if not _table_exists(conn, "reproducibility_artifacts"):
        return []
    return _dicts(conn.execute("SELECT * FROM reproducibility_artifacts"))


def _source_rights(source: dict[str, Any] | None) -> tuple[bool, str, str | None, str | None]:
    source = source or {}
    decision = _pick(
        source,
        "release_archive_decision", "distribution_decision", "redistribution_decision",
        "redistribution_allowed", "include_in_release",
    )
    allowed = _as_bool(_pick(source, "redistribution_allowed", "include_in_release")) or _decision_allows(decision)
    if _decision_restricts(decision):
        allowed = False
    return (
        allowed,
        str(decision or ("include" if allowed else "metadata_only")),
        _pick(source, "terms_url", "license_url", "source_terms_url"),
        _pick(source, "attribution_required", "attribution", "citation_text"),
    )


def _infer_source_role(source_id: str, source: dict[str, Any] | None) -> str:
    explicit = _pick(source, "source_role", "role", "source_type")
    if explicit:
        return str(explicit)
    code = source_id.upper()
    if any(token in code for token in ("WORLD_BANK", "WIPO", "ITU", "ILO", "OECD", "UIS", "UNESCO", "UNDP")):
        return "international_organisation"
    if any(token in code for token in ("ROSSTAT", "NATIONAL", "STATISTICS")):
        return "national_statistics"
    if any(token in code for token in ("QS", "RANK", "PORTULANS", "GTCI")):
        return "specialised_ranking"
    if any(token in code for token in ("SEC", "CORPORATE", "COMPANY")):
        return "corporate_reporting"
    return "research_source"


def _catalog_records(conn: sqlite3.Connection) -> tuple[dict[str, dict[str, Any]], list[dict[str, Any]]]:
    source_rows = _source_rows(conn)
    sources = {str(row.get("source_id")): row for row in source_rows if row.get("source_id")}
    raw_rows = _raw_rows(conn)
    artifact_rows = _artifact_rows(conn)

    raw_by_snapshot: dict[str, dict[str, Any]] = {}
    for item in raw_rows:
        snapshot_id = str(_pick(item, "snapshot_id", "raw_snapshot_id", "id", default=""))
        if snapshot_id:
            raw_by_snapshot[snapshot_id] = item

    artifacts_by_snapshot: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for item in artifact_rows:
        snapshot_id = str(_pick(item, "snapshot_id", "raw_snapshot_id", default=""))
        if snapshot_id:
            artifacts_by_snapshot[snapshot_id].append(item)

    files: list[dict[str, Any]] = []
    for item in raw_rows:
        snapshot_id = str(_pick(item, "snapshot_id", "raw_snapshot_id", "id", default=""))
        if not snapshot_id:
            continue
        source_id = str(_pick(item, "source_id", default="UNKNOWN"))
        source = sources.get(source_id, {})
        source_allowed, source_decision, terms_url, attribution = _source_rights(source)
        associated = artifacts_by_snapshot.get(snapshot_id, [])
        raw_artifact = next(
            (
                artifact for artifact in associated
                if "raw" in str(_pick(artifact, "artifact_type", "distribution_kind", default="")).lower()
            ),
            None,
        )
        artifact_decision = _pick(raw_artifact, "distribution_decision", "release_archive_decision")
        allowed = source_allowed or _decision_allows(artifact_decision)
        if _decision_restricts(artifact_decision) or _decision_restricts(source_decision):
            allowed = False
        raw_path = _pick(item, "raw_snapshot_path", "snapshot_path", "file_path", "path")
        resolved = _safe_path(raw_path)
        content_type = _pick(item, "content_type", "media_type", "mime_type")
        extension = _file_extension(raw_path, content_type)
        sha = _pick(item, "raw_snapshot_sha256", "snapshot_sha256", "sha256") or _sha256(resolved)
        size = _pick(item, "bytes_count", "content_length", "file_size", "bytes")
        if size is None and resolved:
            size = resolved.stat().st_size
        filename = Path(str(raw_path or snapshot_id)).name
        files.append({
            "file_id": _encode_file_id("raw", snapshot_id),
            "record_kind": "raw",
            "record_id": snapshot_id,
            "snapshot_id": snapshot_id,
            "artifact_id": None,
            "source_id": source_id,
            "title_ru": f"Исходный snapshot · {filename}",
            "title_en": f"Source snapshot · {filename}",
            "filename": filename,
            "catalog_path": str(raw_path or ""),
            "resolved_path": str(resolved) if resolved else None,
            "exists_locally": bool(resolved),
            "distribution_kind": "raw_snapshot",
            "artifact_type": "raw_snapshot",
            "extension": extension,
            "media_type": _media_type(raw_path, content_type),
            "bytes_count": int(size or 0),
            "bytes_human": _human_bytes(size),
            "sha256": sha,
            "release_year": _pick(item, "release_year", "source_data_year", "year"),
            "retrieved_at": _pick(item, "retrieved_at", "retrieval_date", "created_at"),
            "official_url": _pick(item, "source_url", "official_url") or _pick(source, "source_url", "official_url", "url"),
            "terms_url": terms_url,
            "attribution_required": attribution,
            "release_archive_decision": str(artifact_decision or source_decision or "metadata_only"),
            "download_allowed": bool(allowed and resolved),
            "preview_allowed": bool(allowed and resolved and extension in {"pdf", "csv", "tsv", "xlsx", "xls", "json", "txt", "md", "html", "htm"}),
            "rows_exported": None,
            "recipe": None,
        })

    for item in artifact_rows:
        artifact_id = str(_pick(item, "artifact_id", "id", default=""))
        if not artifact_id:
            continue
        snapshot_id = str(_pick(item, "snapshot_id", "raw_snapshot_id", default="")) or None
        raw = raw_by_snapshot.get(snapshot_id or "", {})
        source_id = str(_pick(item, "source_id") or _pick(raw, "source_id") or "UNKNOWN")
        source = sources.get(source_id, {})
        source_allowed, source_decision, terms_url, attribution = _source_rights(source)
        decision = _pick(item, "distribution_decision", "release_archive_decision", default=source_decision)
        kind = str(_pick(item, "artifact_type", "distribution_kind", "kind", default="derived_artifact"))
        raw_path = _pick(item, "artifact_path", "file_path", "path", "output_path")
        resolved = _safe_path(raw_path)
        content_type = _pick(item, "content_type", "media_type", "mime_type")
        extension = _file_extension(raw_path, content_type)
        sha = _pick(item, "artifact_sha256", "sha256", "file_sha256") or _sha256(resolved)
        size = _pick(item, "bytes_count", "content_length", "file_size", "bytes")
        if size is None and resolved:
            size = resolved.stat().st_size
        decision_allows = _decision_allows(decision)
        decision_restricts = _decision_restricts(decision)
        inherently_public = any(token in kind.lower() for token in ("recipe", "manifest", "metadata", "documentation"))
        allowed = bool(resolved and not decision_restricts and (decision_allows or inherently_public or source_allowed))
        filename = Path(str(raw_path or artifact_id)).name
        files.append({
            "file_id": _encode_file_id("artifact", artifact_id),
            "record_kind": "artifact",
            "record_id": artifact_id,
            "snapshot_id": snapshot_id,
            "artifact_id": artifact_id,
            "source_id": source_id,
            "title_ru": f"{_human_artifact_kind(kind, 'ru')} · {filename}",
            "title_en": f"{_human_artifact_kind(kind, 'en')} · {filename}",
            "filename": filename,
            "catalog_path": str(raw_path or ""),
            "resolved_path": str(resolved) if resolved else None,
            "exists_locally": bool(resolved),
            "distribution_kind": kind,
            "artifact_type": kind,
            "extension": extension,
            "media_type": _media_type(raw_path, content_type),
            "bytes_count": int(size or 0),
            "bytes_human": _human_bytes(size),
            "sha256": sha,
            "release_year": _pick(item, "release_year") or _pick(raw, "release_year", "source_data_year", "year"),
            "retrieved_at": _pick(item, "created_at", "retrieved_at") or _pick(raw, "retrieved_at", "retrieval_date"),
            "official_url": _pick(item, "source_url", "official_url") or _pick(raw, "source_url", "official_url") or _pick(source, "source_url", "official_url", "url"),
            "terms_url": terms_url,
            "attribution_required": attribution,
            "release_archive_decision": str(decision or "metadata_only"),
            "download_allowed": allowed,
            "preview_allowed": bool(allowed and extension in {"pdf", "csv", "tsv", "xlsx", "xls", "json", "txt", "md", "html", "htm"}),
            "rows_exported": _pick(item, "rows_exported", "row_count", "records_count"),
            "recipe": _pick(item, "recipe", "recipe_path", "transformation_recipe"),
        })

    # Remove exact duplicate distribution paths while preferring explicit artifacts.
    deduplicated: dict[tuple[str, str], dict[str, Any]] = {}
    for item in sorted(files, key=lambda row: 0 if row["record_kind"] == "artifact" else 1):
        key = (item["source_id"], item["catalog_path"] or item["file_id"])
        deduplicated.setdefault(key, item)
    files = list(deduplicated.values())
    files.sort(key=lambda row: (str(row.get("source_id")), str(row.get("release_year") or ""), str(row.get("filename"))))
    return sources, files


def _human_artifact_kind(value: str, lang: str) -> str:
    key = _normalise_decision(value)
    labels = {
        "derived_export": ("Производная выгрузка", "Derived export"),
        "normalized_export": ("Нормализованная выгрузка", "Normalised export"),
        "recipe": ("Рецепт воспроизведения", "Reproduction recipe"),
        "manifest": ("Манифест", "Manifest"),
        "metadata": ("Метаданные", "Metadata"),
        "raw_snapshot": ("Исходный snapshot", "Source snapshot"),
        "extracted_table": ("Извлечённая таблица", "Extracted table"),
    }
    ru, en = labels.get(key, (value.replace("_", " ").title(), value.replace("_", " ").title()))
    return ru if lang == "ru" else en


def _source_observation_stats(conn: sqlite3.Connection) -> dict[str, dict[str, Any]]:
    table = "source_observations"
    columns = _columns(table)
    if not columns:
        return {}
    source_col = _first_existing(columns, "source_id")
    country_col = _first_existing(columns, "iso3", "country_iso3", "country_code")
    year_col = _first_existing(columns, "year", "source_data_year", "data_year")
    if not source_col:
        return {}
    pieces = [f"{_quote_identifier(source_col)} AS source_id", "COUNT(*) AS observation_count"]
    if country_col:
        pieces.append(f"COUNT(DISTINCT {_quote_identifier(country_col)}) AS country_count")
    if year_col:
        pieces.extend([
            f"MIN({_quote_identifier(year_col)}) AS year_min",
            f"MAX({_quote_identifier(year_col)}) AS year_max",
        ])
    sql = f'SELECT {", ".join(pieces)} FROM {_quote_identifier(table)} GROUP BY {_quote_identifier(source_col)}'
    return {str(row["source_id"]): dict(row) for row in conn.execute(sql)}


def _related_indices(conn: sqlite3.Connection) -> dict[str, list[str]]:
    result: dict[str, set[str]] = defaultdict(set)
    for table in ("component_values", "htei_v6_component_values", "training_model_components"):
        columns = _columns(table)
        source_col = _first_existing(columns, "source_id")
        index_col = _first_existing(columns, "index_code", "model_code")
        if not source_col:
            continue
        if index_col:
            query = f'SELECT DISTINCT {_quote_identifier(source_col)}, {_quote_identifier(index_col)} FROM {_quote_identifier(table)} WHERE {_quote_identifier(source_col)} IS NOT NULL'
            for source_id, index_code in conn.execute(query):
                if source_id and index_code:
                    result[str(source_id)].add(str(index_code))
        else:
            label = "HTEI" if table.startswith("htei") else "TRAINING_MODEL"
            query = f'SELECT DISTINCT {_quote_identifier(source_col)} FROM {_quote_identifier(table)} WHERE {_quote_identifier(source_col)} IS NOT NULL'
            for (source_id,) in conn.execute(query):
                if source_id:
                    result[str(source_id)].add(label)
    return {key: sorted(values) for key, values in result.items()}


def _source_catalog(conn: sqlite3.Connection) -> list[dict[str, Any]]:
    source_map, files = _catalog_records(conn)
    stats = _source_observation_stats(conn)
    related = _related_indices(conn)
    file_groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for item in files:
        file_groups[item["source_id"]].append(item)
    source_ids = sorted(set(source_map) | set(file_groups) | set(stats))
    result: list[dict[str, Any]] = []
    for source_id in source_ids:
        source = source_map.get(source_id, {})
        source_files = file_groups.get(source_id, [])
        source_stats = stats.get(source_id, {})
        allowed, decision, terms_url, attribution = _source_rights(source)
        formats = sorted({item["extension"] for item in source_files if item.get("extension")})
        years = [int(item["release_year"]) for item in source_files if str(item.get("release_year") or "").isdigit()]
        year_min = source_stats.get("year_min") or (min(years) if years else None)
        year_max = source_stats.get("year_max") or (max(years) if years else None)
        local_bytes = sum(item["bytes_count"] for item in source_files if item.get("exists_locally"))
        result.append({
            "source_id": source_id,
            "source_name": _pick(source, "source_name", "name", default=source_id),
            "title": _pick(source, "title", "dataset_title", "source_name", default=source_id),
            "owner": _pick(source, "owner", "publisher", "organisation", "organization"),
            "source_url": _pick(source, "source_url", "official_url", "url"),
            "terms_url": terms_url,
            "license_or_terms": _pick(source, "license_or_terms", "license", "terms"),
            "attribution_required": attribution,
            "release_archive_decision": decision,
            "redistribution_allowed": allowed,
            "review_status": _pick(source, "review_status", "license_review_status", default="project_review"),
            "source_role": _infer_source_role(source_id, source),
            "latest_snapshot_id": _pick(source, "latest_snapshot_id"),
            "retrieved_at": _pick(source, "retrieved_at", "latest_retrieved_at"),
            "release_year": _pick(source, "release_year"),
            "distribution_count": len(source_files),
            "downloadable_count": sum(bool(item["download_allowed"]) for item in source_files),
            "formats": formats,
            "local_bytes": local_bytes,
            "local_bytes_human": _human_bytes(local_bytes),
            "observation_count": int(source_stats.get("observation_count") or 0),
            "country_count": int(source_stats.get("country_count") or 0),
            "year_min": year_min,
            "year_max": year_max,
            "related_indices": related.get(source_id, []),
        })
    result.sort(key=lambda row: (-int(row["observation_count"] or 0), -int(row["distribution_count"]), row["source_id"]))
    return result


def _file_by_id(conn: sqlite3.Connection, file_id: str) -> dict[str, Any]:
    kind, identifier = _decode_file_id(file_id)
    _sources, files = _catalog_records(conn)
    for item in files:
        if item["record_kind"] == kind and item["record_id"] == identifier:
            return item
    raise HTTPException(status_code=404, detail="Catalog file not found")


# ---------------------------------------------------------------------------
# Explorer schema and query engine
# ---------------------------------------------------------------------------

_DATASET_BLUEPRINTS: tuple[dict[str, Any], ...] = (
    {
        "id": "index_scores", "table": "index_scores",
        "title_ru": "Итоговые оценки международных индексов", "title_en": "International index scores",
        "description_ru": "Score, место, процентиль и качество данных по индексам, странам и годам.",
        "description_en": "Scores, ranks, percentiles and data quality by index, country and year.",
        "country": ("iso3",), "year": ("year",), "series": ("index_code",), "source": ("source_id",),
        "value_id": ("value_id",), "unit": (),
        "measures": (
            ("score", ("score",), "Оценка", "Score", "index points"),
            ("rank", ("rank",), "Место", "Rank", "rank"),
            ("percentile", ("percentile",), "Процентиль", "Percentile", "percentile"),
            ("data_quality", ("data_quality",), "Качество данных", "Data quality", "share"),
        ),
        "dimensions": (("index_code", ("index_code",), "Индекс", "Index"),),
        "default": {"measure": "score", "index_code": "HTEI", "countries": "RUS,CHN,USA,DEU,KOR,SGP"},
    },
    {
        "id": "component_values", "table": "component_values",
        "title_ru": "Компоненты официальных индексов", "title_en": "Official index components",
        "description_ru": "Исходные и нормированные значения компонентов, из которых строятся индексные оценки.",
        "description_en": "Raw and normalised component values used to construct index scores.",
        "country": ("iso3",), "year": ("year", "source_data_year"), "series": ("component_code",),
        "source": ("source_id",), "value_id": ("value_id",), "unit": ("unit",),
        "series_ru": ("component_name_ru", "name_ru"), "series_en": ("component_name_en", "name_en"),
        "measures": (
            ("score", ("score", "normalized_score"), "Нормированный score", "Normalised score", "points"),
            ("raw_value", ("raw_value", "value"), "Исходное значение", "Raw value", None),
            ("weighted_contribution", ("weighted_contribution",), "Вклад", "Contribution", "points"),
        ),
        "dimensions": (
            ("index_code", ("index_code",), "Индекс", "Index"),
            ("component_code", ("component_code",), "Компонент", "Component"),
        ),
        "default": {"measure": "score", "countries": "RUS,CHN,USA,DEU,KOR,SGP"},
    },
    {
        "id": "source_observations", "table": "source_observations",
        "title_ru": "Исходные наблюдения", "title_en": "Source observations",
        "description_ru": "Нормализованный слой наблюдений до построения итоговых индексов.",
        "description_en": "Normalised observation layer before final index construction.",
        "country": ("iso3", "country_iso3", "country_code"), "year": ("year", "source_data_year", "data_year"),
        "series": ("indicator_code", "series_code", "metric_code"), "source": ("source_id",),
        "value_id": ("value_id", "observation_id"), "unit": ("unit", "unit_name"),
        "series_ru": ("indicator_name_ru", "name_ru", "label_ru"), "series_en": ("indicator_name_en", "name_en", "label_en"),
        "measures": (
            ("value", ("value", "raw_value", "numeric_value"), "Значение", "Value", None),
            ("normalized_score", ("normalized_score", "score"), "Нормированный score", "Normalised score", "points"),
        ),
        "dimensions": (
            ("source_id", ("source_id",), "Источник", "Source"),
            ("indicator_code", ("indicator_code", "series_code", "metric_code"), "Показатель", "Indicator"),
        ),
        "default": {"measure": "value", "countries": "RUS,CHN,USA,DEU,KOR,SGP"},
    },
    {
        "id": "htei_v6_component_values", "table": "htei_v6_component_values",
        "title_ru": "Компоненты HTEI v6", "title_en": "HTEI v6 components",
        "description_ru": "Исходные значения, нормированные scores, веса и вклады компонентов авторского индекса.",
        "description_en": "Raw values, normalised scores, weights and contributions of the original index.",
        "country": ("iso3",), "year": ("release_year", "year"), "series": ("component_code",),
        "source": ("source_id",), "value_id": ("value_id",), "unit": ("unit",),
        "series_ru": ("name_ru", "component_name_ru"), "series_en": ("name_en", "component_name_en"),
        "measures": (
            ("normalized_score", ("normalized_score", "score"), "Нормированный score", "Normalised score", "points"),
            ("raw_value", ("raw_value", "value"), "Исходное значение", "Raw value", None),
            ("weighted_contribution", ("weighted_contribution",), "Вклад в HTEI", "Contribution to HTEI", "points"),
            ("data_lag", ("data_lag",), "Лаг данных", "Data lag", "years"),
        ),
        "dimensions": (
            ("mode", ("mode",), "Режим", "Mode"),
            ("component_code", ("component_code",), "Компонент", "Component"),
            ("source_id", ("source_id",), "Источник", "Source"),
        ),
        "default": {"measure": "normalized_score", "mode": "proxy_extended", "countries": "RUS,CHN,USA,DEU,KOR,SGP"},
    },
    {
        "id": "training_model_scores", "table": "training_model_scores",
        "title_ru": "Модель конкурентоспособности подготовки кадров", "title_en": "Workforce training competitiveness model",
        "description_ru": "Итоговые оценки, места и процентили четырёхблочной модели подготовки технологических кадров.",
        "description_en": "Overall scores, ranks and percentiles for the four-block technology workforce training model.",
        "country": ("iso3",), "year": ("year",), "series": (), "source": (), "value_id": ("value_id",), "unit": (),
        "measures": (
            ("score", ("score",), "Оценка модели", "Model score", "points"),
            ("rank", ("rank",), "Место", "Rank", "rank"),
            ("percentile", ("percentile",), "Процентиль", "Percentile", "percentile"),
            ("data_quality", ("data_quality",), "Качество данных", "Data quality", "share"),
        ),
        "dimensions": (),
        "default": {"measure": "score", "countries": "RUS,CHN,USA,DEU,KOR,SGP"},
    },
    {
        "id": "training_model_components", "table": "training_model_components",
        "title_ru": "Блоки и компоненты модели подготовки кадров", "title_en": "Training-model blocks and components",
        "description_ru": "Компонентная структура институциональной среды, образования, корпоративных стратегий и международной кооперации.",
        "description_en": "Component structure of institutions, education, corporate strategies and international cooperation.",
        "country": ("iso3",), "year": ("year",), "series": ("component_ref", "component_code"),
        "source": ("source_id",), "value_id": ("value_id",), "unit": ("unit",),
        "series_ru": ("component_name_ru",), "series_en": ("component_name_en",),
        "measures": (
            ("normalized_score", ("normalized_score", "score"), "Нормированный score", "Normalised score", "points"),
            ("weighted_contribution", ("weighted_contribution",), "Взвешенный вклад", "Weighted contribution", "points"),
        ),
        "dimensions": (
            ("block_code", ("block_code",), "Блок модели", "Model block"),
            ("component_ref", ("component_ref", "component_code"), "Компонент", "Component"),
            ("source_id", ("source_id",), "Источник", "Source"),
        ),
        "default": {"measure": "normalized_score", "countries": "RUS,CHN,USA,DEU,KOR,SGP"},
    },
    {
        "id": "qs_institution_rankings", "table": "qs_institution_rankings",
        "title_ru": "Инженерно-технологические университеты QS", "title_en": "QS engineering and technology institutions",
        "description_ru": "Позиции и оценки университетов по странам и выпускам рейтинга.",
        "description_en": "Institution positions and scores by country and ranking release.",
        "country": ("iso3",), "year": ("year", "release_year"), "series": ("institution_name", "institution_id"),
        "source": ("source_id",), "value_id": ("value_id", "ranking_id"), "unit": (),
        "series_ru": ("institution_name",), "series_en": ("institution_name",),
        "measures": (
            ("rank", ("rank", "rank_position"), "Место университета", "Institution rank", "rank"),
            ("score", ("score", "overall_score"), "Оценка", "Score", "points"),
        ),
        "dimensions": (),
        "default": {"measure": "rank", "countries": "RUS,CHN,USA,DEU,KOR,SGP"},
    },
    {
        "id": "national_statistics_metrics", "table": "national_statistics_metrics",
        "title_ru": "Национальная статистика", "title_en": "National statistics",
        "description_ru": "Показатели национальных статистических служб, приведённые к единой страново-временной схеме.",
        "description_en": "National statistical-office metrics mapped to a common country-year schema.",
        "country": ("iso3",), "year": ("year", "source_data_year"), "series": ("metric_code", "indicator_code"),
        "source": ("source_id",), "value_id": ("value_id", "metric_id"), "unit": ("unit",),
        "series_ru": ("metric_name_ru", "name_ru"), "series_en": ("metric_name_en", "name_en"),
        "measures": (("value", ("value", "raw_value"), "Значение", "Value", None),),
        "dimensions": (("metric_code", ("metric_code", "indicator_code"), "Показатель", "Metric"),),
        "default": {"measure": "value", "countries": "RUS,CHN,USA,DEU,KOR,SGP"},
    },
    {
        "id": "corporate_metrics", "table": "corporate_metrics",
        "title_ru": "Корпоративные показатели", "title_en": "Corporate metrics",
        "description_ru": "Показатели корпоративной отчётности, используемые для диагностики технологического спроса и участия бизнеса в НИОКР.",
        "description_en": "Corporate disclosure metrics used to diagnose technology demand and business participation in R&D.",
        "country": ("iso3",), "year": ("year", "source_data_year"), "series": ("metric_code", "indicator_code"),
        "source": ("source_id",), "value_id": ("value_id", "metric_id"), "unit": ("unit",),
        "series_ru": ("metric_name_ru", "name_ru"), "series_en": ("metric_name_en", "name_en"),
        "measures": (("value", ("value", "raw_value"), "Значение", "Value", None),),
        "dimensions": (("metric_code", ("metric_code", "indicator_code"), "Показатель", "Metric"),),
        "default": {"measure": "value", "countries": "RUS,CHN,USA,DEU,KOR,SGP"},
    },
)


def _dimension_label(code: str, value: Any, conn: sqlite3.Connection) -> tuple[str, str]:
    text = str(value or "")
    if code == "index_code" and _table_exists(conn, "indices"):
        cols = _columns("indices")
        key = _first_existing(cols, "index_code", "code")
        ru = _first_existing(cols, "name_ru", "title_ru")
        en = _first_existing(cols, "name_en", "title_en")
        if key:
            selected = [key] + [column for column in (ru, en) if column]
            row = conn.execute(
                f'SELECT {", ".join(_quote_identifier(column) for column in selected)} FROM "indices" WHERE {_quote_identifier(key)}=? LIMIT 1',
                (text,),
            ).fetchone()
            if row:
                data = dict(row)
                return str(data.get(ru) or text), str(data.get(en) or data.get(ru) or text)
    mode_labels = {
        "direct_core": ("Строгий слой прямых данных", "Strict direct-data layer"),
        "common_support": ("Основной международный рейтинг", "Core international ranking"),
        "proxy_extended": ("Расширенный международный рейтинг", "Extended international ranking"),
        "asof_diagnostic": ("Диагностический профиль", "Diagnostic profile"),
    }
    block_labels = {
        "institutional_environment": ("Институциональная среда", "Institutional environment"),
        "educational_infrastructure": ("Образовательная инфраструктура", "Educational infrastructure"),
        "corporate_strategies": ("Корпоративные стратегии", "Corporate strategies"),
        "international_cooperation": ("Международное сотрудничество", "International cooperation"),
    }
    if code == "mode" and text in mode_labels:
        return mode_labels[text]
    if code == "block_code" and text in block_labels:
        return block_labels[text]
    if code == "source_id":
        row = conn.execute("SELECT * FROM source_registry WHERE source_id=? LIMIT 1", (text,)).fetchone() if _table_exists(conn, "source_registry") else None
        if row:
            name = _pick(dict(row), "source_name", "name", default=text)
            return str(name), str(name)
    return text.replace("_", " ").strip().title() or text, text.replace("_", " ").strip().title() or text


def _resolved_dataset(conn: sqlite3.Connection, blueprint: dict[str, Any]) -> dict[str, Any] | None:
    table = blueprint["table"]
    columns = _columns(table)
    if not columns:
        return None
    country = _first_existing(columns, *blueprint.get("country", ()))
    year = _first_existing(columns, *blueprint.get("year", ()))
    series = _first_existing(columns, *blueprint.get("series", ()))
    source = _first_existing(columns, *blueprint.get("source", ()))
    value_id = _first_existing(columns, *blueprint.get("value_id", ()))
    unit = _first_existing(columns, *blueprint.get("unit", ()))
    series_ru = _first_existing(columns, *blueprint.get("series_ru", ()))
    series_en = _first_existing(columns, *blueprint.get("series_en", ()))
    if not country:
        return None
    measures: list[dict[str, Any]] = []
    for code, candidates, label_ru, label_en, fixed_unit in blueprint.get("measures", ()):
        column = _first_existing(columns, *candidates)
        if column:
            measures.append({
                "code": code, "column": column, "label_ru": label_ru, "label_en": label_en,
                "fixed_unit": fixed_unit,
            })
    if not measures:
        return None
    dimensions: list[dict[str, Any]] = []
    for code, candidates, label_ru, label_en in blueprint.get("dimensions", ()):
        column = _first_existing(columns, *candidates)
        if not column:
            continue
        query = f'SELECT {_quote_identifier(column)} AS value, COUNT(*) AS n FROM {_quote_identifier(table)} WHERE {_quote_identifier(column)} IS NOT NULL GROUP BY {_quote_identifier(column)} ORDER BY n DESC, value LIMIT 500'
        options = []
        for row in conn.execute(query):
            ru, en = _dimension_label(code, row["value"], conn)
            options.append({"value": str(row["value"]), "label_ru": ru, "label_en": en, "count": int(row["n"])})
        dimensions.append({
            "code": code, "column": column, "label_ru": label_ru, "label_en": label_en,
            "options": options,
        })
    count = int(conn.execute(f'SELECT COUNT(*) FROM {_quote_identifier(table)}').fetchone()[0])
    year_min = year_max = None
    if year:
        year_min, year_max = conn.execute(
            f'SELECT MIN({_quote_identifier(year)}), MAX({_quote_identifier(year)}) FROM {_quote_identifier(table)}'
        ).fetchone()
    default = dict(blueprint.get("default", {}))
    if default.get("measure") not in {item["code"] for item in measures}:
        default["measure"] = measures[0]["code"]
    return {
        "id": blueprint["id"], "table": table,
        "title_ru": blueprint["title_ru"], "title_en": blueprint["title_en"],
        "description_ru": blueprint["description_ru"], "description_en": blueprint["description_en"],
        "country_column": country, "year_column": year, "series_column": series,
        "series_ru_column": series_ru, "series_en_column": series_en,
        "source_column": source, "value_id_column": value_id, "unit_column": unit,
        "measures": measures, "dimensions": dimensions, "row_count": count,
        "year_min": year_min, "year_max": year_max, "default": default,
    }


def _explorer_schema(conn: sqlite3.Connection) -> dict[str, Any]:
    datasets = [item for blueprint in _DATASET_BLUEPRINTS if (item := _resolved_dataset(conn, blueprint))]
    countries: list[dict[str, Any]] = []
    if _table_exists(conn, "countries"):
        columns = _columns("countries")
        select = [column for column in ("iso3", "iso2", "name_ru", "name_en", "region", "income_group", "flag") if column in columns]
        for item in _dicts(conn.execute(f'SELECT {", ".join(_quote_identifier(column) for column in select)} FROM countries ORDER BY name_en, iso3')):
            if item.get("iso3"):
                countries.append(item)
    return {
        "schema_version": "gir-data-explorer-v1",
        "datasets": datasets,
        "countries": countries,
        "max_json_rows": 5000,
        "max_export_rows": 100000,
    }


def _dataset_by_id(conn: sqlite3.Connection, dataset_id: str) -> dict[str, Any]:
    schema = _explorer_schema(conn)
    dataset = next((item for item in schema["datasets"] if item["id"] == dataset_id), None)
    if not dataset:
        raise HTTPException(status_code=404, detail=f"Unknown explorer dataset: {dataset_id}")
    return dataset


def _series_labels(conn: sqlite3.Connection, dataset: dict[str, Any], code: Any, row: dict[str, Any]) -> tuple[str, str]:
    explicit_ru = row.get("series_ru")
    explicit_en = row.get("series_en")
    if explicit_ru or explicit_en:
        return str(explicit_ru or explicit_en or code or ""), str(explicit_en or explicit_ru or code or "")
    if dataset["id"] == "index_scores":
        return _dimension_label("index_code", code, conn)
    if dataset["id"] == "training_model_scores":
        return "Итоговая оценка модели", "Overall model score"
    return _dimension_label("series", code, conn)


def _query_explorer(
    conn: sqlite3.Connection,
    dataset_id: str,
    measure_code: str,
    countries: str | None,
    year_from: int | None,
    year_to: int | None,
    dimensions: dict[str, str],
    x_measure: str | None,
    limit: int,
    count_total: bool = True,
) -> dict[str, Any]:
    dataset = _dataset_by_id(conn, dataset_id)
    measure = next((item for item in dataset["measures"] if item["code"] == measure_code), None)
    if not measure:
        raise HTTPException(status_code=400, detail=f"Measure {measure_code!r} is not available for {dataset_id}")
    x_measure_item = None
    if x_measure:
        x_measure_item = next((item for item in dataset["measures"] if item["code"] == x_measure), None)
        if not x_measure_item:
            raise HTTPException(status_code=400, detail=f"X measure {x_measure!r} is not available for {dataset_id}")
    table = dataset["table"]
    aliases = [
        f't.{_quote_identifier(dataset["country_column"])} AS iso3',
        f't.{_quote_identifier(measure["column"])} AS value',
    ]
    if dataset["year_column"]:
        aliases.append(f't.{_quote_identifier(dataset["year_column"])} AS year')
    else:
        aliases.append("NULL AS year")
    if dataset["series_column"]:
        aliases.append(f't.{_quote_identifier(dataset["series_column"])} AS series_code')
    else:
        aliases.append(f"'{dataset_id}' AS series_code")
    if dataset["series_ru_column"]:
        aliases.append(f't.{_quote_identifier(dataset["series_ru_column"])} AS series_ru')
    else:
        aliases.append("NULL AS series_ru")
    if dataset["series_en_column"]:
        aliases.append(f't.{_quote_identifier(dataset["series_en_column"])} AS series_en')
    else:
        aliases.append("NULL AS series_en")
    if dataset["source_column"]:
        aliases.append(f't.{_quote_identifier(dataset["source_column"])} AS source_id')
    else:
        aliases.append("NULL AS source_id")
    if dataset["value_id_column"]:
        aliases.append(f't.{_quote_identifier(dataset["value_id_column"])} AS value_id')
    else:
        aliases.append("NULL AS value_id")
    if dataset["unit_column"]:
        aliases.append(f't.{_quote_identifier(dataset["unit_column"])} AS unit')
    else:
        unit = measure.get("fixed_unit")
        aliases.append(("? AS unit") if unit else "NULL AS unit")
    if x_measure_item:
        aliases.append(f't.{_quote_identifier(x_measure_item["column"])} AS x_value')
    else:
        aliases.append("NULL AS x_value")
    country_columns = _columns("countries")
    has_countries = _table_exists(conn, "countries") and "iso3" in country_columns
    if has_countries:
        aliases.extend([
            "c.name_ru AS country_ru" if "name_ru" in country_columns else "c.iso3 AS country_ru",
            "c.name_en AS country_en" if "name_en" in country_columns else "c.iso3 AS country_en",
            "c.region AS region" if "region" in country_columns else "NULL AS region",
            "c.income_group AS income_group" if "income_group" in country_columns else "NULL AS income_group",
        ])
    else:
        aliases.extend(["NULL AS country_ru", "NULL AS country_en", "NULL AS region", "NULL AS income_group"])

    select_parameters: list[Any] = []
    filter_parameters: list[Any] = []
    if not dataset["unit_column"] and measure.get("fixed_unit"):
        select_parameters.append(measure["fixed_unit"])
    where = [f't.{_quote_identifier(measure["column"])} IS NOT NULL']
    country_values = [value.strip().upper() for value in str(countries or "").split(",") if value.strip()]
    if country_values:
        placeholders = ",".join("?" for _ in country_values)
        where.append(f't.{_quote_identifier(dataset["country_column"])} IN ({placeholders})')
        filter_parameters.extend(country_values)
    if dataset["year_column"] and year_from is not None:
        where.append(f't.{_quote_identifier(dataset["year_column"])} >= ?')
        filter_parameters.append(year_from)
    if dataset["year_column"] and year_to is not None:
        where.append(f't.{_quote_identifier(dataset["year_column"])} <= ?')
        filter_parameters.append(year_to)
    dimension_map = {item["code"]: item for item in dataset["dimensions"]}
    applied_filters: dict[str, list[str]] = {}
    for code, raw_value in dimensions.items():
        dimension = dimension_map.get(code)
        if not dimension or raw_value in {None, "", "all"}:
            continue
        values = [value.strip() for value in str(raw_value).split(",") if value.strip()]
        if not values:
            continue
        placeholders = ",".join("?" for _ in values)
        where.append(f't.{_quote_identifier(dimension["column"])} IN ({placeholders})')
        filter_parameters.extend(values)
        applied_filters[code] = values
    join = f' LEFT JOIN countries c ON c.iso3=t.{_quote_identifier(dataset["country_column"])}' if has_countries else ""
    base_sql = f'SELECT {", ".join(aliases)} FROM {_quote_identifier(table)} t{join} WHERE {" AND ".join(where)}'
    order = []
    if dataset["country_column"]:
        order.append("iso3")
    if dataset["series_column"]:
        order.append("series_code")
    if dataset["year_column"]:
        order.append("year")
    sql = base_sql + (f' ORDER BY {", ".join(order)}' if order else "") + " LIMIT ?"
    query_parameters = list(select_parameters) + list(filter_parameters) + [max(1, min(int(limit), 100000))]
    raw_rows = _dicts(conn.execute(sql, tuple(query_parameters)))
    total = len(raw_rows)
    if count_total:
        count_sql = f'SELECT COUNT(*) FROM {_quote_identifier(table)} t WHERE {" AND ".join(where)}'
        total = int(conn.execute(count_sql, tuple(filter_parameters)).fetchone()[0])
    output_rows: list[dict[str, Any]] = []
    for item in raw_rows:
        ru, en = _series_labels(conn, dataset, item.get("series_code"), item)
        item["series_ru"] = ru
        item["series_en"] = en
        if not item.get("country_ru"):
            item["country_ru"] = item.get("iso3")
        if not item.get("country_en"):
            item["country_en"] = item.get("iso3")
        try:
            item["value"] = float(item["value"])
        except (TypeError, ValueError):
            continue
        if item.get("x_value") is not None:
            try:
                item["x_value"] = float(item["x_value"])
            except (TypeError, ValueError):
                item["x_value"] = None
        output_rows.append(item)
    values = [item["value"] for item in output_rows]
    years = [int(item["year"]) for item in output_rows if item.get("year") is not None and str(item["year"]).lstrip("-").isdigit()]
    source_ids = sorted({str(item["source_id"]) for item in output_rows if item.get("source_id")})
    source_details = []
    if source_ids and _table_exists(conn, "source_registry"):
        placeholders = ",".join("?" for _ in source_ids)
        for source in _dicts(conn.execute(f"SELECT * FROM source_registry WHERE source_id IN ({placeholders})", tuple(source_ids))):
            source_details.append({
                "source_id": source.get("source_id"),
                "source_name": _pick(source, "source_name", "name"),
                "owner": _pick(source, "owner", "publisher"),
                "source_url": _pick(source, "source_url", "official_url", "url"),
            })
    warnings_ru: list[str] = []
    warnings_en: list[str] = []
    if total > len(output_rows):
        warnings_ru.append(f"Интерактивное представление ограничено {len(output_rows):,} строками из {total:,}. Полная отфильтрованная выборка доступна через CSV-экспорт.".replace(",", " "))
        warnings_en.append(f"The interactive view is limited to {len(output_rows):,} of {total:,} rows. The full filtered selection is available as CSV.")
    units = sorted({str(item["unit"]) for item in output_rows if item.get("unit")})
    if len(units) > 1:
        warnings_ru.append("Выборка содержит несколько единиц измерения. Для содержательного графика сузьте показатель или серию.")
        warnings_en.append("The selection contains multiple units. Narrow the indicator or series before interpreting the chart.")
    return {
        "dataset": dataset_id,
        "dataset_meta": {key: dataset[key] for key in ("id", "title_ru", "title_en", "description_ru", "description_en")},
        "measure_meta": measure,
        "x_measure_meta": x_measure_item,
        "filters": {"countries": country_values, "year_from": year_from, "year_to": year_to, **applied_filters},
        "rows": output_rows,
        "sources": source_details,
        "units": units,
        "summary": {
            "returned_rows": len(output_rows), "total_rows": total,
            "country_count": len({item.get("iso3") for item in output_rows if item.get("iso3")}),
            "series_count": len({item.get("series_code") for item in output_rows if item.get("series_code")}),
            "year_min": min(years) if years else None, "year_max": max(years) if years else None,
            "value_min": min(values) if values else None, "value_max": max(values) if values else None,
            "value_mean": (sum(values) / len(values)) if values else None,
        },
        "warnings_ru": warnings_ru, "warnings_en": warnings_en,
    }


def _query_dimensions_from_request(request: Request, dataset: dict[str, Any]) -> dict[str, str]:
    return {
        dimension["code"]: request.query_params.get(dimension["code"], "")
        for dimension in dataset["dimensions"]
        if request.query_params.get(dimension["code"])
    }


# ---------------------------------------------------------------------------
# HTTP routes
# ---------------------------------------------------------------------------

@router.get("/data-lab", include_in_schema=False)
def data_lab_page() -> FileResponse:
    page = STATIC_DIR / "data-lab.html"
    if not page.exists():
        raise HTTPException(status_code=500, detail="Data Lab frontend is missing")
    return FileResponse(page, media_type="text/html; charset=utf-8")


@router.get("/data", include_in_schema=False)
def data_lab_alias() -> RedirectResponse:
    return RedirectResponse(url="/data-lab", status_code=307)


@router.get("/api/data-catalog/summary")
def catalog_summary() -> dict[str, Any]:
    with _connect() as conn:
        sources, files = _catalog_records(conn)
        formats: dict[str, dict[str, Any]] = defaultdict(lambda: {"count": 0, "bytes": 0})
        access: dict[str, int] = defaultdict(int)
        unique_paths: dict[str, int] = {}
        for item in files:
            formats[item["extension"]]["count"] += 1
            formats[item["extension"]]["bytes"] += int(item["bytes_count"] or 0)
            if item["download_allowed"]:
                access["downloadable"] += 1
            elif item["exists_locally"]:
                access["restricted"] += 1
            else:
                access["metadata_only"] += 1
            if item.get("resolved_path"):
                unique_paths[item["resolved_path"]] = int(item["bytes_count"] or 0)
        database_rows: dict[str, int] = {}
        for table in (
            "countries", "index_scores", "component_values", "source_observations",
            "htei_v6_profiles", "htei_v6_component_values", "training_model_scores",
            "training_model_components", "qs_institution_rankings", "national_statistics_metrics",
            "corporate_metrics",
        ):
            if _table_exists(conn, table):
                database_rows[table] = int(conn.execute(f'SELECT COUNT(*) FROM {_quote_identifier(table)}').fetchone()[0])
        year_values: list[int] = []
        for table in ("source_observations", "index_scores"):
            columns = _columns(table)
            year_col = _first_existing(columns, "year", "source_data_year", "data_year")
            if year_col:
                minimum, maximum = conn.execute(
                    f'SELECT MIN({_quote_identifier(year_col)}), MAX({_quote_identifier(year_col)}) FROM {_quote_identifier(table)}'
                ).fetchone()
                for value in (minimum, maximum):
                    if value is not None:
                        try:
                            year_values.append(int(value))
                        except (TypeError, ValueError):
                            pass
        schema = _explorer_schema(conn)
        return {
            "schema_version": "gir-data-catalog-v1",
            "source_count": len(sources),
            "catalog_distributions": len(files),
            "registered_snapshots": len(_raw_rows(conn)),
            "reproducibility_artifacts": len(_artifact_rows(conn)),
            "total_local_bytes": sum(unique_paths.values()),
            "formats": [
                {"key": key, **value}
                for key, value in sorted(formats.items(), key=lambda item: (-item[1]["count"], item[0]))
            ],
            "access": [{"key": key, "count": value} for key, value in sorted(access.items())],
            "distribution_kinds": [
                {"key": key, "count": value}
                for key, value in sorted(
                    defaultdict(int, {kind: sum(1 for item in files if item["distribution_kind"] == kind) for kind in {item["distribution_kind"] for item in files}}).items(),
                    key=lambda item: (-item[1], item[0]),
                )
            ],
            "database_rows": database_rows,
            "data_year_min": min(year_values) if year_values else None,
            "data_year_max": max(year_values) if year_values else None,
            "explorer_dataset_count": len(schema["datasets"]),
        }


@router.get("/api/data-catalog/sources")
def catalog_sources(
    q: str = "",
    role: str = "all",
    access: str = "all",
) -> dict[str, Any]:
    with _connect() as conn:
        sources = _source_catalog(conn)
    query = q.strip().lower()
    if query:
        sources = [
            item for item in sources
            if query in " ".join(
                str(value or "") for value in (
                    item.get("source_id"), item.get("source_name"), item.get("owner"),
                    item.get("title"), " ".join(item.get("related_indices") or []),
                )
            ).lower()
        ]
    if role != "all":
        sources = [item for item in sources if item.get("source_role") == role]
    if access == "downloadable":
        sources = [item for item in sources if item.get("redistribution_allowed") or item.get("downloadable_count")]
    elif access == "restricted":
        sources = [item for item in sources if not item.get("redistribution_allowed")]
    return {"sources": sources, "total": len(sources)}


@router.get("/api/data-catalog/sources/{source_id}")
def catalog_source(source_id: str) -> dict[str, Any]:
    with _connect() as conn:
        sources = _source_catalog(conn)
        source = next((item for item in sources if item["source_id"] == source_id), None)
        if not source:
            raise HTTPException(status_code=404, detail="Source not found")
        _source_map, files = _catalog_records(conn)
        source["distributions"] = [item for item in files if item["source_id"] == source_id]
        return source


@router.get("/api/data-catalog/files")
def catalog_files(
    q: str = "",
    source_id: str = "all",
    extension: str = "all",
    kind: str = "all",
    access: str = "all",
    year: str = "",
    page: int = Query(1, ge=1),
    page_size: int = Query(30, ge=1, le=100),
) -> dict[str, Any]:
    with _connect() as conn:
        _sources, files = _catalog_records(conn)
    query = q.strip().lower()
    if query:
        files = [
            item for item in files
            if query in " ".join(str(item.get(key) or "") for key in ("filename", "title_ru", "title_en", "source_id", "snapshot_id", "artifact_id")).lower()
        ]
    if source_id != "all":
        files = [item for item in files if item["source_id"] == source_id]
    if extension != "all":
        files = [item for item in files if item["extension"] == extension]
    if kind != "all":
        files = [item for item in files if item["distribution_kind"] == kind]
    if year:
        files = [item for item in files if str(item.get("release_year") or "") == str(year)]
    if access == "downloadable":
        files = [item for item in files if item["download_allowed"]]
    elif access == "restricted":
        files = [item for item in files if item["exists_locally"] and not item["download_allowed"]]
    elif access == "metadata_only":
        files = [item for item in files if not item["exists_locally"]]
    total = len(files)
    pages = max(1, (total + page_size - 1) // page_size)
    page = min(page, pages)
    start = (page - 1) * page_size
    return {"files": files[start:start + page_size], "total": total, "page": page, "page_size": page_size, "pages": pages}


@router.get("/api/data-catalog/files/{file_id}")
def catalog_file_detail(file_id: str) -> dict[str, Any]:
    with _connect() as conn:
        item = _file_by_id(conn, file_id)
        _sources, files = _catalog_records(conn)
        alternatives = [
            candidate for candidate in files
            if candidate["file_id"] != file_id
            and candidate["source_id"] == item["source_id"]
            and (
                candidate.get("snapshot_id") == item.get("snapshot_id")
                or (candidate.get("release_year") is not None and candidate.get("release_year") == item.get("release_year"))
            )
        ][:20]
        public = {key: value for key, value in item.items() if key != "resolved_path"}
        public["alternatives"] = alternatives
        public["content_url"] = f"/api/data-catalog/files/{file_id}/content"
        public["inline_url"] = f"/api/data-catalog/files/{file_id}/content?inline=true"
        return public


@router.get("/api/data-catalog/files/{file_id}/content")
def catalog_file_content(file_id: str, inline: bool = False) -> FileResponse:
    with _connect() as conn:
        item = _file_by_id(conn, file_id)
    if not item["download_allowed"]:
        raise HTTPException(
            status_code=403,
            detail={
                "code": "RAW_REDISTRIBUTION_RESTRICTED",
                "message_ru": "Файл зарегистрирован в доказательной базе, но не разрешён для публичного распространения. Используйте официальный URL или связанную производную выгрузку.",
                "message_en": "The file is registered in the evidence base but is not approved for public redistribution. Use the official URL or an approved derived export.",
                "official_url": item.get("official_url"),
            },
        )
    path = _safe_path(item.get("catalog_path"))
    if not path:
        raise HTTPException(status_code=404, detail="Registered file is not present in this deployment")
    disposition = "inline" if inline else "attachment"
    return FileResponse(path, media_type=item["media_type"], filename=item["filename"], content_disposition_type=disposition)


@router.get("/api/data-catalog/files/{file_id}/preview")
def catalog_file_preview(file_id: str) -> dict[str, Any]:
    with _connect() as conn:
        item = _file_by_id(conn, file_id)
    if not item["preview_allowed"]:
        raise HTTPException(
            status_code=403,
            detail={
                "code": "PREVIEW_NOT_AVAILABLE",
                "message_ru": "Предпросмотр недоступен из-за режима распространения или формата файла.",
                "message_en": "Preview is unavailable because of the distribution policy or file format.",
                "official_url": item.get("official_url"),
            },
        )
    path = _safe_path(item.get("catalog_path"))
    if not path:
        raise HTTPException(status_code=404, detail="Registered file is not present in this deployment")
    extension = item["extension"]
    if extension == "pdf":
        return {"preview_type": "pdf", "inline_url": f"/api/data-catalog/files/{file_id}/content?inline=true"}
    if extension in {"csv", "tsv"}:
        delimiter = "\t" if extension == "tsv" else ","
        text = _read_text(path, 1024 * 1024)
        sample = text[:8192]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            delimiter = dialect.delimiter
        except csv.Error:
            pass
        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        rows = []
        for index, row in enumerate(reader):
            rows.append(row)
            if index >= 50:
                break
        columns = rows[0] if rows else []
        return {"preview_type": "table", "columns": columns, "rows": rows[1:], "truncated": len(rows) >= 51}
    if extension in {"xlsx", "xls"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - environment dependent
            raise HTTPException(status_code=501, detail="Spreadsheet preview requires openpyxl") from exc
        if extension == "xls":
            raise HTTPException(status_code=415, detail="Legacy XLS preview is not supported; download the approved file instead")
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        values = []
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            values.append([_json_safe_cell(cell) for cell in row])
            if index >= 50:
                break
        columns = [str(value or "") for value in (values[0] if values else [])]
        return {"preview_type": "table", "sheet": sheet.title, "sheets": workbook.sheetnames, "columns": columns, "rows": values[1:], "truncated": sheet.max_row > 51}
    if extension == "json":
        with path.open("r", encoding="utf-8-sig", errors="replace") as stream:
            payload = json.load(stream)
        return {"preview_type": "json", "json": _truncate_json(payload)}
    text = _read_text(path, 200_000)
    return {"preview_type": "text", "text": text, "truncated": path.stat().st_size > 200_000}


def _read_text(path: Path, limit: int) -> str:
    raw = path.read_bytes()[:limit]
    for encoding in ("utf-8-sig", "utf-8", "cp1251", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _json_safe_cell(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def _truncate_json(value: Any, depth: int = 0) -> Any:
    if depth > 5:
        return "…"
    if isinstance(value, dict):
        return {str(key): _truncate_json(item, depth + 1) for key, item in list(value.items())[:100]}
    if isinstance(value, list):
        return [_truncate_json(item, depth + 1) for item in value[:100]]
    return _json_safe_cell(value)


@router.get("/api/data-catalog/manifest.csv")
def catalog_manifest() -> StreamingResponse:
    with _connect() as conn:
        _sources, files = _catalog_records(conn)
    stream = io.StringIO()
    columns = [
        "file_id", "source_id", "distribution_kind", "filename", "extension", "media_type",
        "bytes_count", "sha256", "release_year", "retrieved_at", "download_allowed",
        "release_archive_decision", "official_url", "snapshot_id", "artifact_id",
    ]
    writer = csv.DictWriter(stream, fieldnames=columns, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(files)
    content = "\ufeff" + stream.getvalue()
    return StreamingResponse(
        iter([content.encode("utf-8")]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="gir_data_catalog_manifest.csv"'},
    )


@router.get("/api/data-catalog/dcat.jsonld")
def catalog_dcat() -> JSONResponse:
    with _connect() as conn:
        sources = _source_catalog(conn)
        _source_map, files = _catalog_records(conn)
    by_source: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for item in files:
        by_source[item["source_id"]].append(item)
    graph = []
    for source in sources:
        source_id = source["source_id"]
        distributions = []
        for item in by_source.get(source_id, []):
            distribution = {
                "@type": "dcat:Distribution",
                "@id": f"urn:gir:distribution:{item['file_id']}",
                "dct:title": item["title_en"],
                "dct:format": item["extension"],
                "dcat:mediaType": item["media_type"],
                "dcat:byteSize": item["bytes_count"],
                "spdx:checksum": {"@type": "spdx:Checksum", "spdx:algorithm": "spdx:checksumAlgorithm_sha256", "spdx:checksumValue": item.get("sha256")},
                "dct:rights": item.get("release_archive_decision"),
            }
            if item["download_allowed"]:
                distribution["dcat:downloadURL"] = f"/api/data-catalog/files/{item['file_id']}/content"
            elif item.get("official_url"):
                distribution["dcat:accessURL"] = item["official_url"]
            distributions.append(distribution)
        graph.append({
            "@type": "dcat:Dataset",
            "@id": f"urn:gir:dataset:{source_id}",
            "dct:identifier": source_id,
            "dct:title": source["source_name"],
            "dct:publisher": source.get("owner"),
            "dct:temporal": {"start": source.get("year_min"), "end": source.get("year_max")},
            "dcat:distribution": distributions,
        })
    payload = {
        "@context": {
            "dcat": "http://www.w3.org/ns/dcat#", "dct": "http://purl.org/dc/terms/",
            "spdx": "http://spdx.org/rdf/terms#",
        },
        "@type": "dcat:Catalog",
        "@id": "urn:gir:catalog:data",
        "dct:title": "Global Index Research data catalog",
        "dcat:dataset": graph,
    }
    return JSONResponse(payload, media_type="application/ld+json")


@router.get("/api/data-explorer/schema")
def explorer_schema() -> dict[str, Any]:
    with _connect() as conn:
        return _explorer_schema(conn)


@router.get("/api/data-explorer/query")
def explorer_query(
    request: Request,
    dataset: str,
    measure: str,
    countries: str | None = None,
    year_from: int | None = None,
    year_to: int | None = None,
    x_measure: str | None = None,
    limit: int = Query(5000, ge=1, le=5000),
) -> dict[str, Any]:
    with _connect() as conn:
        resolved = _dataset_by_id(conn, dataset)
        dimensions = _query_dimensions_from_request(request, resolved)
        return _query_explorer(conn, dataset, measure, countries, year_from, year_to, dimensions, x_measure, limit)


@router.get("/api/data-explorer/export.csv")
def explorer_export(
    request: Request,
    dataset: str,
    measure: str,
    countries: str | None = None,
    year_from: int | None = None,
    year_to: int | None = None,
    x_measure: str | None = None,
) -> StreamingResponse:
    with _connect() as conn:
        resolved = _dataset_by_id(conn, dataset)
        dimensions = _query_dimensions_from_request(request, resolved)
        payload = _query_explorer(conn, dataset, measure, countries, year_from, year_to, dimensions, x_measure, 100000, count_total=False)
    stream = io.StringIO()
    fields = ["iso3", "country_ru", "country_en", "region", "income_group", "year", "series_code", "series_ru", "series_en", "value", "x_value", "unit", "source_id", "value_id"]
    writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(payload["rows"])
    filename = f"gir_{re.sub(r'[^a-zA-Z0-9_-]+', '_', dataset)}_{re.sub(r'[^a-zA-Z0-9_-]+', '_', measure)}.csv"
    return StreamingResponse(
        iter([("\ufeff" + stream.getvalue()).encode("utf-8")]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/api/data-explorer/datasets/{dataset_id}/metadata.json")
def explorer_metadata(dataset_id: str) -> JSONResponse:
    with _connect() as conn:
        dataset = _dataset_by_id(conn, dataset_id)
    columns = [
        {"name": "iso3", "titles": {"ru": "Код страны", "en": "Country code"}, "datatype": "string"},
        {"name": "year", "titles": {"ru": "Год", "en": "Year"}, "datatype": "integer"},
        {"name": "series_code", "titles": {"ru": "Код серии", "en": "Series code"}, "datatype": "string"},
        {"name": "value", "titles": {"ru": "Значение", "en": "Value"}, "datatype": "number"},
        {"name": "unit", "titles": {"ru": "Единица", "en": "Unit"}, "datatype": "string"},
        {"name": "source_id", "titles": {"ru": "Источник", "en": "Source"}, "datatype": "string"},
        {"name": "value_id", "titles": {"ru": "Идентификатор значения", "en": "Value identifier"}, "datatype": "string"},
    ]
    payload = {
        "@context": "http://www.w3.org/ns/csvw",
        "url": f"/api/data-explorer/export.csv?dataset={dataset_id}&measure={dataset['default']['measure']}",
        "dc:title": dataset["title_en"],
        "dc:description": dataset["description_en"],
        "tableSchema": {"columns": columns},
        "gir:sourceTable": dataset["table"],
        "gir:availableMeasures": dataset["measures"],
        "gir:availableDimensions": dataset["dimensions"],
    }
    return JSONResponse(payload, media_type="application/json")


@router.get("/api/data-explorer/provenance/{value_id:path}")
def explorer_provenance(value_id: str) -> dict[str, Any]:
    with _connect() as conn:
        for blueprint in _DATASET_BLUEPRINTS:
            dataset = _resolved_dataset(conn, blueprint)
            if not dataset or not dataset["value_id_column"]:
                continue
            table = dataset["table"]
            row = conn.execute(
                f'SELECT * FROM {_quote_identifier(table)} WHERE {_quote_identifier(dataset["value_id_column"])}=? LIMIT 1',
                (value_id,),
            ).fetchone()
            if not row:
                continue
            record = dict(row)
            source_id = record.get(dataset["source_column"]) if dataset["source_column"] else None
            source = None
            if source_id and _table_exists(conn, "source_registry"):
                source_row = conn.execute("SELECT * FROM source_registry WHERE source_id=? LIMIT 1", (source_id,)).fetchone()
                source = dict(source_row) if source_row else None
            snapshot = None
            if source_id and _table_exists(conn, "raw_snapshots"):
                raw_columns = _columns("raw_snapshots")
                if "source_id" in raw_columns:
                    order_col = _first_existing(raw_columns, "retrieved_at", "release_year", "snapshot_id")
                    order_sql = f' ORDER BY {_quote_identifier(order_col)} DESC' if order_col else ""
                    raw_row = conn.execute(f'SELECT * FROM raw_snapshots WHERE source_id=?{order_sql} LIMIT 1', (source_id,)).fetchone()
                    snapshot = dict(raw_row) if raw_row else None
            return {
                "value_id": value_id,
                "dataset": dataset["id"],
                "table": table,
                "record": record,
                "source": source,
                "snapshot": snapshot,
                "formula_version": _pick(record, "formula_version", "methodology_version"),
                "transform_id": _pick(record, "transform_id", "transformation_id", "transformation_run_id"),
                "quality_flag": _pick(record, "quality_flag", "data_quality", "proxy_status"),
            }
    raise HTTPException(status_code=404, detail="Value provenance not found")
