from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import re
import subprocess
import uuid
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from statistics import median
from typing import Any

import openpyxl
import pycountry
import requests
from babel import Locale
from charset_normalizer import from_bytes

try:
    import pdfplumber
except Exception:  # pragma: no cover - import is validated by the release gate.
    pdfplumber = None

from .config import DATA_DIR, DB_PATH
from .db import connect, init_db, source_role_for, upsert_many
from .htei_release_finalization import ensure_htei_asof_release

RAW_DIR = DATA_DIR / "raw"
SNAPSHOT_RELEASE_YEAR = 2025
LATEST_RELEASE_YEAR = 2026
HISTORICAL_START_YEAR = 1990
SOURCE_TIMEOUT = 180

UNDP_HDI_CSV_URL = "https://hdr.undp.org/sites/default/files/2025_HDR/HDR25_Composite_indices_complete_time_series.csv"
WIPO_GII_XLSX_URL = "https://tind.wipo.int/record/58882/files/wipo-pub-2000-2025-gii-tech1.xlsx"
ITU_IDI_XLSX_URL = "https://www.itu.int/itu-d/reports/statistics/wp-content/uploads/sites/5/2025/06/IDIDataset_2025.xlsx"
PORTULANS_GTCI_PDF_URL = "https://portulansinstitute.org/wp-content/uploads/2025/11/GTCI_2025_report.pdf"
WORLD_BANK_API_BASE = "https://api.worldbank.org/v2"
WORLD_BANK_COUNTRIES_URL = f"{WORLD_BANK_API_BASE}/country?format=json&per_page=400"
UIS_DATA_API_BASE = "https://api.uis.unesco.org/api/public"
UIS_DEFAULT_VERSION_URL = f"{UIS_DATA_API_BASE}/versions/default"
UIS_STEM_INDICATOR = "FOSGP.5T8.F500600700"
UIS_STEM_DATA_URL = f"{UIS_DATA_API_BASE}/data/indicators"
QS_OFFICIAL_URL = "https://www.topuniversities.com/university-subject-rankings/engineering-technology"
QS_YEARS = [2023, 2024, 2025, 2026]
GII_XLSX_URLS = {
    2022: "https://www.wipo.int/edocs/pubdocs/en/wipo-pub-2000-2022-tech1.xlsx",
    2023: "https://www.wipo.int/edocs/pubdocs/en/wipo-pub-2000-2023-tech1.xlsx",
    2024: "https://www.wipo.int/edocs/pubdocs/en/wipo-pub-2000-2024-tech1.xlsx",
    2025: WIPO_GII_XLSX_URL,
}
ILOSTAT_HTEI_URLS = {
    "employment_by_activity": "https://rplumber.ilo.org/data/indicator/?id=EMP_TEMP_SEX_ECO_NB_A",
    "employment_by_occupation": "https://rplumber.ilo.org/data/indicator/?id=EMP_TEMP_SEX_OCU_NB_A",
    "employment_by_activity_occupation": "https://rplumber.ilo.org/data/indicator/?id=EMP_TEMP_ECO_OCU_NB_A",
}
OECD_HTEI_URLS = {
    "msti": "https://sdmx.oecd.org/public/rest/v1/data/OECD.STI.STP,DSD_MSTI@DF_MSTI/.?startPeriod=1990&endPeriod=2025",
    "rd_personnel": "https://sdmx.oecd.org/public/rest/v1/data/OECD.STI.STP,DSD_RDS_PERS@DF_PERS_FUNC/.?startPeriod=1990&endPeriod=2025",
    "berd_industry": "https://sdmx.oecd.org/public/rest/v1/data/OECD.STI.STP,DSD_RDS_BERD@DF_BERD_INDU/.?startPeriod=1990&endPeriod=2025",
}
EUROSTAT_HTEI_URLS = {
    "htec_emp_nat2": "https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/htec_emp_nat2?format=JSON&lang=EN&sinceTimePeriod=1997&sex=T&nace_r2=HTC&unit=PC_EMP",
    "hrst_st_nsec2": "https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/hrst_st_nsec2?format=JSON&lang=EN&sinceTimePeriod=1997&category=HRSTO&nace_r2=TOTAL&age=Y15-74&unit=THS_PER",
    "lfsa_egan2": "https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/lfsa_egan2?format=JSON&lang=EN&sinceTimePeriod=1997&sex=T&age=Y15-74&nace_r2=TOTAL&unit=THS_PER",
    "rd_p_persocc": "https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/rd_p_persocc?format=JSON&lang=EN&sinceTimePeriod=1997&sectperf=TOTAL&prof_pos=RSE&prof_pos=TEC&sex=T&unit=FTE",
}
HTEI_V3_WEIGHTS = {
    "HT_EMPLOYMENT_SHARE": 0.22,
    "HIGH_TECH_OCCUPATIONS": 0.18,
    "RND_PERSONNEL": 0.18,
    "STEM_PIPELINE": 0.16,
    "TECH_OUTPUTS": 0.14,
    "CORPORATE_STRATEGY_AND_DEMAND": 0.12,
}
HTEI_SOURCE_PRIORITY = {
    "NATIONAL_STATS": 1,
    "EUROSTAT_NATIONAL_STATS": 2,
    "OECD": 3,
    "ILOSTAT": 4,
    "WORLD_BANK_UIS": 5,
    "CORPORATE_REPORTS": 6,
    "SPECIALIZED_RATINGS": 7,
}
HTEI_RELEASE_MIN_WEIGHT = 0.55
HTEI_RELEASE_MIN_COMPONENTS = 3
HTEI_RELEASE_MIN_SOURCE_GROUPS = 2
TRAINING_MODEL_VERSION = "training-system-competitiveness-v1"
TRAINING_MODEL_BLOCKS = [
    {
        "code": "institutional_environment",
        "name_ru": "Институциональная среда",
        "name_en": "Institutional environment",
        "weight": 0.25,
        "components": [("GTCI", "ENABLE"), ("GII", "INST"), ("IDI", "MEANING")],
    },
    {
        "code": "educational_infrastructure",
        "name_ru": "Образовательная инфраструктура",
        "name_en": "Educational infrastructure",
        "weight": 0.30,
        "components": [("HCI", "SCHOOL"), ("HCI", "LEARN"), ("QS_ET", "TOP_COUNT"), ("HTEI", "STEM_PIPELINE")],
    },
    {
        "code": "corporate_strategies",
        "name_ru": "Корпоративные стратегии",
        "name_en": "Corporate strategies",
        "weight": 0.25,
        "components": [("GII", "BUS"), ("HTEI", "CORPORATE_STRATEGY_AND_DEMAND"), ("HTEI", "TECH_OUTPUTS")],
    },
    {
        "code": "international_cooperation",
        "name_ru": "Международное сотрудничество",
        "name_en": "International cooperation",
        "weight": 0.20,
        "components": [("GTCI", "ATTRACT"), ("QS_ET", "QS_SCORE"), ("GII", "KTO")],
    },
]

HEADERS = {"User-Agent": "GIIP release data ingestion/1.0 (+official source archival)"}

CENTRAL_ASIA_ISO3 = {"KAZ", "KGZ", "TJK", "TKM", "UZB"}
ECA_TO_EUROPE_EXCEPTIONS = CENTRAL_ASIA_ISO3
COUNTRY_METADATA_FALLBACK_REGION = "Other / not classified"
COUNTRY_METADATA_FALLBACK_INCOME = "Not classified"


@dataclass(frozen=True)
class Snapshot:
    snapshot_id: str
    source_id: str
    release_year: int
    retrieved_at: str
    source_url: str
    raw_snapshot_path: str
    raw_snapshot_sha256: str
    content_type: str
    bytes_count: int
    license_or_terms: str


@dataclass(frozen=True)
class Transform:
    transformation_run_id: str
    transform_id: str
    source_id: str
    snapshot_id: str
    started_at: str
    completed_at: str
    code_version: str
    rows_loaded: int
    notes: str


SOURCE_DEFS: dict[str, dict[str, Any]] = {
    "UNDP_HDR": {
        "source_name": "UNDP Human Development Reports Data Center",
        "owner": "UNDP",
        "url": "https://hdr.undp.org/data-center",
        "source_url": UNDP_HDI_CSV_URL,
        "access_mode": "official downloadable CSV",
        "update_frequency": "annual",
        "license_or_terms": "UNDP Human Development Reports public data center terms",
        "automation_status": "implemented_official_download",
        "release_year": 2025,
    },
    "WORLD_BANK_HCI": {
        "source_name": "World Bank Human Capital Index API",
        "owner": "World Bank",
        "url": "https://www.worldbank.org/en/publication/human-capital",
        "source_url": f"{WORLD_BANK_API_BASE}/source/63",
        "access_mode": "official REST API",
        "update_frequency": "irregular",
        "license_or_terms": "World Bank Open Data terms",
        "automation_status": "implemented_official_api",
        "release_year": 2020,
    },
    "WORLD_BANK_COUNTRIES": {
        "source_name": "World Bank Countries API metadata",
        "owner": "World Bank",
        "url": "https://api.worldbank.org/v2/country",
        "source_url": WORLD_BANK_COUNTRIES_URL,
        "access_mode": "official REST API",
        "update_frequency": "regular",
        "license_or_terms": "World Bank Open Data terms",
        "automation_status": "implemented_official_api",
        "release_year": 2026,
        "source_role": "reference_metadata",
    },
    "PORTULANS_GTCI": {
        "source_name": "Global Talent Competitiveness Index 2025 report",
        "owner": "INSEAD / Portulans Institute",
        "url": "https://portulansinstitute.org/reports/",
        "source_url": PORTULANS_GTCI_PDF_URL,
        "access_mode": "official PDF report extraction",
        "update_frequency": "annual",
        "license_or_terms": "INSEAD and Portulans Institute report terms; platform stores extracted country rank and score with provenance",
        "automation_status": "implemented_official_pdf_extract",
        "release_year": 2025,
    },
    "WIPO_GII": {
        "source_name": "WIPO Global Innovation Index 2025 database",
        "owner": "WIPO",
        "url": "https://www.wipo.int/global_innovation_index/en/",
        "source_url": WIPO_GII_XLSX_URL,
        "access_mode": "official downloadable XLSX",
        "update_frequency": "annual",
        "license_or_terms": "WIPO Global Innovation Index data publication terms",
        "automation_status": "implemented_official_download",
        "release_year": 2025,
    },
    "ITU_IDI": {
        "source_name": "ITU ICT Development Index 2025 dataset",
        "owner": "ITU",
        "url": "https://www.itu.int/itu-d/reports/statistics/idi2025/",
        "source_url": ITU_IDI_XLSX_URL,
        "access_mode": "official downloadable XLSX",
        "update_frequency": "annual",
        "license_or_terms": "ITU dataset terms",
        "automation_status": "implemented_official_download",
        "release_year": 2025,
    },
    "WORLD_BANK_HTEI": {
        "source_name": "World Bank indicators for high-technology workforce capacity",
        "owner": "World Bank",
        "url": "https://api.worldbank.org/v2/",
        "source_url": f"{WORLD_BANK_API_BASE}/indicator",
        "access_mode": "official REST API",
        "update_frequency": "regular",
        "license_or_terms": "World Bank Open Data terms",
        "automation_status": "implemented_official_api",
        "release_year": 2025,
    },
    "UIS_HTEI": {
        "source_name": "UNESCO UIS Data API: STEM tertiary graduates",
        "owner": "UNESCO Institute for Statistics",
        "url": "https://databrowser.uis.unesco.org/",
        "source_url": UIS_STEM_DATA_URL,
        "access_mode": "official version-pinned UIS Data API",
        "update_frequency": "regular",
        "license_or_terms": "UNESCO UIS Open Data terms",
        "automation_status": "implemented_official_api",
        "release_year": 2026,
    },
    "ILOSTAT_HTEI": {
        "source_name": "ILOSTAT bulk labour force indicators for HTEI",
        "owner": "International Labour Organization",
        "url": "https://ilostat.ilo.org/data/",
        "source_url": "https://rplumber.ilo.org/files/website/bulk/indicator.html",
        "access_mode": "official bulk CSV API",
        "update_frequency": "regular",
        "license_or_terms": "ILOSTAT bulk data terms",
        "automation_status": "implemented_official_bulk_csv",
        "release_year": 2025,
    },
    "OECD_HTEI": {
        "source_name": "OECD Data Explorer SDMX datasets for R&D and business technology capacity",
        "owner": "OECD",
        "url": "https://data-explorer.oecd.org/",
        "source_url": "https://sdmx.oecd.org/public/rest/v1/",
        "access_mode": "official SDMX CSV API",
        "update_frequency": "regular",
        "license_or_terms": "OECD Data Explorer terms",
        "automation_status": "implemented_official_sdmx_api",
        "release_year": 2025,
    },
    "EUROSTAT_HTEC": {
        "source_name": "Eurostat high-technology sectors, HRST and R&D datasets",
        "owner": "Eurostat",
        "url": "https://ec.europa.eu/eurostat/",
        "source_url": "https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/",
        "access_mode": "official Eurostat Statistics API",
        "update_frequency": "regular",
        "license_or_terms": "Eurostat reuse policy",
        "automation_status": "implemented_official_json_api",
        "release_year": 2025,
    },
    "NATIONAL_STATS_HTEI": {
        "source_name": "Curated national statistical office evidence register for HTEI",
        "owner": "National statistical offices",
        "url": "local:methodology/national-statistics-register",
        "source_url": "local:methodology/national-statistics-register",
        "access_mode": "checksum-verified official-file register",
        "update_frequency": "curated per release",
        "license_or_terms": "Source-specific national statistical office terms recorded per snapshot",
        "automation_status": "implemented_audit_register",
        "release_year": 2026,
    },
    "CORPORATE_REPORTS_HTEI": {
        "source_name": "Corporate reporting evidence register for high-technology workforce strategy",
        "owner": "Public companies and reporting issuers",
        "url": "local:methodology/corporate-report-register",
        "source_url": "local:methodology/corporate-report-register",
        "access_mode": "checksum-verified official report register",
        "update_frequency": "curated per release",
        "license_or_terms": "Source-specific issuer terms recorded per report",
        "automation_status": "implemented_audit_register",
        "release_year": 2026,
    },
    "SPECIALIZED_RATINGS_HTEI": {
        "source_name": "Specialized rankings evidence used by HTEI and training-system model",
        "owner": "QS, WIPO, Portulans and other official ranking publishers",
        "url": "local:methodology/specialized-ratings-register",
        "source_url": "local:methodology/specialized-ratings-register",
        "access_mode": "derived from archived official ranking snapshots",
        "update_frequency": "annual",
        "license_or_terms": "Source-specific ranking terms; numeric layers keep original source provenance",
        "automation_status": "implemented_audit_register",
        "release_year": 2026,
    },
    "HTEI_MULTI_SOURCE": {
        "source_name": "Legacy HTEI observation-staging manifest (superseded by scientific v5)",
        "owner": "MGIMO / FNISC RAS",
        "url": "local:methodology/htei-observation-staging-v3",
        "source_url": "local:methodology/htei-observation-staging-v3",
        "access_mode": "computed manifest from official archived snapshots",
        "update_frequency": "per release build",
        "license_or_terms": "Project methodology; underlying values retain original official source terms",
        "automation_status": "implemented_reproducible_transform",
        "release_year": 2026,
    },
    "QS_ET": {
        "source_name": "QS World University Rankings by Subject: Engineering & Technology",
        "owner": "QS Quacquarelli Symonds",
        "url": QS_OFFICIAL_URL,
        "source_url": QS_OFFICIAL_URL,
        "access_mode": "official public JSON endpoint archive",
        "update_frequency": "annual",
        "license_or_terms": "QS TopUniversities ranking terms; public endpoint responses archived for reproducible non-commercial analysis.",
        "automation_status": "implemented_official_web_endpoint",
        "release_year": 2026,
    },
}


INDEX_DEFINITIONS = [
    {
        "code": "HDI",
        "name": "Human Development Index",
        "short_name": "HDI",
        "name_ru": "Индекс человеческого развития",
        "short_name_ru": "ИЧР",
        "name_en": "Human Development Index",
        "short_name_en": "HDI",
        "theme": "human_development",
        "authority": "UNDP",
        "url": "https://hdr.undp.org/data-center",
        "score_scale": "0-1 official score shown on a 0-100 display scale",
        "rank_direction": "desc",
        "recompute_mode": "official score with diagnostic components normalized from official fields",
        "official_or_derived": "official_index",
        "is_tz_index": 0,
        "description_ru": "Официальный индекс ООН, объединяющий здоровье, образование и уровень жизни.",
        "description_en": "Official UN index combining health, education and living standards.",
    },
    {
        "code": "HCI",
        "name": "Human Capital Index",
        "short_name": "HCI",
        "name_ru": "Индекс человеческого капитала",
        "short_name_ru": "Индекс человеческого капитала",
        "name_en": "Human Capital Index",
        "short_name_en": "HCI",
        "theme": "human_capital",
        "authority": "World Bank",
        "url": "https://www.worldbank.org/en/publication/human-capital",
        "score_scale": "0-1 official score shown on a 0-100 display scale",
        "rank_direction": "desc",
        "recompute_mode": "official API score with available official HCI component indicators",
        "official_or_derived": "official_index",
        "is_tz_index": 0,
        "description_ru": "Официальная оценка человеческого капитала и будущей продуктивности поколения.",
        "description_en": "Official measure of human capital and expected productivity of the next generation.",
    },
    {
        "code": "GTCI",
        "name": "Global Talent Competitiveness Index",
        "short_name": "GTCI",
        "name_ru": "Индекс глобальной конкурентоспособности талантов",
        "short_name_ru": "Индекс глобальной конкурентоспособности талантов",
        "name_en": "Global Talent Competitiveness Index",
        "short_name_en": "GTCI",
        "theme": "talent",
        "authority": "INSEAD / Portulans",
        "url": "https://portulansinstitute.org/reports/",
        "score_scale": "0-100",
        "rank_direction": "desc",
        "recompute_mode": "official report rank and score extraction",
        "official_or_derived": "official_index",
        "is_tz_index": 0,
        "description_ru": "Официальный индекс конкурентоспособности стран в создании, привлечении, развитии и удержании талантов.",
        "description_en": "Official index of country competitiveness in enabling, attracting, growing and retaining talent.",
    },
    {
        "code": "GII",
        "name": "Global Innovation Index",
        "short_name": "GII",
        "name_ru": "Глобальный инновационный индекс",
        "short_name_ru": "Глобальный инновационный индекс",
        "name_en": "Global Innovation Index",
        "short_name_en": "GII",
        "theme": "innovation",
        "authority": "WIPO",
        "url": "https://www.wipo.int/global_innovation_index/en/",
        "score_scale": "0-100",
        "rank_direction": "desc",
        "recompute_mode": "official score and pillar import from WIPO database",
        "official_or_derived": "official_index",
        "is_tz_index": 0,
        "description_ru": "Официальная оценка инновационной экосистемы, входных ресурсов и результатов.",
        "description_en": "Official assessment of innovation ecosystem inputs and outputs.",
    },
    {
        "code": "IDI",
        "name": "ICT Development Index",
        "short_name": "IDI",
        "name_ru": "Индекс развития ИКТ",
        "short_name_ru": "Индекс развития ИКТ",
        "name_en": "ICT Development Index",
        "short_name_en": "IDI",
        "theme": "digital_connectivity",
        "authority": "ITU",
        "url": "https://www.itu.int/itu-d/reports/statistics/idi2025/",
        "score_scale": "0-100",
        "rank_direction": "desc",
        "recompute_mode": "official aggregate and pillar scores from ITU dataset",
        "official_or_derived": "official_index",
        "is_tz_index": 0,
        "description_ru": "Официальный индекс цифрового развития и качества значимой связности.",
        "description_en": "Official ICT development and meaningful connectivity index.",
    },
    {
        "code": "QS_ET",
        "name": "QS Engineering & Technology",
        "short_name": "QS E&T",
        "name_ru": "QS: инженерия и технологии",
        "short_name_ru": "QS: инженерия и технологии",
        "name_en": "QS Engineering & Technology",
        "short_name_en": "QS E&T",
        "theme": "engineering_education",
        "authority": "QS",
        "url": QS_OFFICIAL_URL,
        "score_scale": "0-100 country aggregation from archived official endpoint rows",
        "rank_direction": "desc",
        "recompute_mode": "derived_country_aggregation from archived official QS endpoint rows",
        "official_or_derived": "derived_country_aggregation",
        "is_tz_index": 0,
        "description_ru": "Страновая агрегация рейтинга QS по инженерии и технологиям из архивированных строк официального endpoint.",
        "description_en": "Country aggregation of QS Engineering & Technology ranking from archived official endpoint rows.",
    },
    {
        "code": "HTEI",
        "name": "High-Tech Employment Index",
        "short_name": "HTEI",
        "name_ru": "Индекс занятости в высокотехнологичных отраслях",
        "short_name_ru": "Индекс занятости в высокотехнологичных отраслях",
        "name_en": "High-Tech Employment Index",
        "short_name_en": "HTEI",
        "theme": "hightech_workforce",
        "authority": "MGIMO / FNISC RAS",
        "url": "local:methodology/htei",
        "score_scale": "0-100",
        "rank_direction": "desc",
        "recompute_mode": "computed from official workforce, R&D, STEM and technology output indicators",
        "official_or_derived": "project_index_required_by_TZ",
        "is_tz_index": 1,
        "description_ru": "Индекс занятости в высокотехнологичных отраслях (High-tech employment index) в системе индикаторов для мониторинга развития технологических кадров, рассчитанный на официальных международных рядах.",
        "description_en": "High-Tech Employment Index for monitoring technological workforce development, computed from official international data series.",
    },
]


COMPONENTS: dict[str, list[tuple[str, str, str, float, str]]] = {
    "HDI": [
        ("LE", "Ожидаемая продолжительность жизни", "Life expectancy at birth", 1 / 3, "UNDP HDR component"),
        ("EDU", "Образование", "Education index components", 1 / 3, "UNDP expected and mean years of schooling"),
        ("GNI", "ВНД на душу населения", "GNI per capita", 1 / 3, "UNDP GNI per capita component"),
    ],
    "HCI": [
        ("SURV", "Выживаемость", "Survival", 0.25, "World Bank HCI survival indicators"),
        ("SCHOOL", "Ожидаемые годы обучения", "Expected years of school", 0.25, "World Bank HCI schooling"),
        ("LEARN", "Результаты обучения", "Learning outcomes", 0.25, "World Bank HCI harmonized test scores"),
        ("HEALTH", "Здоровье детей", "Child health", 0.25, "World Bank HCI not-stunted indicator"),
    ],
    "GTCI": [
        ("ENABLE", "Среда для талантов", "Enable", 1 / 6, "GTCI model pillar"),
        ("ATTRACT", "Привлечение талантов", "Attract", 1 / 6, "GTCI model pillar"),
        ("GROW", "Развитие талантов", "Grow", 1 / 6, "GTCI model pillar"),
        ("RETAIN", "Удержание талантов", "Retain", 1 / 6, "GTCI model pillar"),
        ("VT", "Профессионально-технические навыки", "Vocational and technical skills", 1 / 6, "GTCI model pillar"),
        ("GA", "Адаптивные навыки высокого уровня", "Generalist adaptive skills", 1 / 6, "GTCI model pillar"),
    ],
    "GII": [
        ("INST", "Институты", "Institutions", 1 / 7, "WIPO pillar score"),
        ("HCR", "Человеческий капитал и исследования", "Human capital and research", 1 / 7, "WIPO pillar score"),
        ("INFRA", "Инфраструктура", "Infrastructure", 1 / 7, "WIPO pillar score"),
        ("MARKET", "Развитость рынков", "Market sophistication", 1 / 7, "WIPO pillar score"),
        ("BUS", "Развитость бизнеса", "Business sophistication", 1 / 7, "WIPO pillar score"),
        ("KTO", "Знания и технологии: результаты", "Knowledge and technology outputs", 1 / 7, "WIPO pillar score"),
        ("CRE", "Креативные результаты", "Creative outputs", 1 / 7, "WIPO pillar score"),
    ],
    "IDI": [
        ("ACCESS", "Универсальная связность", "Universal connectivity", 0.40, "ITU pillar score"),
        ("USE", "Использование ИКТ", "ICT use", 0.20, "ITU normalized use indicators"),
        ("MEANING", "Качество значимой связности", "Meaningful connectivity", 0.40, "ITU pillar score"),
    ],
    "QS_ET": [
        ("TOP_COUNT", "Число вузов в рейтинге", "Ranked institutions count", 0.25, "Official QS export field"),
        ("BEST_RANK", "Лучшая позиция вуза", "Best institutional rank", 0.25, "Official QS export field"),
        ("MEDIAN_RANK", "Медианная позиция вузов", "Median institutional rank", 0.25, "Official QS export field"),
        ("QS_SCORE", "Суммарная оценка вузов", "Aggregated institutional score", 0.25, "Official QS export field"),
    ],
    "HTEI": [
        ("HT_EMPLOYMENT_SHARE", "Доля занятости в высокотехнологичных секторах", "Employment share in high-technology sectors", 0.22, "Eurostat high-tech sectors, ILOSTAT ISIC activity and national statistical office rows"),
        ("HIGH_TECH_OCCUPATIONS", "Профессиональные и технические STEM-занятия", "STEM and high-technology occupations", 0.18, "ILOSTAT ISCO occupations and Eurostat HRST scientists and engineers"),
        ("RND_PERSONNEL", "Персонал исследований и разработок", "R&D personnel", 0.18, "OECD MSTI/R&D personnel, Eurostat R&D personnel and World Bank/UIS R&D indicators"),
        ("STEM_PIPELINE", "Выпускники STEM-направлений", "STEM tertiary graduates pipeline", 0.16, "UNESCO UIS STEM tertiary graduates indicator"),
        ("TECH_OUTPUTS", "Технологические результаты", "Technology outputs", 0.14, "World Bank high-technology exports, ICT services exports and resident patents"),
        ("CORPORATE_STRATEGY_AND_DEMAND", "Корпоративные стратегии и технологический спрос", "Corporate strategies and technology demand", 0.12, "OECD BERD/business R&D and checksum-verified corporate reporting evidence"),
    ],
}

HCI_INDICATORS = [
    "HD.HCI.OVRL",
    "HD.HCI.MORT",
    "HD.HCI.AMRT",
    "HD.HCI.EYRS",
    "HD.HCI.HLOS",
    "HD.HCI.STNT",
]

HTEI_INDICATORS = [
    "JI.EMP.PROF.ZS",
    "JI.EMP.TECH.ZS",
    "SP.POP.SCIE.RD.P6",
    "SP.POP.TECH.RD.P6",
    "UIS.FOSGP.5T8.F500600700",
    "TX.VAL.TECH.MF.ZS",
    "BX.GSR.CCIS.ZS",
    "IP.PAT.RESD",
    "SP.POP.TOTL",
]

RU_COUNTRY_NAMES = {
    "AUS": "Австралия", "AUT": "Австрия", "BEL": "Бельгия", "BRA": "Бразилия", "CAN": "Канада",
    "CHE": "Швейцария", "CHL": "Чили", "CHN": "Китай", "CZE": "Чехия", "DEU": "Германия",
    "DNK": "Дания", "EGY": "Египет", "ESP": "Испания", "EST": "Эстония", "FIN": "Финляндия",
    "FRA": "Франция", "GBR": "Великобритания", "GRC": "Греция", "HKG": "Гонконг", "IND": "Индия",
    "IRL": "Ирландия", "ISR": "Израиль", "ITA": "Италия", "JPN": "Япония", "KAZ": "Казахстан",
    "KOR": "Республика Корея", "MEX": "Мексика", "MYS": "Малайзия", "NLD": "Нидерланды",
    "NOR": "Норвегия", "PHL": "Филиппины", "POL": "Польша", "PRT": "Португалия",
    "RUS": "Россия", "SAU": "Саудовская Аравия", "SGP": "Сингапур", "SWE": "Швеция",
    "THA": "Таиланд", "TUR": "Турция", "USA": "США", "VNM": "Вьетнам", "ZAF": "ЮАР",
}
RU_LOCALE = Locale.parse("ru")
SPECIAL_ISO2 = {"WBG": "PS"}
SPECIAL_COUNTRY_NAMES_EN = {"WBG": "West Bank and Gaza"}
SPECIAL_COUNTRY_NAMES_RU = {"WBG": "Западный берег и сектор Газа"}

ISO_ALIASES = {
    "Bolivia, Plurinational St.": "BOL",
    "Brunei Darussalam": "BRN",
    "Cabo Verde": "CPV",
    "Congo, Dem. Rep.": "COD",
    "Côte d’Ivoire": "CIV",
    "Côte d'Ivoire": "CIV",
    "Czech Republic": "CZE",
    "Gambia": "GMB",
    "Hong Kong, China": "HKG",
    "Iran, Islamic Rep.": "IRN",
    "Korea, Rep.": "KOR",
    "Lao PDR": "LAO",
    "Moldova, Rep.": "MDA",
    "Russian Federation": "RUS",
    "Tanzania, United Rep.": "TZA",
    "Türkiye": "TUR",
    "United States of America": "USA",
    "Venezuela, Bolivarian Republic of": "VEN",
    "Viet Nam": "VNM",
    "Yemen": "YEM",
}

ISO_ALIASES.update({
    "China (Mainland)": "CHN",
    "Czechia": "CZE",
    "Hong Kong SAR": "HKG",
    "Hong Kong SAR, China": "HKG",
    "Iran": "IRN",
    "Macau SAR": "MAC",
    "Macau SAR, China": "MAC",
    "Russia": "RUS",
    "South Korea": "KOR",
    "Taiwan": "TWN",
    "United Kingdom": "GBR",
    "United States": "USA",
    "Vietnam": "VNM",
})


def now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def rel_path(path: Path) -> str:
    return path.resolve().relative_to(DATA_DIR.parent.resolve()).as_posix()


def flag_emoji(iso2: str) -> str:
    return "".join(chr(127397 + ord(c)) for c in iso2.upper()) if len(iso2) == 2 else ""


def iso2_for(iso3: str) -> str:
    if iso3 in SPECIAL_ISO2:
        return SPECIAL_ISO2[iso3]
    try:
        country = pycountry.countries.get(alpha_3=iso3)
        return country.alpha_2 if country else ""
    except Exception:
        return ""


def country_name_en(iso3: str, fallback: str | None = None) -> str:
    if iso3 in SPECIAL_COUNTRY_NAMES_EN:
        return SPECIAL_COUNTRY_NAMES_EN[iso3]
    try:
        country = pycountry.countries.get(alpha_3=iso3)
        if country:
            return getattr(country, "common_name", None) or country.name
    except Exception:
        pass
    return fallback or iso3


def iso3_from_name(name: str) -> str | None:
    cleaned = " ".join(name.replace("’", "'").split())
    if cleaned in ISO_ALIASES:
        return ISO_ALIASES[cleaned]
    try:
        result = pycountry.countries.lookup(cleaned)
        return result.alpha_3
    except LookupError:
        return None


def safe_float(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if math.isnan(float(value)):
            return None
        return float(value)
    text = str(value).strip()
    if not text or text.lower() in {"n.a.", "n.a", "na", "n.p.", "n.p", "..", "-"}:
        return None
    text = re.sub(r"[†‡*]", "", text).replace(",", "")
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    return float(match.group(0)) if match else None


def normalize_values(raw_values: dict[str, float], reverse: bool = False) -> dict[str, float]:
    vals = [v for v in raw_values.values() if v is not None]
    if not vals:
        return {}
    lo, hi = min(vals), max(vals)
    if hi == lo:
        return {k: 100.0 for k in raw_values}
    out: dict[str, float] = {}
    for key, value in raw_values.items():
        score = (value - lo) / (hi - lo) * 100
        out[key] = 100 - score if reverse else score
    return out


def bounded_score(value: float | None) -> float | None:
    if value is None:
        return None
    return max(0.0, min(100.0, float(value)))


def group_scores_by_year(scores: list[dict[str, Any]]) -> dict[int, list[dict[str, Any]]]:
    grouped: dict[int, list[dict[str, Any]]] = {}
    for score in scores:
        grouped.setdefault(int(score.get("year") or score.get("source_data_year") or SNAPSHOT_RELEASE_YEAR), []).append(score)
    return grouped


def ensure_download(source_id: str, release_year: int, url: str, filename: str, refresh: bool = False) -> Snapshot:
    source = SOURCE_DEFS[source_id]
    source_dir = RAW_DIR / source_id / str(release_year)
    existing = sorted(source_dir.glob(f"*/{filename}"))
    existing_snapshot = None
    if existing:
        path = existing[-1]
        meta_path = path.with_suffix(path.suffix + ".metadata.json")
        if meta_path.exists():
            existing_snapshot = Snapshot(**json.loads(meta_path.read_text(encoding="utf-8")))
            if not refresh:
                return existing_snapshot

    retrieved_at = now_iso()
    response = requests.get(url, timeout=SOURCE_TIMEOUT, headers=HEADERS)
    response.raise_for_status()
    content = response.content
    if existing_snapshot and sha256_bytes(content) == existing_snapshot.raw_snapshot_sha256:
        return existing_snapshot
    stamp = retrieved_at.replace(":", "").replace("+", "Z")
    target_dir = source_dir / stamp
    target_dir.mkdir(parents=True, exist_ok=True)
    path = target_dir / filename
    path.write_bytes(content)
    snapshot = Snapshot(
        snapshot_id=f"{source_id}:{release_year}:{sha256_bytes(content)[:16]}",
        source_id=source_id,
        release_year=release_year,
        retrieved_at=retrieved_at,
        source_url=url,
        raw_snapshot_path=rel_path(path),
        raw_snapshot_sha256=sha256_bytes(content),
        content_type=response.headers.get("content-type", ""),
        bytes_count=len(content),
        license_or_terms=source["license_or_terms"],
    )
    path.with_suffix(path.suffix + ".metadata.json").write_text(json.dumps(snapshot.__dict__, ensure_ascii=False, indent=2), encoding="utf-8")
    return snapshot


def ensure_world_bank_snapshot(source_id: str, release_year: int, indicators: list[str], filename: str, refresh: bool = False) -> Snapshot:
    source = SOURCE_DEFS[source_id]
    source_dir = RAW_DIR / source_id / str(release_year)
    existing = sorted(source_dir.glob(f"*/{filename}"))
    existing_snapshot = None
    existing_payload = None
    if existing:
        path = existing[-1]
        meta_path = path.with_suffix(path.suffix + ".metadata.json")
        if meta_path.exists():
            existing_snapshot = Snapshot(**json.loads(meta_path.read_text(encoding="utf-8")))
            if not refresh:
                return existing_snapshot
            existing_payload = json.loads(path.read_text(encoding="utf-8"))

    retrieved_at = now_iso()
    payload: dict[str, Any] = {"retrieved_at": retrieved_at, "source": source, "indicators": {}}
    for indicator in indicators:
        rows: list[dict[str, Any]] = []
        page = 1
        while True:
            url = f"{WORLD_BANK_API_BASE}/country/all/indicator/{indicator}?format=json&per_page=20000&page={page}"
            response = requests.get(url, timeout=SOURCE_TIMEOUT, headers=HEADERS)
            response.raise_for_status()
            data = response.json()
            if not isinstance(data, list) or len(data) < 2:
                break
            meta, records = data[0], data[1] or []
            rows.extend(records)
            if page >= int(meta.get("pages") or 1):
                break
            page += 1
        payload["indicators"][indicator] = rows
    if existing_snapshot and existing_payload and payload["indicators"] == existing_payload.get("indicators"):
        return existing_snapshot
    content = json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")
    stamp = retrieved_at.replace(":", "").replace("+", "Z")
    target_dir = source_dir / stamp
    target_dir.mkdir(parents=True, exist_ok=True)
    path = target_dir / filename
    path.write_bytes(content)
    snapshot = Snapshot(
        snapshot_id=f"{source_id}:{release_year}:{sha256_bytes(content)[:16]}",
        source_id=source_id,
        release_year=release_year,
        retrieved_at=retrieved_at,
        source_url=source["source_url"],
        raw_snapshot_path=rel_path(path),
        raw_snapshot_sha256=sha256_bytes(content),
        content_type="application/json",
        bytes_count=len(content),
        license_or_terms=source["license_or_terms"],
    )
    path.with_suffix(path.suffix + ".metadata.json").write_text(json.dumps(snapshot.__dict__, ensure_ascii=False, indent=2), encoding="utf-8")
    return snapshot


def ensure_world_bank_countries_snapshot(refresh: bool = False) -> Snapshot:
    return ensure_download("WORLD_BANK_COUNTRIES", 2026, WORLD_BANK_COUNTRIES_URL, "world_bank_countries.json", refresh=refresh)


def ensure_uis_snapshot(refresh: bool = False) -> Snapshot:
    source_id = "UIS_HTEI"
    release_year = 2026
    filename = "uis_stem_pipeline_api.json"
    source = SOURCE_DEFS[source_id]
    source_dir = RAW_DIR / source_id / str(release_year)
    existing = sorted(source_dir.glob(f"*/{filename}"))
    existing_snapshot = None
    existing_payload = None
    if existing:
        path = existing[-1]
        meta_path = path.with_suffix(path.suffix + ".metadata.json")
        if meta_path.exists():
            existing_snapshot = Snapshot(**json.loads(meta_path.read_text(encoding="utf-8")))
            if not refresh:
                return existing_snapshot
            existing_payload = json.loads(path.read_text(encoding="utf-8"))

    version_response = requests.get(UIS_DEFAULT_VERSION_URL, timeout=SOURCE_TIMEOUT, headers=HEADERS)
    version_response.raise_for_status()
    version_metadata = version_response.json()
    version = str(version_metadata["version"])
    params = {
        "indicator": UIS_STEM_INDICATOR,
        "geoUnitType": "NATIONAL",
        "start": HISTORICAL_START_YEAR,
        "end": LATEST_RELEASE_YEAR,
        "indicatorMetadata": "true",
        "version": version,
    }
    response = requests.get(UIS_STEM_DATA_URL, params=params, timeout=SOURCE_TIMEOUT, headers=HEADERS)
    response.raise_for_status()
    response_payload = response.json()
    records = response_payload.get("records") or []
    if not records:
        raise RuntimeError("UNESCO UIS Data API returned no STEM pipeline records")
    payload = {
        "api_version": version,
        "version_metadata": version_metadata,
        "query_url": response.url,
        "retrieved_at": now_iso(),
        "indicator": UIS_STEM_INDICATOR,
        "response": response_payload,
    }
    if existing_snapshot and existing_payload:
        if (
            existing_payload.get("api_version") == payload["api_version"]
            and existing_payload.get("response") == payload["response"]
        ):
            return existing_snapshot

    content = json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")
    stamp = payload["retrieved_at"].replace(":", "").replace("+", "Z")
    target_dir = source_dir / stamp
    target_dir.mkdir(parents=True, exist_ok=True)
    path = target_dir / filename
    path.write_bytes(content)
    snapshot = Snapshot(
        snapshot_id=f"{source_id}:{release_year}:{sha256_bytes(content)[:16]}",
        source_id=source_id,
        release_year=release_year,
        retrieved_at=payload["retrieved_at"],
        source_url=response.url,
        raw_snapshot_path=rel_path(path),
        raw_snapshot_sha256=sha256_bytes(content),
        content_type="application/json",
        bytes_count=len(content),
        license_or_terms=source["license_or_terms"],
    )
    path.with_suffix(path.suffix + ".metadata.json").write_text(
        json.dumps(snapshot.__dict__, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return snapshot


def world_bank_region_for_country(iso3: str, region: str | None) -> str:
    region = (region or "").strip()
    if not region or region == "Aggregates":
        return COUNTRY_METADATA_FALLBACK_REGION
    if region == "Europe & Central Asia":
        return "Central Asia" if iso3 in ECA_TO_EUROPE_EXCEPTIONS else "Europe"
    return region


def parse_world_bank_country_metadata(snapshot: Snapshot) -> tuple[dict[str, str], dict[str, str], dict[str, Any]]:
    path = snapshot_path(snapshot)
    payload = json.loads(path.read_text(encoding="utf-8"))
    records = payload[1] if isinstance(payload, list) and len(payload) > 1 else []
    iso_to_region: dict[str, str] = {}
    iso_to_income: dict[str, str] = {}
    skipped: list[str] = []
    for record in records or []:
        iso3 = (record.get("id") or "").upper()
        if not iso2_for(iso3):
            skipped.append(iso3)
            continue
        wb_region = ((record.get("region") or {}).get("value") or "").strip()
        income = ((record.get("incomeLevel") or {}).get("value") or "").strip()
        if not wb_region or wb_region == "Aggregates":
            skipped.append(iso3)
            continue
        iso_to_region[iso3] = world_bank_region_for_country(iso3, wb_region)
        iso_to_income[iso3] = income if income and income != "Aggregates" else COUNTRY_METADATA_FALLBACK_INCOME
    audit = {
        "source_id": snapshot.source_id,
        "snapshot_id": snapshot.snapshot_id,
        "retrieved_at": snapshot.retrieved_at,
        "source_url": snapshot.source_url,
        "raw_snapshot_path": snapshot.raw_snapshot_path,
        "raw_snapshot_sha256": snapshot.raw_snapshot_sha256,
        "transform_id": "world_bank_country_metadata_to_country_dimension",
        "rows_loaded": len(iso_to_region),
        "skipped_codes": sorted({code for code in skipped if code}),
        "region_split_rule": "World Bank Europe & Central Asia is split into Europe vs Central Asia for platform region filters; Central Asia ISO3 set: KAZ, KGZ, TJK, TKM, UZB.",
    }
    return iso_to_region, iso_to_income, audit


def ensure_multifile_snapshot(
    source_id: str,
    release_year: int,
    files: dict[str, str],
    filename: str,
    accept: str | None = None,
    refresh: bool = False,
) -> Snapshot:
    source = SOURCE_DEFS[source_id]
    source_dir = RAW_DIR / source_id / str(release_year)
    existing = sorted(source_dir.glob(f"*/{filename}"))
    existing_snapshot = None
    existing_manifest = None
    if existing:
        path = existing[-1]
        meta_path = path.with_suffix(path.suffix + ".metadata.json")
        if meta_path.exists():
            existing_snapshot = Snapshot(**json.loads(meta_path.read_text(encoding="utf-8")))
            existing_manifest = json.loads(path.read_text(encoding="utf-8"))
            if not refresh and set(files).issubset(existing_manifest.get("files", {})):
                return existing_snapshot

    retrieved_at = now_iso()
    file_meta: dict[str, dict[str, Any]] = {}
    downloaded: dict[str, tuple[bytes, str]] = {}
    request_headers = dict(HEADERS)
    if accept:
        request_headers["Accept"] = accept
    for key, url in files.items():
        response = requests.get(url, timeout=SOURCE_TIMEOUT, headers=request_headers)
        response.raise_for_status()
        content = response.content
        suffix = ".json" if "json" in (response.headers.get("content-type") or "").lower() else ".csv"
        file_meta[key] = {
            "url": url,
            "path": "",
            "sha256": sha256_bytes(content),
            "content_type": response.headers.get("content-type", ""),
            "bytes_count": len(content),
        }
        downloaded[key] = (content, suffix)
    if existing_snapshot and existing_manifest:
        old_hashes = {key: item.get("sha256") for key, item in existing_manifest.get("files", {}).items()}
        new_hashes = {key: item["sha256"] for key, item in file_meta.items()}
        if old_hashes == new_hashes:
            return existing_snapshot
    stamp = retrieved_at.replace(":", "").replace("+", "Z")
    target_dir = source_dir / stamp
    target_dir.mkdir(parents=True, exist_ok=True)
    for key, (content, suffix) in downloaded.items():
        path = target_dir / f"{key}{suffix}"
        path.write_bytes(content)
        file_meta[key]["path"] = rel_path(path)
    manifest = {
        "source_id": source_id,
        "source_name": source["source_name"],
        "retrieved_at": retrieved_at,
        "release_year": release_year,
        "files": file_meta,
    }
    content = json.dumps(manifest, ensure_ascii=False, indent=2).encode("utf-8")
    manifest_path = target_dir / filename
    manifest_path.write_bytes(content)
    snapshot = Snapshot(
        snapshot_id=f"{source_id}:{release_year}:{sha256_bytes(content)[:16]}",
        source_id=source_id,
        release_year=release_year,
        retrieved_at=retrieved_at,
        source_url=source["source_url"],
        raw_snapshot_path=rel_path(manifest_path),
        raw_snapshot_sha256=sha256_bytes(content),
        content_type="application/json",
        bytes_count=len(content),
        license_or_terms=source["license_or_terms"],
    )
    manifest_path.with_suffix(manifest_path.suffix + ".metadata.json").write_text(json.dumps(snapshot.__dict__, ensure_ascii=False, indent=2), encoding="utf-8")
    return snapshot


def ensure_json_snapshot(source_id: str, release_year: int, payload: dict[str, Any], filename: str) -> Snapshot:
    source = SOURCE_DEFS[source_id]
    source_dir = RAW_DIR / source_id / str(release_year)
    existing = sorted(source_dir.glob(f"*/{filename}"))
    existing_snapshot = None
    existing_manifest = None
    if existing:
        path = existing[-1]
        meta_path = path.with_suffix(path.suffix + ".metadata.json")
        if meta_path.exists():
            existing_snapshot = Snapshot(**json.loads(meta_path.read_text(encoding="utf-8")))
            existing_manifest = json.loads(path.read_text(encoding="utf-8"))

    expected_payload = {**payload, "source_id": source_id, "release_year": release_year}
    if existing_snapshot and existing_manifest:
        comparable_existing = {key: value for key, value in existing_manifest.items() if key != "retrieved_at"}
        if comparable_existing == expected_payload:
            return existing_snapshot

    retrieved_at = now_iso()
    stamp = retrieved_at.replace(":", "").replace("+", "Z")
    target_dir = source_dir / stamp
    target_dir.mkdir(parents=True, exist_ok=True)
    manifest = {**expected_payload, "retrieved_at": retrieved_at}
    content = json.dumps(manifest, ensure_ascii=False, indent=2).encode("utf-8")
    path = target_dir / filename
    path.write_bytes(content)
    snapshot = Snapshot(
        snapshot_id=f"{source_id}:{release_year}:{sha256_bytes(content)[:16]}",
        source_id=source_id,
        release_year=release_year,
        retrieved_at=retrieved_at,
        source_url=source["source_url"],
        raw_snapshot_path=rel_path(path),
        raw_snapshot_sha256=sha256_bytes(content),
        content_type="application/json",
        bytes_count=len(content),
        license_or_terms=source["license_or_terms"],
    )
    path.with_suffix(path.suffix + ".metadata.json").write_text(json.dumps(snapshot.__dict__, ensure_ascii=False, indent=2), encoding="utf-8")
    return snapshot


def ensure_optional_download(source_id: str, release_year: int, url: str, filename: str) -> Snapshot | None:
    try:
        return ensure_download(source_id, release_year, url, filename)
    except Exception:
        return None


def archived_snapshot_metadata() -> list[Snapshot]:
    snapshots: list[Snapshot] = []
    required = set(Snapshot.__dataclass_fields__)
    for meta_path in RAW_DIR.rglob("*.metadata.json"):
        try:
            payload = json.loads(meta_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        if not required.issubset(payload):
            continue
        try:
            snapshots.append(Snapshot(**{key: payload[key] for key in required}))
        except (TypeError, ValueError):
            continue
    return snapshots


def qs_manifest_path(year: int) -> Path | None:
    manifests = sorted((RAW_DIR / "QS_ET" / str(year)).glob(f"*/qs_engineering_technology_{year}.manifest.json"))
    return manifests[-1] if manifests else None


def qs_manifest_snapshot(manifest_path: Path) -> Snapshot:
    meta_path = manifest_path.with_suffix(manifest_path.suffix + ".metadata.json")
    if meta_path.exists():
        return Snapshot(**json.loads(meta_path.read_text(encoding="utf-8")))
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    content = manifest_path.read_bytes()
    snapshot = Snapshot(
        snapshot_id=f"QS_ET:{manifest['year']}:{sha256_bytes(content)[:16]}",
        source_id="QS_ET",
        release_year=int(manifest["year"]),
        retrieved_at=manifest["retrieved_at"],
        source_url=manifest["page_url"],
        raw_snapshot_path=rel_path(manifest_path),
        raw_snapshot_sha256=sha256_bytes(content),
        content_type="application/json",
        bytes_count=len(content),
        license_or_terms=SOURCE_DEFS["QS_ET"]["license_or_terms"],
    )
    meta_path.write_text(json.dumps(snapshot.__dict__, ensure_ascii=False, indent=2), encoding="utf-8")
    return snapshot


def ensure_qs_snapshots(years: list[int] | None = None) -> dict[int, Snapshot]:
    years = years or QS_YEARS
    snapshots: dict[int, Snapshot] = {}
    missing: list[int] = []
    for year in years:
        manifest = qs_manifest_path(year)
        if manifest:
            snapshots[year] = qs_manifest_snapshot(manifest)
        else:
            missing.append(year)
    if missing:
        script = DATA_DIR.parent / "scripts" / "fetch_qs_rankings.cjs"
        result = subprocess.run(
            ["node", str(script), "--root", str(DATA_DIR.parent), "--years", ",".join(str(y) for y in missing)],
            cwd=DATA_DIR.parent,
            capture_output=True,
            text=True,
            check=False,
        )
        if result.returncode != 0:
            raise RuntimeError(f"QS TopUniversities endpoint archival failed: {result.stderr or result.stdout}")
        for year in missing:
            manifest = qs_manifest_path(year)
            if not manifest:
                raise RuntimeError(f"QS manifest was not created for {year}")
            snapshots[year] = qs_manifest_snapshot(manifest)
    return snapshots


def snapshot_path(snapshot: Snapshot) -> Path:
    return DATA_DIR.parent / snapshot.raw_snapshot_path


def transform_for(snapshot: Snapshot, transform_id: str, rows_loaded: int, notes: str) -> Transform:
    ts = now_iso()
    return Transform(
        transformation_run_id=f"{transform_id}:{snapshot.raw_snapshot_sha256[:12]}",
        transform_id=transform_id,
        source_id=snapshot.source_id,
        snapshot_id=snapshot.snapshot_id,
        started_at=ts,
        completed_at=ts,
        code_version="giip-production-data-v1",
        rows_loaded=rows_loaded,
        notes=notes,
    )


def provenance(snapshot: Snapshot, transform: Transform, formula_version: str, quality_flag: str, is_official: bool, is_recomputed: bool) -> dict[str, Any]:
    source = SOURCE_DEFS[snapshot.source_id]
    return {
        "source_id": snapshot.source_id,
        "source_name": source["source_name"],
        "source_url": snapshot.source_url,
        "license_or_terms": snapshot.license_or_terms,
        "retrieved_at": snapshot.retrieved_at,
        "release_year": snapshot.release_year,
        "raw_snapshot_path": snapshot.raw_snapshot_path,
        "raw_snapshot_sha256": snapshot.raw_snapshot_sha256,
        "transform_id": transform.transform_id,
        "transformation_run_id": transform.transformation_run_id,
        "formula_version": formula_version,
        "quality_flag": quality_flag,
        "is_official": is_official,
        "is_recomputed": is_recomputed,
    }


def parse_undp_hdi(snapshot: Snapshot) -> tuple[list[dict[str, Any]], list[dict[str, Any]], Transform]:
    content = snapshot_path(snapshot).read_bytes()
    encoding = from_bytes(content).best().encoding or "utf-8"
    text = content.decode(encoding, errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    scores: list[dict[str, Any]] = []
    components: list[dict[str, Any]] = []
    fieldnames = reader.fieldnames or []
    years = sorted(
        int(match.group(1))
        for name in fieldnames
        if (match := re.match(r"^hdi_(\d{4})$", name))
        and HISTORICAL_START_YEAR <= int(match.group(1)) <= 2023
    )
    raw_components: dict[int, dict[str, dict[str, float]]] = {
        year: {"LE": {}, "EDU": {}, "GNI": {}} for year in years
    }
    rows_cache: dict[str, dict[str, Any]] = {}
    for row in reader:
        iso3 = row.get("iso3", "").strip().upper()
        if len(iso3) != 3:
            continue
        rows_cache[iso3] = row
        for year in years:
            hdi = safe_float(row.get(f"hdi_{year}"))
            if hdi is None:
                continue
            scores.append({"iso3": iso3, "year": year, "score": hdi * 100, "rank": 0, "source_data_year": year})
            le = safe_float(row.get(f"le_{year}"))
            if le is not None:
                raw_components[year]["LE"][iso3] = le
            eys = safe_float(row.get(f"eys_{year}"))
            mys = safe_float(row.get(f"mys_{year}"))
            if eys is not None and mys is not None:
                raw_components[year]["EDU"][iso3] = (eys + mys) / 2
            gni = safe_float(row.get(f"gnipc_{year}"))
            if gni is not None:
                raw_components[year]["GNI"][iso3] = math.log(max(gni, 1.0))
    for year_scores in group_scores_by_year(scores).values():
        assign_ranks(year_scores)
    normalized = {
        year: {code: normalize_values(vals) for code, vals in values.items()}
        for year, values in raw_components.items()
    }
    transform = transform_for(snapshot, "undp_hdr_2025_csv_to_hdi_timeseries", len(scores), "UNDP HDR composite time-series CSV parsed for 1990-2023 HDI scores and components.")
    for iso3, row in rows_cache.items():
        for year in years:
            if not any(s["iso3"] == iso3 and s["year"] == year for s in scores):
                continue
            for code, raw_map in raw_components[year].items():
                raw = raw_map.get(iso3)
                norm = normalized[year][code].get(iso3)
                if raw is None or norm is None:
                    continue
                raw_value = raw
                unit = "years_or_log_ppp_usd" if code in {"LE", "GNI"} else "years"
                if code == "EDU":
                    raw_value = ((safe_float(row.get(f"eys_{year}")) or 0.0) + (safe_float(row.get(f"mys_{year}")) or 0.0)) / 2
                elif code == "GNI":
                    raw_value = safe_float(row.get(f"gnipc_{year}")) or 0.0
                    unit = "2021 PPP USD"
                elif code == "LE":
                    unit = "years"
                components.append({"iso3": iso3, "year": year, "component_code": code, "raw_value": raw_value, "normalized_score": norm, "unit": unit, "source_data_year": year})
    return scores, components, transform


def load_world_bank_records(snapshot: Snapshot) -> dict[str, dict[str, dict[str, Any]]]:
    payload = json.loads(snapshot_path(snapshot).read_text(encoding="utf-8"))
    out: dict[str, dict[str, dict[str, Any]]] = {}
    for indicator, records in payload["indicators"].items():
        latest: dict[str, dict[str, Any]] = {}
        for record in records:
            iso3 = (record.get("countryiso3code") or "").upper()
            value = safe_float(record.get("value"))
            date = record.get("date")
            if len(iso3) != 3 or value is None or not date:
                continue
            old = latest.get(iso3)
            if old is None or int(date) > int(old["date"]):
                latest[iso3] = {"value": value, "date": int(date), "indicator": indicator}
        out[indicator] = latest
    return out


def load_world_bank_series(snapshot: Snapshot) -> dict[str, dict[str, dict[int, dict[str, Any]]]]:
    payload = json.loads(snapshot_path(snapshot).read_text(encoding="utf-8"))
    out: dict[str, dict[str, dict[int, dict[str, Any]]]] = {}
    for indicator, records in payload["indicators"].items():
        by_iso: dict[str, dict[int, dict[str, Any]]] = {}
        for record in records:
            iso3 = (record.get("countryiso3code") or "").upper()
            value = safe_float(record.get("value"))
            date = record.get("date")
            if len(iso3) != 3 or value is None or not date:
                continue
            year = int(date)
            if year < HISTORICAL_START_YEAR:
                continue
            by_iso.setdefault(iso3, {})[year] = {"value": value, "date": year, "indicator": indicator}
        out[indicator] = by_iso
    return out


def manifest_file_path(snapshot: Snapshot, key: str) -> Path:
    manifest = json.loads(snapshot_path(snapshot).read_text(encoding="utf-8"))
    file_meta = manifest["files"][key]
    return DATA_DIR.parent / file_meta["path"]


def iso3_from_iso2(code: str) -> str | None:
    clean = (code or "").upper()
    if clean == "EL":
        return "GRC"
    if clean == "UK":
        return "GBR"
    if len(clean) != 2:
        return None
    try:
        country = pycountry.countries.get(alpha_2=clean)
        return country.alpha_3 if country else None
    except Exception:
        return None


def jsonstat_records(payload: dict[str, Any]) -> list[tuple[dict[str, str], float]]:
    ids = payload.get("id") or []
    sizes = payload.get("size") or []
    if not ids or not sizes:
        return []
    inverses: list[dict[int, str]] = []
    for dim_id in ids:
        category = payload["dimension"][dim_id]["category"]
        index = category.get("index") or {}
        if isinstance(index, dict):
            inverses.append({int(pos): code for code, pos in index.items()})
        else:
            labels = list((category.get("label") or {}).keys())
            inverses.append({int(pos): labels[int(pos)] for pos in range(len(labels)) if int(pos) < len(labels)})
    values = payload.get("value") or {}
    items = enumerate(values) if isinstance(values, list) else ((int(k), v) for k, v in values.items())
    out: list[tuple[dict[str, str], float]] = []
    for offset, value in items:
        raw_value = safe_float(value)
        if raw_value is None:
            continue
        rem = int(offset)
        coords: list[int] = []
        for size in reversed(sizes):
            coords.append(rem % int(size))
            rem //= int(size)
        coords.reverse()
        dims = {dim_id: inverses[i].get(coords[i], "") for i, dim_id in enumerate(ids)}
        out.append((dims, raw_value))
    return out


def htei_base_observation(
    source_id: str,
    source_group: str,
    component_code: str,
    iso3: str,
    year: int,
    raw_value: float,
    unit: str,
    indicator_code: str,
    dimensions: dict[str, Any],
) -> dict[str, Any] | None:
    if len(iso3) != 3 or not iso2_for(iso3) or year < HISTORICAL_START_YEAR or raw_value is None:
        return None
    return {
        "index_code": "HTEI",
        "component_code": component_code,
        "iso3": iso3.upper(),
        "year": int(year),
        "source_data_year": int(year),
        "raw_value": float(raw_value),
        "unit": unit,
        "source_id": source_id,
        "source_group": source_group,
        "source_priority": HTEI_SOURCE_PRIORITY[source_group],
        "indicator_code": indicator_code,
        "dimensions_json": json.dumps(dimensions, ensure_ascii=False, sort_keys=True),
        "selection_rule": "national_official_then_eurostat_then_oecd_then_ilostat_then_world_bank_uis",
        "quality_flag": "official_observation",
    }


def finalize_htei_observations(
    observations: list[dict[str, Any]],
    snapshot: Snapshot,
    transform_id: str,
    notes: str,
) -> tuple[list[dict[str, Any]], Transform]:
    transform = transform_for(snapshot, transform_id, len(observations), notes)
    finalized: list[dict[str, Any]] = []
    for obs in observations:
        digest = sha256_bytes(
            json.dumps(
                [obs["source_id"], obs["component_code"], obs["iso3"], obs["year"], obs["indicator_code"], obs["dimensions_json"]],
                ensure_ascii=False,
                sort_keys=True,
            ).encode("utf-8")
        )[:14]
        prov = {
            "source_id": snapshot.source_id,
            "source_name": SOURCE_DEFS[snapshot.source_id]["source_name"],
            "source_group": obs["source_group"],
            "source_priority": obs["source_priority"],
            "source_url": snapshot.source_url,
            "retrieved_at": snapshot.retrieved_at,
            "raw_snapshot_path": snapshot.raw_snapshot_path,
            "raw_snapshot_sha256": snapshot.raw_snapshot_sha256,
            "transform_id": transform.transform_id,
            "transformation_run_id": transform.transformation_run_id,
            "selection_rule": obs["selection_rule"],
            "dimensions_json": obs["dimensions_json"],
        }
        finalized.append({
            **obs,
            "observation_id": f"HTEI:{obs['component_code']}:{obs['iso3']}:{obs['year']}:{obs['source_id']}:{digest}",
            "normalized_score": None,
            "selected": 0,
            "source_url": snapshot.source_url,
            "retrieved_at": snapshot.retrieved_at,
            "release_year": snapshot.release_year,
            "raw_snapshot_path": snapshot.raw_snapshot_path,
            "raw_snapshot_sha256": snapshot.raw_snapshot_sha256,
            "transformation_run_id": transform.transformation_run_id,
            "transform_id": transform.transform_id,
            "provenance_json": json.dumps(prov, ensure_ascii=False, sort_keys=True),
        })
    return finalized, transform


def parse_ilostat_htei(snapshot: Snapshot) -> tuple[list[dict[str, Any]], Transform]:
    observations: list[dict[str, Any]] = []
    activity: dict[tuple[str, int], dict[str, float]] = {}
    text = manifest_file_path(snapshot, "employment_by_activity").read_text(encoding="utf-8-sig", errors="replace")
    for row in csv.DictReader(io.StringIO(text)):
        if row.get("sex") != "SEX_T":
            continue
        iso3 = (row.get("ref_area") or "").upper()
        year = int(row.get("time") or 0)
        value = safe_float(row.get("obs_value"))
        if value is None:
            continue
        activity.setdefault((iso3, year), {})[row.get("classif1", "")] = value
    for (iso3, year), values in activity.items():
        total = values.get("ECO_ISIC4_TOTAL") or values.get("ECO_SECTOR_TOTAL")
        tech_bits = [values.get("ECO_ISIC4_J"), values.get("ECO_ISIC4_M")]
        if total and any(v is not None for v in tech_bits):
            raw = sum(v for v in tech_bits if v is not None) / total * 100
            obs = htei_base_observation(
                "ILOSTAT_HTEI",
                "ILOSTAT",
                "HT_EMPLOYMENT_SHARE",
                iso3,
                year,
                raw,
                "percent_of_total_employment",
                "EMP_TEMP_SEX_ECO_NB_A:J+M/TOTAL",
                {"sex": "SEX_T", "activity": ["ECO_ISIC4_J", "ECO_ISIC4_M"], "denominator": "ECO_ISIC4_TOTAL"},
            )
            if obs:
                observations.append(obs)

    occupations: dict[tuple[str, int], dict[str, float]] = {}
    text = manifest_file_path(snapshot, "employment_by_occupation").read_text(encoding="utf-8-sig", errors="replace")
    for row in csv.DictReader(io.StringIO(text)):
        if row.get("sex") != "SEX_T":
            continue
        iso3 = (row.get("ref_area") or "").upper()
        year = int(row.get("time") or 0)
        value = safe_float(row.get("obs_value"))
        if value is None:
            continue
        occupations.setdefault((iso3, year), {})[row.get("classif1", "")] = value
    for (iso3, year), values in occupations.items():
        total = values.get("OCU_ISCO08_TOTAL") or values.get("OCU_SKILL_TOTAL")
        professional = values.get("OCU_ISCO08_2")
        technician = values.get("OCU_ISCO08_3")
        high_skill = values.get("OCU_SKILL_L3-4")
        raw = None
        numerator_code: Any = ["OCU_ISCO08_2", "OCU_ISCO08_3"]
        if total and (professional is not None or technician is not None):
            raw = sum(v for v in [professional, technician] if v is not None) / total * 100
        elif total and high_skill is not None:
            raw = high_skill / total * 100
            numerator_code = "OCU_SKILL_L3-4"
        if raw is not None:
            obs = htei_base_observation(
                "ILOSTAT_HTEI",
                "ILOSTAT",
                "HIGH_TECH_OCCUPATIONS",
                iso3,
                year,
                raw,
                "percent_of_total_employment",
                "EMP_TEMP_SEX_OCU_NB_A:ISCO08_2+3/TOTAL",
                {"sex": "SEX_T", "occupation": numerator_code, "denominator": "OCU_ISCO08_TOTAL"},
            )
            if obs:
                observations.append(obs)
    return finalize_htei_observations(observations, snapshot, "ilostat_bulk_to_htei_v3_observations", "ILOSTAT employment by activity and occupation transformed into HTEI v3 source observations.")


def parse_eurostat_htei(snapshot: Snapshot, population: dict[str, dict[int, float]]) -> tuple[list[dict[str, Any]], Transform]:
    observations: list[dict[str, Any]] = []
    htec = json.loads(manifest_file_path(snapshot, "htec_emp_nat2").read_text(encoding="utf-8"))
    for dims, value in jsonstat_records(htec):
        iso3 = iso3_from_iso2(dims.get("geo", ""))
        if not iso3 or dims.get("sex") != "T" or dims.get("nace_r2") != "HTC" or dims.get("unit") != "PC_EMP":
            continue
        obs = htei_base_observation(
            "EUROSTAT_HTEC",
            "EUROSTAT_NATIONAL_STATS",
            "HT_EMPLOYMENT_SHARE",
            iso3,
            int(dims["time"]),
            value,
            "percent_of_total_employment",
            "htec_emp_nat2:HTC:PC_EMP",
            dims,
        )
        if obs:
            observations.append(obs)

    employment = json.loads(manifest_file_path(snapshot, "lfsa_egan2").read_text(encoding="utf-8"))
    employment_totals: dict[tuple[str, int], float] = {}
    for dims, value in jsonstat_records(employment):
        iso3 = iso3_from_iso2(dims.get("geo", ""))
        if (
            iso3
            and dims.get("sex") == "T"
            and dims.get("age") == "Y15-74"
            and dims.get("nace_r2") == "TOTAL"
            and dims.get("unit") == "THS_PER"
        ):
            employment_totals[(iso3, int(dims["time"]))] = value

    hrst = json.loads(manifest_file_path(snapshot, "hrst_st_nsec2").read_text(encoding="utf-8"))
    for dims, value in jsonstat_records(hrst):
        iso3 = iso3_from_iso2(dims.get("geo", ""))
        if (
            not iso3
            or dims.get("category") != "HRSTO"
            or dims.get("nace_r2") != "TOTAL"
            or dims.get("age") != "Y15-74"
            or dims.get("unit") != "THS_PER"
        ):
            continue
        year = int(dims["time"])
        total_employment = employment_totals.get((iso3, year))
        if not total_employment:
            continue
        obs = htei_base_observation(
            "EUROSTAT_HTEC",
            "EUROSTAT_NATIONAL_STATS",
            "HIGH_TECH_OCCUPATIONS",
            iso3,
            year,
            value / total_employment * 100.0,
            "percent_of_total_employment",
            "hrst_st_nsec2:HRSTO:TOTAL:Y15-74:THS_PER/lfsa_egan2:TOTAL:Y15-74:THS_PER",
            {**dims, "denominator_dataset": "lfsa_egan2", "denominator_value_thousand_persons": total_employment},
        )
        if obs:
            observations.append(obs)

    rd = json.loads(manifest_file_path(snapshot, "rd_p_persocc").read_text(encoding="utf-8"))
    rd_personnel: dict[tuple[str, int], dict[str, float]] = {}
    for dims, value in jsonstat_records(rd):
        iso3 = iso3_from_iso2(dims.get("geo", ""))
        if (
            not iso3
            or dims.get("sectperf") != "TOTAL"
            or dims.get("prof_pos") not in {"RSE", "TEC"}
            or dims.get("sex") != "T"
            or dims.get("unit") != "FTE"
        ):
            continue
        rd_personnel.setdefault((iso3, int(dims["time"])), {})[dims["prof_pos"]] = value
    for (iso3, year), positions in rd_personnel.items():
        country_population = population.get(iso3, {}).get(year)
        if not country_population:
            continue
        total_fte = sum(positions.values())
        obs = htei_base_observation(
            "EUROSTAT_HTEC",
            "EUROSTAT_NATIONAL_STATS",
            "RND_PERSONNEL",
            iso3,
            year,
            total_fte / country_population * 1_000_000,
            "rd_personnel_fte_per_million_population",
            "rd_p_persocc:TOTAL:RSE+TEC:T:FTE",
            {"positions": sorted(positions), "total_fte": total_fte, "population": country_population},
        )
        if obs:
            observations.append(obs)
    return finalize_htei_observations(observations, snapshot, "eurostat_api_to_htei_v3_observations", "Eurostat high-technology employment, HRST/LFS occupation share and R&D personnel datasets transformed into HTEI v3 source observations.")


def population_by_iso_year(world_bank_snapshot: Snapshot) -> dict[str, dict[int, float]]:
    records = load_world_bank_series(world_bank_snapshot)
    return {
        iso3: {year: rec["value"] for year, rec in by_year.items()}
        for iso3, by_year in records.get("SP.POP.TOTL", {}).items()
    }


def parse_oecd_htei(snapshot: Snapshot, population: dict[str, dict[int, float]]) -> tuple[list[dict[str, Any]], Transform]:
    observations: list[dict[str, Any]] = []
    rd_text = manifest_file_path(snapshot, "rd_personnel").read_text(encoding="utf-8", errors="replace")
    for row in csv.DictReader(io.StringIO(rd_text)):
        iso3 = (row.get("REF_AREA") or "").upper()
        year = int(row.get("TIME_PERIOD") or 0)
        value = safe_float(row.get("OBS_VALUE"))
        pop = population.get(iso3, {}).get(year)
        if (
            not value
            or not pop
            or row.get("FREQ") != "A"
            or row.get("MEASURE") != "T_RD"
            or row.get("SECT_PERF") != "_T"
            or row.get("FUNCTION") != "_T"
            or row.get("SEX") != "_T"
            or row.get("UNIT_MEASURE") not in {"PS_FTE", "PS"}
        ):
            continue
        obs = htei_base_observation(
            "OECD_HTEI",
            "OECD",
            "RND_PERSONNEL",
            iso3,
            year,
            value / pop * 1_000_000,
            "rd_personnel_per_million_population",
            "OECD_RDS_PERS_FUNC:T_RD:_T:_T",
            {k: row.get(k) for k in ["DATAFLOW", "MEASURE", "SECT_PERF", "FUNCTION", "UNIT_MEASURE", "EMP_STATUS", "SEX"]},
        )
        if obs:
            observations.append(obs)

    msti_text = manifest_file_path(snapshot, "msti").read_text(encoding="utf-8", errors="replace")
    for row in csv.DictReader(io.StringIO(msti_text)):
        iso3 = (row.get("REF_AREA") or "").upper()
        year = int(row.get("TIME_PERIOD") or 0)
        value = safe_float(row.get("OBS_VALUE"))
        if (
            value is None
            or row.get("FREQ") != "A"
            or row.get("MEASURE") != "B"
            or row.get("UNIT_MEASURE") != "PT_GERD"
            or row.get("TRANSFORMATION") != "_Z"
        ):
            continue
        obs = htei_base_observation(
            "OECD_HTEI",
            "OECD",
            "CORPORATE_STRATEGY_AND_DEMAND",
            iso3,
            year,
            value,
            "business_rd_percent_of_gerd",
            "OECD_MSTI:B:PT_GERD",
            {k: row.get(k) for k in ["DATAFLOW", "MEASURE", "UNIT_MEASURE", "TRANSFORMATION"]},
        )
        if obs:
            observations.append(obs)
    return finalize_htei_observations(observations, snapshot, "oecd_sdmx_to_htei_v3_observations", "OECD SDMX MSTI and R&D personnel datasets transformed into HTEI v3 source observations.")


def parse_uis_htei(snapshot: Snapshot) -> tuple[list[dict[str, Any]], Transform]:
    observations: list[dict[str, Any]] = []
    payload = json.loads(snapshot_path(snapshot).read_text(encoding="utf-8"))
    api_response = payload.get("response") if isinstance(payload, dict) else None
    if isinstance(api_response, dict) and isinstance(api_response.get("records"), list):
        for rec in api_response["records"]:
            iso3 = str(rec.get("geoUnit") or "").upper()
            year = int(rec.get("year") or 0)
            value = safe_float(rec.get("value"))
            if value is None or not iso2_for(iso3):
                continue
            obs = htei_base_observation(
                "UIS_HTEI",
                "WORLD_BANK_UIS",
                "STEM_PIPELINE",
                iso3,
                year,
                value,
                "percent_of_tertiary_graduates",
                UIS_STEM_INDICATOR,
                {
                    "indicator": UIS_STEM_INDICATOR,
                    "api_version": payload.get("api_version"),
                    "magnitude": rec.get("magnitude"),
                    "qualifier": rec.get("qualifier"),
                },
            )
            if obs:
                observations.append(obs)
    else:
        records = load_world_bank_series(snapshot)
        for iso3, by_year in records.get("UIS.FOSGP.5T8.F500600700", {}).items():
            for year, rec in by_year.items():
                obs = htei_base_observation(
                    "UIS_HTEI",
                    "WORLD_BANK_UIS",
                    "STEM_PIPELINE",
                    iso3,
                    year,
                    rec["value"],
                    "percent_of_tertiary_graduates",
                    "UIS.FOSGP.5T8.F500600700",
                    {"indicator": "UIS.FOSGP.5T8.F500600700", "legacy_channel": "World Bank Indicators API"},
                )
                if obs:
                    observations.append(obs)
    return finalize_htei_observations(observations, snapshot, "uis_indicator_to_htei_v3_observations", "UNESCO UIS STEM tertiary graduates indicator transformed into HTEI v3 source observations.")


def parse_world_bank_htei_observations(snapshot: Snapshot) -> tuple[list[dict[str, Any]], Transform]:
    observations: list[dict[str, Any]] = []
    records = load_world_bank_series(snapshot)
    iso_set = set().union(*(set(v.keys()) for v in records.values()))
    tech_inputs: dict[int, dict[str, dict[str, float]]] = {}
    for iso3 in iso_set:
        years = sorted({year for by_year in records.values() for year in by_year.get(iso3, {})})
        for year in years:
            researchers = records.get("SP.POP.SCIE.RD.P6", {}).get(iso3, {}).get(year, {}).get("value")
            technicians = records.get("SP.POP.TECH.RD.P6", {}).get(iso3, {}).get(year, {}).get("value")
            if researchers is not None or technicians is not None:
                obs = htei_base_observation(
                    "WORLD_BANK_HTEI",
                    "WORLD_BANK_UIS",
                    "RND_PERSONNEL",
                    iso3,
                    year,
                    sum(v for v in [researchers, technicians] if v is not None),
                    "rd_personnel_per_million_population",
                    "SP.POP.SCIE.RD.P6+SP.POP.TECH.RD.P6",
                    {"indicators": ["SP.POP.SCIE.RD.P6", "SP.POP.TECH.RD.P6"]},
                )
                if obs:
                    observations.append(obs)
            ht_export = records.get("TX.VAL.TECH.MF.ZS", {}).get(iso3, {}).get(year, {}).get("value")
            ict_exports = records.get("BX.GSR.CCIS.ZS", {}).get(iso3, {}).get(year, {}).get("value")
            patents = records.get("IP.PAT.RESD", {}).get(iso3, {}).get(year, {}).get("value")
            pop = records.get("SP.POP.TOTL", {}).get(iso3, {}).get(year, {}).get("value")
            if ht_export is not None:
                tech_inputs.setdefault(year, {}).setdefault("TX.VAL.TECH.MF.ZS", {})[iso3] = ht_export
            if ict_exports is not None:
                tech_inputs.setdefault(year, {}).setdefault("BX.GSR.CCIS.ZS", {})[iso3] = ict_exports
            if patents is not None and pop:
                tech_inputs.setdefault(year, {}).setdefault("IP.PAT.RESD_per_million", {})[iso3] = patents / pop * 1_000_000

    for year, indicator_values in sorted(tech_inputs.items()):
        normalized = {code: normalize_values(values) for code, values in indicator_values.items()}
        countries = sorted(set().union(*(set(values) for values in indicator_values.values())))
        for iso3 in countries:
            available = [code for code in normalized if iso3 in normalized[code]]
            if not available:
                continue
            normalized_values = {code: normalized[code][iso3] for code in available}
            raw_values = {code: indicator_values[code][iso3] for code in available}
            obs = htei_base_observation(
                "WORLD_BANK_HTEI",
                "WORLD_BANK_UIS",
                "TECH_OUTPUTS",
                iso3,
                year,
                sum(normalized_values.values()) / len(normalized_values),
                "equal_weight_mean_of_individually_minmax_normalized_indicators_0_100",
                "+".join(available),
                {
                    "indicators": available,
                    "raw_values": raw_values,
                    "normalized_values": normalized_values,
                    "aggregation": "equal weight over available indicators after country-year min-max normalization",
                },
            )
            if obs:
                observations.append(obs)
    return finalize_htei_observations(observations, snapshot, "world_bank_api_to_htei_v3_observations", "World Bank high-technology exports, ICT services, patents and R&D personnel transformed into HTEI v3 source observations.")


def parse_hci(snapshot: Snapshot) -> tuple[list[dict[str, Any]], list[dict[str, Any]], Transform]:
    records = load_world_bank_series(snapshot)
    official = records["HD.HCI.OVRL"]
    scores: list[dict[str, Any]] = []
    for iso3, by_year in official.items():
        for year, rec in by_year.items():
            scores.append({"iso3": iso3, "year": year, "score": rec["value"] * 100, "rank": 0, "source_data_year": year})
    for year_scores in group_scores_by_year(scores).values():
        assign_ranks(year_scores)
    components: list[dict[str, Any]] = []
    for iso3, by_year in official.items():
        for year in by_year:
            mort = records.get("HD.HCI.MORT", {}).get(iso3, {}).get(year, {}).get("value")
            amrt = records.get("HD.HCI.AMRT", {}).get(iso3, {}).get(year, {}).get("value")
            eyrs = records.get("HD.HCI.EYRS", {}).get(iso3, {}).get(year, {}).get("value")
            hlos = records.get("HD.HCI.HLOS", {}).get(iso3, {}).get(year, {}).get("value")
            stnt = records.get("HD.HCI.STNT", {}).get(iso3, {}).get(year, {}).get("value")
            surv_vals = [v for v in [mort, amrt] if v is not None]
            component_defs = {
                "SURV": ((sum(surv_vals) / len(surv_vals) * 100) if surv_vals else None, "0-1 survival scale"),
                "SCHOOL": (eyrs / 14 * 100 if eyrs is not None else None, "years"),
                "LEARN": (hlos / 625 * 100 if hlos is not None else None, "harmonized test score"),
                "HEALTH": (stnt * 100 if stnt is not None else None, "0-1 not-stunted scale"),
            }
            for code, (norm, unit) in component_defs.items():
                if norm is None:
                    continue
                raw_lookup = {"SURV": norm / 100, "SCHOOL": eyrs, "LEARN": hlos, "HEALTH": stnt}
                components.append({
                    "iso3": iso3,
                    "year": year,
                    "component_code": code,
                    "raw_value": raw_lookup[code],
                    "normalized_score": bounded_score(norm),
                    "unit": unit,
                    "source_data_year": year,
                })
    transform = transform_for(snapshot, "world_bank_hci_api_to_timeseries", len(scores), "World Bank HCI API indicators parsed for all non-empty official HCI years and components.")
    return scores, components, transform


def parse_gii(snapshot: Snapshot) -> tuple[list[dict[str, Any]], list[dict[str, Any]], Transform]:
    wb = openpyxl.load_workbook(snapshot_path(snapshot), read_only=True, data_only=True)
    ws = wb["Data"]
    component_map = {
        "IN.1": "INST",
        "IN.2": "HCR",
        "IN.3": "INFRA",
        "IN.4": "MARKET",
        "IN.5": "BUS",
        "OUT.6": "KTO",
        "OUT.7": "CRE",
    }
    scores: list[dict[str, Any]] = []
    components: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        iso3, _economy, num, _name, data_year, _value_screen, score, rank, *_rest = row
        iso3 = str(iso3).upper() if iso3 else ""
        if len(iso3) != 3:
            continue
        if num is None and score is not None and rank is not None:
            scores.append({"iso3": iso3, "year": snapshot.release_year, "score": float(score), "rank": int(rank), "source_data_year": snapshot.release_year})
        elif num in component_map and score is not None:
            components.append({
                "iso3": iso3,
                "year": snapshot.release_year,
                "component_code": component_map[num],
                "raw_value": float(score),
                "normalized_score": float(score),
                "unit": "0-100 WIPO pillar score",
                "source_data_year": int(data_year) if isinstance(data_year, int) else snapshot.release_year,
            })
    transform = transform_for(snapshot, f"wipo_gii_{snapshot.release_year}_xlsx_to_scores", len(scores), f"WIPO GII {snapshot.release_year} database parsed for official score, rank and pillar scores.")
    return scores, components, transform


def parse_idi(snapshot: Snapshot) -> tuple[list[dict[str, Any]], list[dict[str, Any]], Transform]:
    wb = openpyxl.load_workbook(snapshot_path(snapshot), read_only=True, data_only=True)
    scores: list[dict[str, Any]] = []
    components: list[dict[str, Any]] = []
    for sheet_name in [name for name in wb.sheetnames if re.match(r"IDI \d{4} Data", name)]:
        ws = wb[sheet_name]
        release_year = int(re.search(r"(\d{4})", sheet_name).group(1))
        source_data_year = release_year - 2
        rows_list = list(ws.iter_rows(values_only=True))
        header_idx = None
        for idx, row in enumerate(rows_list):
            if row and str(row[0]).strip().lower() == "iso3":
                header_idx = idx
                break
        if header_idx is None:
            continue
        for row in rows_list[header_idx + 2:]:
            iso3 = str(row[0]).upper() if row and row[0] else ""
            if len(iso3) != 3:
                continue
            score = safe_float(row[37] if len(row) > 37 else None)
            if score is None and len(row) > 39:
                score = safe_float(row[39])
            if score is None:
                continue
            scores.append({"iso3": iso3, "year": release_year, "score": score, "rank": 0, "source_data_year": source_data_year})
            access = safe_float(row[35] if len(row) > 35 else None)
            meaning = safe_float(row[36] if len(row) > 36 else None)
            use_candidates = [safe_float(row[i]) for i in [25, 26, 27, 33] if len(row) > i]
            use_vals = [v for v in use_candidates if v is not None]
            comp_defs = {
                "ACCESS": (access, "0-100 ITU pillar score"),
                "USE": ((sum(use_vals) / len(use_vals)) if use_vals else None, "0-100 ITU normalized use score"),
                "MEANING": (meaning, "0-100 ITU pillar score"),
            }
            for code, (val, unit) in comp_defs.items():
                if val is None:
                    continue
                components.append({
                    "iso3": iso3,
                    "year": release_year,
                    "component_code": code,
                    "raw_value": val,
                    "normalized_score": val,
                    "unit": unit,
                    "source_data_year": source_data_year,
                })
    for year_scores in group_scores_by_year(scores).values():
        assign_ranks(year_scores)
    transform = transform_for(snapshot, "itu_idi_2025_xlsx_to_timeseries", len(scores), "ITU IDI workbook parsed for 2023, 2024 and 2025 aggregate and pillar scores.")
    return scores, components, transform


def parse_gtci(snapshot: Snapshot) -> tuple[list[dict[str, Any]], list[dict[str, Any]], Transform]:
    if pdfplumber is None:
        raise RuntimeError("pdfplumber is required to parse the official GTCI report")
    scores: list[dict[str, Any]] = []
    components: list[dict[str, Any]] = []
    with pdfplumber.open(snapshot_path(snapshot)) as pdf:
        ranking_text = "\n".join((pdf.pages[p - 1].extract_text() or "") for p in [26, 27, 28])
        pillar_text = "\n".join((pdf.pages[p].extract_text() or "") for p in [79, 80, 81])

    pattern = re.compile(r"^\s*(\d{1,3})\s+(.+?)\s+(\d{2}\.\d{2})\s+(High|Upper middle|Lower middle|Low)\s+(.+)$")
    for line in ranking_text.splitlines():
        match = pattern.match(line)
        if not match:
            continue
        rank, name, score, _income, _region = match.groups()
        iso3 = iso3_from_name(name)
        if not iso3:
            continue
        scores.append({"iso3": iso3, "year": snapshot.release_year, "score": float(score), "rank": int(rank), "source_data_year": snapshot.release_year})

    # Official GTCI report, Table 1 "Rankings by Pillar". The PDF exposes pillar ranks,
    # not pillar scores; the platform therefore stores the official rank as raw_value
    # and derives a transparent 0-100 diagnostic score from the official rank.
    pillar_codes = ["ENABLE", "ATTRACT", "GROW", "RETAIN", "VT", "GA"]
    pillar_rank_rows: list[tuple[str, dict[str, int]]] = []
    pillar_line = re.compile(r"^(.+?)\s+(\d{1,3})\s+(\d{1,3})\s+(\d{1,3})\s+(\d{1,3})\s+(\d{1,3})\s+(\d{1,3})\s+(\d{1,3})\s*$")
    ignore_prefixes = {"Economy", "Table", "Detailed", "Talent", "stluseR", "deliateD", "noitpursiD", "fo", "arE", "na", "gnitagivaN", ":ecneiliseR", "dna", "tnelaT"}
    for raw_line in pillar_text.splitlines():
        line = " ".join(raw_line.split())
        if not line or any(line.startswith(prefix) for prefix in ignore_prefixes) or line.isdigit():
            continue
        match = pillar_line.match(line)
        if not match:
            continue
        name = match.group(1).strip()
        nums = [int(match.group(i)) for i in range(2, 9)]
        iso3 = iso3_from_name(name)
        if not iso3:
            continue
        pillar_rank_rows.append((iso3, dict(zip(pillar_codes, nums[1:]))))
    max_rank = max((rank for _iso, ranks in pillar_rank_rows for rank in ranks.values()), default=135)
    for iso3, ranks in pillar_rank_rows:
        if not any(s["iso3"] == iso3 for s in scores):
            continue
        for code, rank_value in ranks.items():
            components.append({
                "iso3": iso3,
                "year": snapshot.release_year,
                "component_code": code,
                "raw_value": rank_value,
                "unit": "official_pillar_rank_lower_is_better",
                "normalized_score": round((max_rank - rank_value + 1) / max_rank * 100, 6),
                "source_data_year": snapshot.release_year,
            })
    transform = transform_for(
        snapshot,
        "portulans_gtci_2025_pdf_rank_and_pillar_extract",
        len(scores) + len(components),
        "Official GTCI 2025 PDF ranking table and Table 1 pillar ranks extracted; pillar rank is converted to a transparent diagnostic 0-100 scale.",
    )
    return scores, components, transform


def parse_htei(snapshot: Snapshot) -> tuple[list[dict[str, Any]], list[dict[str, Any]], Transform]:
    records = load_world_bank_series(snapshot)
    raw: dict[int, dict[str, dict[str, float]]] = {}
    iso_set = set().union(*(set(v.keys()) for v in records.values()))
    all_years = sorted({
        year
        for by_iso in records.values()
        for by_year in by_iso.values()
        for year in by_year
        if year >= HISTORICAL_START_YEAR
    })
    for year in all_years:
        for iso3 in iso_set:
            prof = records.get("JI.EMP.PROF.ZS", {}).get(iso3, {}).get(year, {}).get("value")
            tech = records.get("JI.EMP.TECH.ZS", {}).get(iso3, {}).get(year, {}).get("value")
            researchers = records.get("SP.POP.SCIE.RD.P6", {}).get(iso3, {}).get(year, {}).get("value")
            technicians = records.get("SP.POP.TECH.RD.P6", {}).get(iso3, {}).get(year, {}).get("value")
            stem = records.get("UIS.FOSGP.5T8.F500600700", {}).get(iso3, {}).get(year, {}).get("value")
            ht_export = records.get("TX.VAL.TECH.MF.ZS", {}).get(iso3, {}).get(year, {}).get("value")
            ict_exports = records.get("BX.GSR.CCIS.ZS", {}).get(iso3, {}).get(year, {}).get("value")
            patents = records.get("IP.PAT.RESD", {}).get(iso3, {}).get(year, {}).get("value")
            pop = records.get("SP.POP.TOTL", {}).get(iso3, {}).get(year, {}).get("value")
            comp: dict[str, float] = {}
            if prof is not None or tech is not None:
                comp["HT_EMP"] = sum(v for v in [prof, tech] if v is not None)
            if tech is not None:
                comp["STEM_OCC"] = tech
            if researchers is not None or technicians is not None:
                comp["RD_PERSONNEL"] = sum(v for v in [researchers, technicians] if v is not None)
            if stem is not None:
                comp["STEM_PIPE"] = stem
            demand_bits = []
            if ict_exports is not None:
                demand_bits.append(ict_exports)
            if patents is not None and pop:
                demand_bits.append(patents / pop * 1_000_000)
            if demand_bits:
                comp["DEMAND_OUTPUT"] = sum(demand_bits) / len(demand_bits)
            if ht_export is not None:
                comp["HT_EXPORT"] = ht_export
            if comp:
                raw.setdefault(year, {})[iso3] = comp
    by_component: dict[int, dict[str, dict[str, float]]] = {}
    for year, by_iso in raw.items():
        for iso3, values in by_iso.items():
            for code, value in values.items():
                by_component.setdefault(year, {}).setdefault(code, {})[iso3] = value
    normalized = {
        year: {code: normalize_values(values) for code, values in comps.items()}
        for year, comps in by_component.items()
    }
    weights = {code: weight for code, _ru, _en, weight, _note in COMPONENTS["HTEI"]}
    scores: list[dict[str, Any]] = []
    components: list[dict[str, Any]] = []
    for year, by_iso in raw.items():
        for iso3, values in by_iso.items():
            available_weight = sum(weights[code] for code in values if code in normalized.get(year, {}) and iso3 in normalized[year][code])
            if available_weight < 0.45:
                continue
            score = sum(normalized[year][code][iso3] * weights[code] for code in values if code in normalized.get(year, {}) and iso3 in normalized[year][code]) / available_weight
            scores.append({"iso3": iso3, "year": year, "score": score, "rank": 0, "source_data_year": year, "data_quality": min(1.0, available_weight)})
            for code, raw_value in values.items():
                norm = normalized.get(year, {}).get(code, {}).get(iso3)
                if norm is None:
                    continue
                components.append({
                    "iso3": iso3,
                    "year": year,
                    "component_code": code,
                    "raw_value": raw_value,
                    "normalized_score": norm,
                    "unit": "official source native unit",
                    "source_data_year": year,
                    "data_quality": min(1.0, available_weight),
                })
    for year_scores in group_scores_by_year(scores).values():
        assign_ranks(year_scores)
    transform = transform_for(snapshot, "world_bank_indicators_to_htei_v2_tz_indicators", len(scores), "HTEI v2 computed by year from official World Bank, Jobs Indicators and UIS-coded indicators as the TЗ High-tech employment index without carry-forward.")
    return scores, components, transform


def parse_htei_multisource(
    snapshots: dict[str, Snapshot],
    htei_snapshot: Snapshot,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]], list[Transform]]:
    population = population_by_iso_year(snapshots["WORLD_BANK_HTEI"])
    source_observations: list[dict[str, Any]] = []
    transforms: list[Transform] = []
    for obs, transform in [
        parse_ilostat_htei(snapshots["ILOSTAT_HTEI"]),
        parse_eurostat_htei(snapshots["EUROSTAT_HTEC"], population),
        parse_oecd_htei(snapshots["OECD_HTEI"], population),
        parse_uis_htei(snapshots["UIS_HTEI"]),
        parse_world_bank_htei_observations(snapshots["WORLD_BANK_HTEI"]),
    ]:
        source_observations.extend(obs)
        transforms.append(transform)

    grouped: dict[tuple[str, int], dict[str, float]] = {}
    for obs in source_observations:
        grouped.setdefault((obs["component_code"], obs["year"]), {})[obs["observation_id"]] = obs["raw_value"]
    norm_by_id: dict[str, float] = {}
    for raw_values in grouped.values():
        norm_by_id.update(normalize_values(raw_values))
    for obs in source_observations:
        obs["normalized_score"] = round(norm_by_id.get(obs["observation_id"], 0.0), 6)

    selected_by_key: dict[tuple[str, str, int], dict[str, Any]] = {}
    for obs in sorted(
        source_observations,
        key=lambda item: (item["source_priority"], -int(item["source_data_year"]), item["indicator_code"], item["observation_id"]),
    ):
        key = (obs["component_code"], obs["iso3"], obs["year"])
        if key not in selected_by_key:
            selected_by_key[key] = obs
            obs["selected"] = 1

    selected_by_country_year: dict[tuple[str, int], list[dict[str, Any]]] = {}
    for obs in selected_by_key.values():
        selected_by_country_year.setdefault((obs["iso3"], obs["year"]), []).append(obs)

    weights = HTEI_V3_WEIGHTS
    scores: list[dict[str, Any]] = []
    components: list[dict[str, Any]] = []
    for (iso3, year), items in sorted(selected_by_country_year.items()):
        available_weight = sum(weights.get(item["component_code"], 0.0) for item in items)
        groups = sorted({item["source_group"] for item in items})
        if (
            available_weight < HTEI_RELEASE_MIN_WEIGHT
            or len(items) < HTEI_RELEASE_MIN_COMPONENTS
            or len(groups) < HTEI_RELEASE_MIN_SOURCE_GROUPS
        ):
            continue
        weighted = sum(item["normalized_score"] * weights[item["component_code"]] for item in items)
        weighted_average = weighted / available_weight
        group_factor = min(1.0, len(groups) / 4)
        data_quality = min(1.0, 0.5 * available_weight + 0.5 * group_factor)
        quality_adjustment = 0.85 + 0.15 * data_quality
        score = bounded_score(weighted_average * quality_adjustment) or 0.0
        scores.append({
            "iso3": iso3,
            "year": year,
            "score": score,
            "rank": 0,
            "source_data_year": year,
            "data_quality": data_quality,
            "source_group_coverage": groups,
            "available_weight": available_weight,
        })
        for item in items:
            components.append({
                "iso3": iso3,
                "year": year,
                "component_code": item["component_code"],
                "raw_value": item["raw_value"],
                "normalized_score": item["normalized_score"],
                "unit": item["unit"],
                "source_data_year": item["source_data_year"],
                "data_quality": data_quality,
                "source_id": item["source_id"],
                "source_group": item["source_group"],
                "source_priority": item["source_priority"],
                "source_url": item["source_url"],
                "retrieved_at": item["retrieved_at"],
                "release_year": item["release_year"],
                "raw_snapshot_path": item["raw_snapshot_path"],
                "raw_snapshot_sha256": item["raw_snapshot_sha256"],
                "transformation_run_id": item["transformation_run_id"],
                "transform_id": item["transform_id"],
                "quality_flag": "selected_official_observation",
                "provenance_json": item["provenance_json"],
                "source_note": f"Selected from {item['source_group']} by HTEI v3 source cascade.",
            })
    for year_scores in group_scores_by_year(scores).values():
        assign_ranks(year_scores)
    htei_transform = transform_for(
        htei_snapshot,
        "htei_v3_multisource_country_score",
        len(scores) + len(components),
        "HTEI v3 computed from ILOSTAT, OECD, Eurostat, UIS and World Bank source observations with no hidden carry-forward.",
    )
    transforms.insert(0, htei_transform)
    return scores, components, source_observations, transforms


def parse_rank(value: Any) -> int | None:
    if value is None:
        return None
    text = str(value).strip()
    match = re.search(r"\d+", text.replace("=", ""))
    return int(match.group(0)) if match else None


def flatten_qs_indicators(node: dict[str, Any]) -> dict[str, Any]:
    out: dict[str, Any] = {}
    scores = node.get("scores") or {}
    mapping = {
        "Employer Reputation": "EMPLOYER_REPUTATION",
        "Academic Reputation": "ACADEMIC_REPUTATION",
        "H-index Citations": "H_INDEX_CITATIONS",
        "Citations per Paper": "CITATIONS_PER_PAPER",
        "International Research Network": "INTERNATIONAL_RESEARCH_NETWORK",
    }
    for group in scores.values():
        if not isinstance(group, dict):
            continue
        for name, value in group.items():
            code = mapping.get(str(name), re.sub(r"[^A-Z0-9]+", "_", str(name).upper()).strip("_"))
            out[code] = value
    return out


def parse_qs_endpoint(snapshot: Snapshot) -> tuple[list[dict[str, Any]], list[dict[str, Any]], Transform, list[dict[str, Any]], list[str]]:
    manifest_path = snapshot_path(snapshot)
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    year = int(manifest["year"])
    base_dir = manifest_path.parent
    institution_rows: list[dict[str, Any]] = []
    unmatched: list[str] = []
    country_values: dict[str, list[dict[str, float]]] = {}
    loaded_records = 0
    for file_name in manifest.get("page_files", []):
        page_path = base_dir / file_name
        page_digest = sha256_bytes(page_path.read_bytes())
        page = json.loads(page_path.read_text(encoding="utf-8"))
        nodes = page.get("score_nodes") or []
        loaded_records += len(nodes)
        for node in nodes:
            country = str(node.get("country") or "").strip()
            title = str(node.get("title") or node.get("name") or "").strip()
            if not country or not title:
                continue
            iso3 = iso3_from_name(country)
            rank = parse_rank(node.get("rank") or node.get("rank_display"))
            score = safe_float(node.get("overall_score"))
            indicators = flatten_qs_indicators(node)
            institution_rows.append({
                "year": year,
                "score_nid": str(node.get("score_nid") or ""),
                "nid": str(node.get("nid") or ""),
                "core_id": str(node.get("core_id") or ""),
                "title": title,
                "qs_path": node.get("path"),
                "region": node.get("region"),
                "country": country,
                "iso3": iso3,
                "city": node.get("city"),
                "overall_score": score,
                "rank_display": str(node.get("rank_display") or node.get("rank") or ""),
                "rank": rank,
                "indicators_json": json.dumps(indicators, ensure_ascii=False, sort_keys=True),
                "source_id": "QS_ET",
                "source_url": manifest["endpoint_base"],
                "retrieved_at": snapshot.retrieved_at,
                "raw_snapshot_path": rel_path(page_path),
                "raw_snapshot_sha256": page_digest,
                "quality_flag": "official_value",
            })
            if iso3 and rank is not None and score is not None:
                country_values.setdefault(iso3, []).append({"rank": float(rank), "score": float(score)})
            elif country and not iso3:
                unmatched.append(country)
    expected = int(manifest.get("total_record") or 0)
    if expected and loaded_records != expected:
        raise RuntimeError(f"QS {year} endpoint row count mismatch: loaded {loaded_records}, expected {expected}")

    raw_components: dict[str, dict[str, float]] = {
        "TOP_COUNT": {},
        "BEST_RANK": {},
        "MEDIAN_RANK": {},
        "QS_SCORE": {},
    }
    for iso3, values in country_values.items():
        ranks = [item["rank"] for item in values]
        scores = [item["score"] for item in values]
        raw_components["TOP_COUNT"][iso3] = float(len(values))
        raw_components["BEST_RANK"][iso3] = min(ranks)
        raw_components["MEDIAN_RANK"][iso3] = float(median(ranks))
        raw_components["QS_SCORE"][iso3] = sum(scores)
    normalized = {
        "TOP_COUNT": normalize_values(raw_components["TOP_COUNT"]),
        "BEST_RANK": normalize_values(raw_components["BEST_RANK"], reverse=True),
        "MEDIAN_RANK": normalize_values(raw_components["MEDIAN_RANK"], reverse=True),
        "QS_SCORE": normalize_values(raw_components["QS_SCORE"]),
    }
    weights = {code: weight for code, _ru, _en, weight, _note in COMPONENTS["QS_ET"]}
    scores: list[dict[str, Any]] = []
    components: list[dict[str, Any]] = []
    for iso3 in sorted(country_values):
        available = [code for code in raw_components if iso3 in normalized[code]]
        if len(available) != len(raw_components):
            continue
        country_score = sum(normalized[code][iso3] * weights[code] for code in available)
        scores.append({"iso3": iso3, "year": year, "score": country_score, "rank": 0, "source_data_year": year})
        for code in available:
            components.append({
                "iso3": iso3,
                "year": year,
                "component_code": code,
                "raw_value": raw_components[code][iso3],
                "normalized_score": normalized[code][iso3],
                "unit": "QS public endpoint country aggregation input",
                "source_data_year": year,
            })
    assign_ranks(scores)
    transform = transform_for(snapshot, f"qs_topuniversities_endpoint_{year}_to_country_aggregation", len(scores), "QS TopUniversities public endpoint archived and aggregated to country scores with TOP_COUNT, BEST_RANK, MEDIAN_RANK and QS_SCORE.")
    return scores, components, transform, institution_rows, sorted(set(unmatched))


def qs_page_snapshots(snapshot: Snapshot) -> list[Snapshot]:
    manifest_path = snapshot_path(snapshot)
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    base_dir = manifest_path.parent
    out: list[Snapshot] = []
    for file_name in manifest.get("page_files", []):
        page_path = base_dir / file_name
        content = page_path.read_bytes()
        page_no = Path(file_name).stem.replace("page_", "")
        page_snapshot = Snapshot(
            snapshot_id=f"QS_ET:{manifest['year']}:page:{page_no}:{sha256_bytes(content)[:16]}",
            source_id="QS_ET",
            release_year=int(manifest["year"]),
            retrieved_at=snapshot.retrieved_at,
            source_url=f"{manifest['endpoint_base']}&page={int(page_no)}",
            raw_snapshot_path=rel_path(page_path),
            raw_snapshot_sha256=sha256_bytes(content),
            content_type="application/json",
            bytes_count=len(content),
            license_or_terms=SOURCE_DEFS["QS_ET"]["license_or_terms"],
        )
        page_path.with_suffix(page_path.suffix + ".metadata.json").write_text(json.dumps(page_snapshot.__dict__, ensure_ascii=False, indent=2), encoding="utf-8")
        out.append(page_snapshot)
    return out


def assign_ranks(scores: list[dict[str, Any]]) -> None:
    ordered = sorted(scores, key=lambda item: item["score"], reverse=True)
    total = len(ordered)
    for idx, item in enumerate(ordered, start=1):
        item["rank"] = idx
        item["percentile"] = round((total - idx + 1) / total * 100, 4) if total else None


def formula_rows() -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    formula_text = {
        "HDI": (
            "Официальная оценка ИЧР UNDP; диагностические компоненты нормируются из life expectancy, schooling и GNI.",
            "Official UNDP HDI score; diagnostic components are normalized from life expectancy, schooling and GNI.",
            "official-hdi-timeseries-2025",
            "UNDP_HDR",
            1,
        ),
        "HCI": (
            "Официальная оценка World Bank HCI; компоненты берутся из доступных индикаторов HCI.",
            "Official World Bank HCI score; components use available HCI indicators.",
            "official-hci-series",
            "WORLD_BANK_HCI",
            1,
        ),
        "GTCI": (
            "Официальная оценка GTCI 2025 извлечена из таблицы рейтинга; шесть направлений диагностики рассчитаны из официальных рангов по направлениям Enable, Attract, Grow, Retain, VT и GA из таблицы Rankings by Pillar.",
            "Official GTCI 2025 score is extracted from the ranking table; six diagnostic pillars are derived from official pillar ranks in the Rankings by Pillar table.",
            "official-gtci-2025-pillars",
            "PORTULANS_GTCI",
            1,
        ),
        "GII": (
            "Официальная оценка WIPO GII и семь направлений из базы GII 2025.",
            "Official WIPO GII score and seven pillar scores from the GII 2025 database.",
            "official-gii-2025",
            "WIPO_GII",
            1,
        ),
        "IDI": (
            "Официальная оценка ITU IDI и направления из набора IDI 2025.",
            "Official ITU IDI score and pillar scores from the IDI 2025 dataset.",
            "official-idi-series-2023-2025",
            "ITU_IDI",
            1,
        ),
        "QS_ET": (
            "Страновая агрегация рассчитывается из архивированных строк официального endpoint QS: число вузов, лучшая позиция, медианная позиция и суммарная оценка вузов.",
            "Country aggregation is computed from archived official QS endpoint rows: count, best rank, median rank and aggregated institutional score.",
            "qs-country-aggregation-v1",
            "QS_ET",
            0,
        ),
        "HTEI": (
            "Индекс занятости в высокотехнологичных отраслях v3 = сумма нормированных компонентов ILOSTAT, OECD, Eurostat/национальной статистики, UIS и World Bank, умноженных на веса и коэффициент качества покрытия; значения не переносятся между годами без явного source_data_year.",
            "High-Tech Employment Index v3 = weighted normalized components from ILOSTAT, OECD, Eurostat/national statistics, UIS and World Bank, adjusted by source coverage quality; values are not silently carried across years.",
            "htei-v3-multisource",
            "HTEI_MULTI_SOURCE",
            1,
        ),
    }
    for index in INDEX_DEFINITIONS:
        code = index["code"]
        ru, en, version, source_id, official_formula = formula_text[code]
        weights = {c: w for c, _ru, _en, w, _note in COMPONENTS[code]}
        rows.append({
            "index_code": code,
            "formula_version": version,
            "formula_text_ru": ru,
            "formula_text_en": en,
            "method_notes_ru": index["description_ru"],
            "method_notes_en": index["description_en"],
            "weights_json": json.dumps(weights, ensure_ascii=False, sort_keys=True),
            "normalization_ru": "Официальные оценки 0-100 используются напрямую; внешние исходные компоненты нормируются min-max по доступным странам.",
            "normalization_en": "Official 0-100 scores are used directly; external raw components are min-max normalized across available countries.",
            "official_formula_available": official_formula,
            "source_id": source_id,
        })
    return rows


def source_registry_rows(snapshots: dict[str, Snapshot]) -> list[dict[str, Any]]:
    rows_out: list[dict[str, Any]] = []
    for source_id, source in SOURCE_DEFS.items():
        snap = snapshots.get(source_id)
        source_role = source.get("source_role") or source_role_for(source_id, source["automation_status"])
        rows_out.append({
            "source_id": source_id,
            "source_code": source_id,
            "source_name": source["source_name"],
            "title": source["source_name"],
            "owner": source["owner"],
            "url": source["url"],
            "source_url": source["source_url"],
            "access_mode": source["access_mode"],
            "update_frequency": source["update_frequency"],
            "license_or_terms": source["license_or_terms"],
            "license_note": source["license_or_terms"],
            "automation_status": source["automation_status"],
            "source_role": source_role,
            "retrieved_at": snap.retrieved_at if snap else None,
            "release_year": source["release_year"],
            "latest_snapshot_id": snap.snapshot_id if snap else None,
            "is_official": int(source_role not in {"audit_register", "qualitative_evidence"}),
            "free_access": 1,
        })
    return rows_out


def country_rows(iso_to_name: dict[str, str], iso_to_region: dict[str, str], iso_to_income: dict[str, str]) -> list[dict[str, Any]]:
    rows_out = []
    for iso3 in sorted(iso_to_name):
        iso2 = iso2_for(iso3)
        if not iso2:
            continue
        name_en = country_name_en(iso3, iso_to_name.get(iso3))
        name_ru = SPECIAL_COUNTRY_NAMES_RU.get(iso3) or RU_COUNTRY_NAMES.get(iso3) or RU_LOCALE.territories.get(iso2) or name_en
        rows_out.append({
            "iso3": iso3,
            "iso2": iso2,
            "name_ru": name_ru,
            "name_en": name_en,
            "region": iso_to_region.get(iso3, "Not classified"),
            "income_group": iso_to_income.get(iso3, "Not classified"),
            "population_m": None,
            "flag": flag_emoji(iso2),
        })
    return rows_out


def index_rows() -> list[dict[str, Any]]:
    return INDEX_DEFINITIONS


def component_rows() -> list[dict[str, Any]]:
    rows_out = []
    for index_code, comps in COMPONENTS.items():
        for code, name_ru, name_en, weight, source_note in comps:
            rows_out.append({
                "index_code": index_code,
                "component_code": code,
                "name_ru": name_ru,
                "name_en": name_en,
                "weight": weight,
                "source_note": source_note,
            })
    return rows_out


def alias_rows() -> list[dict[str, Any]]:
    rows_out = []
    for idx in INDEX_DEFINITIONS:
        rows_out.extend([
            {"alias": f"{idx['code']}:ru:primary", "index_code": idx["code"], "lang": "ru", "display_label": idx["short_name_ru"], "is_primary": 1},
            {"alias": f"{idx['code']}:en:primary", "index_code": idx["code"], "lang": "en", "display_label": idx["short_name_en"], "is_primary": 1},
            {"alias": f"{idx['code']}:en:name", "index_code": idx["code"], "lang": "en", "display_label": idx["name_en"], "is_primary": 0},
        ])
    return rows_out


def quality_flag_rows() -> list[dict[str, Any]]:
    return [
        {
            "quality_flag": "official_value",
            "label_ru": "официальное значение",
            "label_en": "official value",
            "severity": 0,
            "description_ru": "Значение взято напрямую из официального источника.",
            "description_en": "Value is taken directly from the official source.",
        },
        {
            "quality_flag": "recomputed_from_official_components",
            "label_ru": "пересчитано из официальных компонентов",
            "label_en": "recomputed from official components",
            "severity": 1,
            "description_ru": "Компонент или диагностика рассчитаны воспроизводимо из официальных рядов.",
            "description_en": "Component or diagnostic is reproducibly computed from official series.",
        },
        {
            "quality_flag": "official_partial_proxy",
            "label_ru": "официальные прокси с неполным покрытием",
            "label_en": "official proxies with partial coverage",
            "severity": 2,
            "description_ru": "Использованы официальные доступные прокси-компоненты; покрытие не полное.",
            "description_en": "Available official proxy components are used; coverage is partial.",
        },
        {
            "quality_flag": "official_file_required",
            "label_ru": "требуется официальный файл",
            "label_en": "official file required",
            "severity": 2,
            "description_ru": "Числовой слой включается после импорта официального файла источника.",
            "description_en": "The numeric layer is enabled after importing an official source file.",
        },
        {
            "quality_flag": "official_observation",
            "label_ru": "официальное исходное наблюдение",
            "label_en": "official source observation",
            "severity": 0,
            "description_ru": "Исходное наблюдение получено из официального снимка источника.",
            "description_en": "Raw observation is loaded from an official source snapshot.",
        },
        {
            "quality_flag": "selected_official_observation",
            "label_ru": "выбранное официальное наблюдение",
            "label_en": "selected official observation",
            "severity": 1,
            "description_ru": "Наблюдение выбрано из официальных источников по каскаду HTEI v3.",
            "description_en": "Observation selected from official sources by the HTEI v3 cascade.",
        },
        {
            "quality_flag": "training_model_score",
            "label_ru": "оценка модели подготовки кадров",
            "label_en": "training-system model score",
            "severity": 1,
            "description_ru": "Значение рассчитано как отдельная модель на основе проверяемых индексов и компонентов.",
            "description_en": "Value is computed as a separate model from auditable index and component inputs.",
        },
    ]


def ui_translation_rows() -> list[dict[str, Any]]:
    values = {
        "country_profile": {"ru": "Профиль страны", "en": "Country profile"},
        "source": {"ru": "Источник", "en": "Source"},
        "retrieved_at": {"ru": "Получено", "en": "Retrieved at"},
        "snapshot": {"ru": "Исходный файл", "en": "Snapshot"},
        "formula": {"ru": "Формула", "en": "Formula"},
    }
    return [{"key": key, "lang": lang, "value": value} for key, langs in values.items() for lang, value in langs.items()]


def build_value_rows(
    index_code: str,
    scores: list[dict[str, Any]],
    components: list[dict[str, Any]],
    snapshot: Snapshot,
    transform: Transform,
    formula_version: str,
    score_quality: str = "official_value",
    component_quality: str = "recomputed_from_official_components",
    score_is_official: bool = True,
    score_is_recomputed: bool = False,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    source = SOURCE_DEFS[snapshot.source_id]
    weight_map = {code: weight for code, _ru, _en, weight, _note in COMPONENTS[index_code]}
    score_rows: list[dict[str, Any]] = []
    ranking_rows: list[dict[str, Any]] = []
    component_rows_out: list[dict[str, Any]] = []
    diagnostics: list[dict[str, Any]] = []
    recommendations: list[dict[str, Any]] = []
    def item_year(item: dict[str, Any]) -> int:
        return int(item.get("year") or item.get("source_data_year") or SNAPSHOT_RELEASE_YEAR)

    score_by_key = {(s["iso3"], item_year(s)): s for s in scores}
    comps_by_key: dict[tuple[str, int], list[dict[str, Any]]] = {}
    for comp in components:
        comps_by_key.setdefault((comp["iso3"], item_year(comp)), []).append(comp)

    scores_by_year: dict[int, list[dict[str, Any]]] = {}
    for score in scores:
        scores_by_year.setdefault(item_year(score), []).append(score)

    for year_scores in scores_by_year.values():
        assign_ranks(year_scores)

    top10_by_component: dict[tuple[int, str], float] = {}
    for year, year_scores in scores_by_year.items():
        top_iso = {s["iso3"] for s in sorted(year_scores, key=lambda x: x["rank"])[:10]}
        for code in weight_map:
            vals = []
            for comp in components:
                if item_year(comp) == year and comp["component_code"] == code and comp["iso3"] in top_iso:
                    vals.append(comp["normalized_score"])
            if vals:
                top10_by_component[(year, code)] = sum(vals) / len(vals)

    for score in scores:
        iso3 = score["iso3"]
        value_year = item_year(score)
        value_id = f"{index_code}:{iso3}:{value_year}:score"
        score_source_id = score.get("source_id", snapshot.source_id)
        score_source = SOURCE_DEFS[score_source_id]
        score_source_url = score.get("source_url", snapshot.source_url)
        score_retrieved_at = score.get("retrieved_at", snapshot.retrieved_at)
        score_release_year = score.get("release_year", snapshot.release_year)
        score_raw_snapshot_path = score.get("raw_snapshot_path", snapshot.raw_snapshot_path)
        score_raw_snapshot_sha256 = score.get("raw_snapshot_sha256", snapshot.raw_snapshot_sha256)
        score_transformation_run_id = score.get("transformation_run_id", transform.transformation_run_id)
        score_transform_id = score.get("transform_id", transform.transform_id)
        prov = provenance(snapshot, transform, formula_version, score_quality, score_is_official, score_is_recomputed)
        prov.update({
            "source_id": score_source_id,
            "source_name": score_source["source_name"],
            "source_url": score_source_url,
            "retrieved_at": score_retrieved_at,
            "release_year": score_release_year,
            "raw_snapshot_path": score_raw_snapshot_path,
            "raw_snapshot_sha256": score_raw_snapshot_sha256,
            "transform_id": score_transform_id,
            "transformation_run_id": score_transformation_run_id,
            "source_group_coverage": score.get("source_group_coverage"),
            "available_weight": score.get("available_weight"),
        })
        data_quality = float(score.get("data_quality", 1.0))
        row = {
            "value_id": value_id,
            "index_code": index_code,
            "iso3": iso3,
            "year": value_year,
            "source_data_year": score.get("source_data_year"),
            "score": score["score"],
            "rank": score["rank"],
            "percentile": score.get("percentile"),
            "rank_delta_1y": None,
            "rank_delta_5y": None,
            "data_quality": data_quality,
            "source_type": "official_snapshot" if score_is_official else "recomputed_from_official_components",
            "source_id": score_source_id,
            "source_url": score_source_url,
            "retrieved_at": score_retrieved_at,
            "release_year": score_release_year,
            "raw_snapshot_path": score_raw_snapshot_path,
            "raw_snapshot_sha256": score_raw_snapshot_sha256,
            "transformation_run_id": score_transformation_run_id,
            "transform_id": score_transform_id,
            "formula_version": formula_version,
            "quality_flag": score_quality,
            "is_official": int(score_is_official),
            "is_recomputed": int(score_is_recomputed),
            "provenance_json": json.dumps(prov, ensure_ascii=False, sort_keys=True),
        }
        score_rows.append(row)
        ranking_rows.append({
            "index_code": index_code,
            "iso3": iso3,
            "year": value_year,
            "rank": score["rank"],
            "score": score["score"],
            "percentile": score.get("percentile"),
            "value_id": value_id,
        })

    for comp in components:
        iso3 = comp["iso3"]
        value_year = item_year(comp)
        score_for_component = score_by_key.get((iso3, value_year))
        if not score_for_component:
            continue
        weight = weight_map.get(comp["component_code"], 0.0)
        gap = max(0.0, 100.0 - float(comp["normalized_score"]))
        peer = top10_by_component.get((value_year, comp["component_code"]), 100.0)
        gap_peer = max(0.0, peer - float(comp["normalized_score"]))
        weighted_gap = gap * weight
        rank_leverage = min(1.0, weighted_gap / 35.0 + weight / 2)
        actionability = {"HDI": 0.45, "HCI": 0.62, "GTCI": 0.70, "GII": 0.68, "IDI": 0.75, "QS_ET": 0.62, "HTEI": 0.82}.get(index_code, 0.6)
        data_quality = float(comp.get("data_quality", score_for_component.get("data_quality", 1.0)))
        priority = gap * rank_leverage * actionability * data_quality
        comp_source_id = comp.get("source_id", snapshot.source_id)
        comp_source = SOURCE_DEFS[comp_source_id]
        comp_source_url = comp.get("source_url", snapshot.source_url)
        comp_retrieved_at = comp.get("retrieved_at", snapshot.retrieved_at)
        comp_release_year = comp.get("release_year", snapshot.release_year)
        comp_raw_snapshot_path = comp.get("raw_snapshot_path", snapshot.raw_snapshot_path)
        comp_raw_snapshot_sha256 = comp.get("raw_snapshot_sha256", snapshot.raw_snapshot_sha256)
        comp_transformation_run_id = comp.get("transformation_run_id", transform.transformation_run_id)
        comp_transform_id = comp.get("transform_id", transform.transform_id)
        if comp.get("provenance_json"):
            c_prov = json.loads(comp["provenance_json"])
            c_prov.update({
                "formula_version": formula_version,
                "quality_flag": comp.get("quality_flag", component_quality),
                "is_official": False,
                "is_recomputed": True,
            })
        else:
            c_prov = provenance(snapshot, transform, formula_version, component_quality, False, True)
        c_prov.update({
            "source_id": comp_source_id,
            "source_name": comp_source["source_name"],
            "source_group": comp.get("source_group"),
            "source_priority": comp.get("source_priority"),
            "source_url": comp_source_url,
            "retrieved_at": comp_retrieved_at,
            "release_year": comp_release_year,
            "raw_snapshot_path": comp_raw_snapshot_path,
            "raw_snapshot_sha256": comp_raw_snapshot_sha256,
            "transform_id": comp_transform_id,
            "transformation_run_id": comp_transformation_run_id,
        })
        c_value_id = f"{index_code}:{iso3}:{value_year}:{comp['component_code']}"
        component_rows_out.append({
            "value_id": c_value_id,
            "index_code": index_code,
            "component_code": comp["component_code"],
            "iso3": iso3,
            "year": value_year,
            "source_data_year": comp.get("source_data_year"),
            "raw_value": comp["raw_value"],
            "unit": comp["unit"],
            "normalized_score": comp["normalized_score"],
            "weighted_contribution": comp["normalized_score"] * weight,
            "gap_to_frontier": gap,
            "gap_to_peer_group": gap_peer,
            "weighted_gap": weighted_gap,
            "rank_leverage": rank_leverage,
            "actionability": actionability,
            "priority_score": priority,
            "source_id": comp_source_id,
            "source_url": comp_source_url,
            "retrieved_at": comp_retrieved_at,
            "release_year": comp_release_year,
            "raw_snapshot_path": comp_raw_snapshot_path,
            "raw_snapshot_sha256": comp_raw_snapshot_sha256,
            "transformation_run_id": comp_transformation_run_id,
            "transform_id": comp_transform_id,
            "formula_version": formula_version,
            "quality_flag": comp.get("quality_flag", component_quality),
            "is_official": 0,
            "is_recomputed": 1,
            "source_note": comp.get("source_note") or "Official source component transformed for platform diagnostics.",
            "provenance_json": json.dumps(c_prov, ensure_ascii=False, sort_keys=True),
        })

    comp_meta = {code: (ru, en) for code, ru, en, _w, _note in COMPONENTS[index_code]}
    for (iso3, value_year), comps in comps_by_key.items():
        score_for_key = score_by_key.get((iso3, value_year))
        if not score_for_key or not comps:
            continue
        ordered = sorted(
            (c for c in component_rows_out if c["iso3"] == iso3 and c["index_code"] == index_code and c["year"] == value_year),
            key=lambda item: item["priority_score"],
            reverse=True,
        )
        if not ordered:
            continue
        priority = ordered[0]
        best = max(ordered, key=lambda item: item["normalized_score"])
        ru, en = comp_meta.get(priority["component_code"], (priority["component_code"], priority["component_code"]))
        if index_code == "HTEI":
            explanation_ru = (
                f"Для выполнения блока ТЗ о формировании системы индикаторов для мониторинга развития технологических кадров "
                f"первым приоритетом является компонент «{ru}»: разрыв до фронтира составляет {priority['gap_to_frontier']:.1f} пункта. "
                f"Этот компонент напрямую связан с практическими рекомендациями по формированию трудовых ресурсов в высокотехнологичных отраслях Российской Федерации."
            )
            explanation_en = (
                f"For the High-Tech Employment Index, the first policy priority is {en}: the gap to frontier is {priority['gap_to_frontier']:.1f} points. "
                f"It is directly linked to recommendations on forming high-technology workforce capacity."
            )
        else:
            explanation_ru = f"Приоритет выбран из-за разрыва {priority['gap_to_frontier']:.1f} пункта до фронтира, веса компонента в методике и потенциального влияния на место страны."
            explanation_en = f"Priority is selected because the component is {priority['gap_to_frontier']:.1f} points below the frontier, carries methodological weight and can affect the country rank."
        diagnostics.append({
            "diagnostic_id": f"{index_code}:{iso3}:{value_year}:{priority['component_code']}",
            "iso3": iso3,
            "index_code": index_code,
            "component_code": priority["component_code"],
            "year": value_year,
            "gap_to_frontier": priority["gap_to_frontier"],
            "gap_to_peer_group": priority["gap_to_peer_group"],
            "weighted_gap": priority["weighted_gap"],
            "rank_leverage": priority["rank_leverage"],
            "actionability": priority["actionability"],
            "data_quality": score_for_key.get("data_quality", 1.0),
            "priority_score": priority["priority_score"],
            "weakest_component": priority["component_code"],
            "highest_leverage_component": best["component_code"],
            "explanation_ru": explanation_ru,
            "explanation_en": explanation_en,
            "source_value_id": priority["value_id"],
        })
        recommendations.append({
            "iso3": iso3,
            "index_code": index_code,
            "component_code": priority["component_code"],
            "year": value_year,
            "priority_score": priority["priority_score"],
            "title_ru": f"Приоритет: {ru}",
            "title_en": f"Priority: {en}",
            "text_ru": explanation_ru,
            "text_en": explanation_en,
            "horizon": "2026-2030",
            "policy_area": index_code,
        })
    return score_rows, component_rows_out, ranking_rows, diagnostics, recommendations

    top10_by_component: dict[str, float] = {}
    for code in weight_map:
        vals = []
        top_iso = {s["iso3"] for s in sorted(scores, key=lambda x: x["rank"])[:10]}
        for comp in components:
            if comp["component_code"] == code and comp["iso3"] in top_iso:
                vals.append(comp["normalized_score"])
        if vals:
            top10_by_component[code] = sum(vals) / len(vals)

    for score in scores:
        iso3 = score["iso3"]
        value_id = f"{index_code}:{iso3}:{SNAPSHOT_RELEASE_YEAR}:score"
        prov = provenance(snapshot, transform, formula_version, score_quality, score_is_official, score_is_recomputed)
        data_quality = float(score.get("data_quality", 1.0))
        row = {
            "value_id": value_id,
            "index_code": index_code,
            "iso3": iso3,
            "year": SNAPSHOT_RELEASE_YEAR,
            "source_data_year": score.get("source_data_year"),
            "score": score["score"],
            "rank": score["rank"],
            "percentile": score.get("percentile"),
            "rank_delta_1y": None,
            "rank_delta_5y": None,
            "data_quality": data_quality,
            "source_type": "official_snapshot" if score_is_official else "recomputed_from_official_components",
            "source_id": snapshot.source_id,
            "source_url": snapshot.source_url,
            "retrieved_at": snapshot.retrieved_at,
            "release_year": snapshot.release_year,
            "raw_snapshot_path": snapshot.raw_snapshot_path,
            "raw_snapshot_sha256": snapshot.raw_snapshot_sha256,
            "transformation_run_id": transform.transformation_run_id,
            "transform_id": transform.transform_id,
            "formula_version": formula_version,
            "quality_flag": score_quality,
            "is_official": int(score_is_official),
            "is_recomputed": int(score_is_recomputed),
            "provenance_json": json.dumps({**prov, "source_name": source["source_name"]}, ensure_ascii=False, sort_keys=True),
        }
        score_rows.append(row)
        ranking_rows.append({
            "index_code": index_code,
            "iso3": iso3,
            "year": SNAPSHOT_RELEASE_YEAR,
            "rank": score["rank"],
            "score": score["score"],
            "percentile": score.get("percentile"),
            "value_id": value_id,
        })

    for comp in components:
        iso3 = comp["iso3"]
        if iso3 not in score_by_iso:
            continue
        weight = weight_map.get(comp["component_code"], 0.0)
        gap = max(0.0, 100.0 - float(comp["normalized_score"]))
        peer = top10_by_component.get(comp["component_code"], 100.0)
        gap_peer = max(0.0, peer - float(comp["normalized_score"]))
        weighted_gap = gap * weight
        rank_leverage = min(1.0, weighted_gap / 35.0 + weight / 2)
        actionability = {"HDI": 0.45, "HCI": 0.62, "GTCI": 0.70, "GII": 0.68, "IDI": 0.75, "HTEI": 0.82}.get(index_code, 0.6)
        data_quality = float(comp.get("data_quality", score_by_iso[iso3].get("data_quality", 1.0)))
        priority = gap * rank_leverage * actionability * data_quality
        c_prov = provenance(snapshot, transform, formula_version, component_quality, False, True)
        c_value_id = f"{index_code}:{iso3}:{SNAPSHOT_RELEASE_YEAR}:{comp['component_code']}"
        component_rows_out.append({
            "value_id": c_value_id,
            "index_code": index_code,
            "component_code": comp["component_code"],
            "iso3": iso3,
            "year": SNAPSHOT_RELEASE_YEAR,
            "source_data_year": comp.get("source_data_year"),
            "raw_value": comp["raw_value"],
            "unit": comp["unit"],
            "normalized_score": comp["normalized_score"],
            "weighted_contribution": comp["normalized_score"] * weight,
            "gap_to_frontier": gap,
            "gap_to_peer_group": gap_peer,
            "weighted_gap": weighted_gap,
            "rank_leverage": rank_leverage,
            "actionability": actionability,
            "priority_score": priority,
            "source_id": snapshot.source_id,
            "source_url": snapshot.source_url,
            "retrieved_at": snapshot.retrieved_at,
            "release_year": snapshot.release_year,
            "raw_snapshot_path": snapshot.raw_snapshot_path,
            "raw_snapshot_sha256": snapshot.raw_snapshot_sha256,
            "transformation_run_id": transform.transformation_run_id,
            "transform_id": transform.transform_id,
            "formula_version": formula_version,
            "quality_flag": component_quality,
            "is_official": 0,
            "is_recomputed": 1,
            "source_note": "Official source component transformed for platform diagnostics.",
            "provenance_json": json.dumps({**c_prov, "source_name": source["source_name"]}, ensure_ascii=False, sort_keys=True),
        })

    comp_meta = {code: (ru, en) for code, ru, en, _w, _note in COMPONENTS[index_code]}
    for iso3, comps in comps_by_iso.items():
        if iso3 not in score_by_iso or not comps:
            continue
        ordered = sorted(
            (c for c in component_rows_out if c["iso3"] == iso3 and c["index_code"] == index_code),
            key=lambda item: item["priority_score"],
            reverse=True,
        )
        if not ordered:
            continue
        priority = ordered[0]
        best = max(ordered, key=lambda item: item["normalized_score"])
        ru, en = comp_meta.get(priority["component_code"], (priority["component_code"], priority["component_code"]))
        explanation_ru = f"Приоритет выбран из-за разрыва {priority['gap_to_frontier']:.1f} пункта до фронтира и веса компонента в методике."
        explanation_en = f"Priority is selected because the component is {priority['gap_to_frontier']:.1f} points below the frontier and carries methodological weight."
        diagnostics.append({
            "diagnostic_id": f"{index_code}:{iso3}:{SNAPSHOT_RELEASE_YEAR}:{priority['component_code']}",
            "iso3": iso3,
            "index_code": index_code,
            "component_code": priority["component_code"],
            "year": SNAPSHOT_RELEASE_YEAR,
            "gap_to_frontier": priority["gap_to_frontier"],
            "gap_to_peer_group": priority["gap_to_peer_group"],
            "weighted_gap": priority["weighted_gap"],
            "rank_leverage": priority["rank_leverage"],
            "actionability": priority["actionability"],
            "data_quality": score_by_iso[iso3].get("data_quality", 1.0),
            "priority_score": priority["priority_score"],
            "weakest_component": priority["component_code"],
            "highest_leverage_component": best["component_code"],
            "explanation_ru": explanation_ru,
            "explanation_en": explanation_en,
            "source_value_id": priority["value_id"],
        })
        recommendations.append({
            "iso3": iso3,
            "index_code": index_code,
            "component_code": priority["component_code"],
            "priority_score": priority["priority_score"],
            "title_ru": f"Приоритет: {ru}",
            "title_en": f"Priority: {en}",
            "text_ru": explanation_ru,
            "text_en": explanation_en,
            "horizon": "2026-2030",
            "policy_area": index_code,
        })
    return score_rows, component_rows_out, ranking_rows, diagnostics, recommendations



def apply_rank_deltas(conn) -> None:
    """Populate rank_delta_1y and rank_delta_5y from actual rank history.

    Positive delta means the country improved its rank position (e.g. from 20 to 15 = +5).
    """
    rows_in = conn.execute("SELECT value_id, index_code, iso3, year, rank FROM index_scores ORDER BY index_code, iso3, year").fetchall()
    by_key: dict[tuple[str, str], list] = {}
    for r in rows_in:
        by_key.setdefault((r["index_code"], r["iso3"]), []).append(r)
    for (_idx, _iso3), vals in by_key.items():
        by_year = {int(v["year"]): v for v in vals}
        for v in vals:
            year = int(v["year"]); rank = int(v["rank"])
            d1 = None; d5 = None
            if year - 1 in by_year:
                d1 = int(by_year[year - 1]["rank"]) - rank
            if year - 5 in by_year:
                d5 = int(by_year[year - 5]["rank"]) - rank
            conn.execute("UPDATE index_scores SET rank_delta_1y=?, rank_delta_5y=? WHERE value_id=?", (d1, d5, v["value_id"]))


def build_training_model_rows(
    score_rows: list[dict[str, Any]],
    component_rows: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    component_meta = {
        (index_code, code): {"name_ru": ru, "name_en": en}
        for index_code, comps in COMPONENTS.items()
        for code, ru, en, _weight, _note in comps
    }
    by_ref_iso: dict[tuple[str, str, str], list[dict[str, Any]]] = {}
    for comp in component_rows:
        by_ref_iso.setdefault((comp["index_code"], comp["component_code"], comp["iso3"]), []).append(comp)
    for items in by_ref_iso.values():
        items.sort(key=lambda item: int(item["year"]))

    countries = sorted({row["iso3"] for row in score_rows})
    years = sorted({int(row["year"]) for row in score_rows if int(row["year"]) >= 2023})
    model_scores: list[dict[str, Any]] = []
    model_components: list[dict[str, Any]] = []
    block_scores_for_rank: dict[int, list[dict[str, Any]]] = {}

    def latest_component(index_code: str, component_code: str, iso3: str, year: int) -> dict[str, Any] | None:
        candidates = [item for item in by_ref_iso.get((index_code, component_code, iso3), []) if int(item["year"]) <= year]
        return candidates[-1] if candidates else None

    for year in years:
        for iso3 in countries:
            block_rows: list[dict[str, Any]] = []
            available_block_weight = 0.0
            weighted_model_score = 0.0
            source_groups: set[str] = set()
            for block in TRAINING_MODEL_BLOCKS:
                per_component_weight = 1.0 / len(block["components"])
                component_inputs: list[dict[str, Any]] = []
                for index_code, component_code in block["components"]:
                    source = latest_component(index_code, component_code, iso3, year)
                    if not source:
                        continue
                    normalized = float(source["normalized_score"])
                    component_inputs.append({
                        "source": source,
                        "index_code": index_code,
                        "component_code": component_code,
                        "normalized": normalized,
                        "component_weight": per_component_weight,
                    })
                component_coverage = sum(item["component_weight"] for item in component_inputs)
                if component_coverage < 0.5:
                    continue
                block_score = sum(item["normalized"] * item["component_weight"] for item in component_inputs) / component_coverage
                available_block_weight += float(block["weight"])
                weighted_model_score += block_score * float(block["weight"])
                for item in component_inputs:
                    source = item["source"]
                    source_group = None
                    try:
                        source_group = json.loads(source.get("provenance_json") or "{}").get("source_group")
                    except Exception:
                        source_group = None
                    source_group = source_group or source["source_id"]
                    source_groups.add(source_group)
                    meta = component_meta.get((item["index_code"], item["component_code"]), {"name_ru": item["component_code"], "name_en": item["component_code"]})
                    block_rows.append({
                        "iso3": iso3,
                        "year": year,
                        "block_code": block["code"],
                        "block_name_ru": block["name_ru"],
                        "block_name_en": block["name_en"],
                        "component_ref": f"{item['index_code']}.{item['component_code']}",
                        "component_name_ru": meta["name_ru"],
                        "component_name_en": meta["name_en"],
                        "source_value_id": source["value_id"],
                        "source_data_year": source.get("source_data_year") or source.get("year"),
                        "raw_value": source.get("raw_value"),
                        "normalized_score": item["normalized"],
                        "weight": float(block["weight"]) * item["component_weight"],
                        "weighted_contribution": item["normalized"] * float(block["weight"]) * item["component_weight"],
                        "source_id": source["source_id"],
                        "source_group": source_group,
                        "quality_flag": source["quality_flag"],
                        "provenance_json": json.dumps({
                            "source_value_id": source["value_id"],
                            "source_id": source["source_id"],
                            "source_group": source_group,
                            "source_data_year": source.get("source_data_year") or source.get("year"),
                            "formula_version": TRAINING_MODEL_VERSION,
                        }, ensure_ascii=False, sort_keys=True),
                    })
            if available_block_weight < 0.65 or len({row["block_code"] for row in block_rows}) < 3:
                continue
            score = weighted_model_score / available_block_weight
            data_quality = min(1.0, 0.5 * available_block_weight + 0.5 * min(1.0, len(source_groups) / 5))
            value_id = f"TRAINING_SYSTEM:{iso3}:{year}:score"
            model_row = {
                "value_id": value_id,
                "iso3": iso3,
                "year": year,
                "requested_year": year,
                "score": score * (0.9 + 0.1 * data_quality),
                "rank": 0,
                "data_quality": data_quality,
                "available_block_weight": available_block_weight,
                "source_group_coverage": json.dumps(sorted(source_groups), ensure_ascii=False),
                "formula_version": TRAINING_MODEL_VERSION,
                "quality_flag": "training_model_score",
                "provenance_json": json.dumps({
                    "formula_version": TRAINING_MODEL_VERSION,
                    "block_weights": {block["code"]: block["weight"] for block in TRAINING_MODEL_BLOCKS},
                    "available_block_weight": available_block_weight,
                    "source_group_coverage": sorted(source_groups),
                }, ensure_ascii=False, sort_keys=True),
            }
            block_scores_for_rank.setdefault(year, []).append(model_row)
            model_scores.append(model_row)
            for row_in in block_rows:
                model_components.append({
                    **row_in,
                    "value_id": f"TRAINING_SYSTEM:{iso3}:{year}:{row_in['block_code']}:{row_in['component_ref']}",
                    "model_value_id": value_id,
                })
    for year_rows in block_scores_for_rank.values():
        assign_ranks(year_rows)
    return model_scores, model_components


def build_production_database(db_path: Path | str | None = DB_PATH, reset: bool = True) -> dict[str, Any]:
    db_path = Path(db_path) if db_path is not None else DB_PATH
    snapshots: dict[str, Snapshot] = {
        "UNDP_HDR": ensure_download("UNDP_HDR", 2025, UNDP_HDI_CSV_URL, "HDR25_Composite_indices_complete_time_series.csv"),
        "WORLD_BANK_HCI": ensure_world_bank_snapshot("WORLD_BANK_HCI", 2020, HCI_INDICATORS, "world_bank_hci_indicators.json"),
        "WORLD_BANK_COUNTRIES": ensure_world_bank_countries_snapshot(),
        "PORTULANS_GTCI": ensure_download("PORTULANS_GTCI", 2025, PORTULANS_GTCI_PDF_URL, "GTCI_2025_report.pdf"),
        "WIPO_GII": ensure_download("WIPO_GII", 2025, WIPO_GII_XLSX_URL, "wipo-pub-2000-2025-gii-tech1.xlsx"),
        "ITU_IDI": ensure_download("ITU_IDI", 2025, ITU_IDI_XLSX_URL, "IDIDataset_2025.xlsx"),
        "WORLD_BANK_HTEI": ensure_world_bank_snapshot("WORLD_BANK_HTEI", 2025, HTEI_INDICATORS, "world_bank_htei_indicators.json"),
        "UIS_HTEI": ensure_uis_snapshot(),
        "ILOSTAT_HTEI": ensure_multifile_snapshot("ILOSTAT_HTEI", 2025, ILOSTAT_HTEI_URLS, "ilostat_htei_manifest.json"),
        "OECD_HTEI": ensure_multifile_snapshot("OECD_HTEI", 2025, OECD_HTEI_URLS, "oecd_htei_manifest.json", accept="text/csv"),
        "EUROSTAT_HTEC": ensure_multifile_snapshot("EUROSTAT_HTEC", 2025, EUROSTAT_HTEI_URLS, "eurostat_htei_manifest.json", accept="application/json"),
        "NATIONAL_STATS_HTEI": ensure_json_snapshot("NATIONAL_STATS_HTEI", 2026, {
            "purpose": "Curated official national statistical office register. Numeric country observations enter HTEI only when official files are archived with source-specific checksums.",
            "initial_scope": ["Russia", "G20", "BRICS", "OECD/EU gaps", "technology benchmark countries"],
            "numeric_rows_loaded": 0,
        }, "national_stats_htei_register.json"),
        "CORPORATE_REPORTS_HTEI": ensure_json_snapshot("CORPORATE_REPORTS_HTEI", 2026, {
            "purpose": "Corporate reporting evidence register. Reports are audit evidence unless country coverage is sufficient for numeric scoring.",
            "numeric_rows_loaded": 0,
            "selection_policy": "official issuer report, archived path, checksum and reproducible extraction required before scoring",
        }, "corporate_reports_htei_register.json"),
        "SPECIALIZED_RATINGS_HTEI": ensure_json_snapshot("SPECIALIZED_RATINGS_HTEI", 2026, {
            "purpose": "Specialized rankings evidence register for the training-system model.",
            "linked_official_snapshots": ["QS_ET", "WIPO_GII", "PORTULANS_GTCI"],
        }, "specialized_ratings_htei_register.json"),
    }
    snapshots["HTEI_MULTI_SOURCE"] = ensure_json_snapshot("HTEI_MULTI_SOURCE", 2026, {
        "formula_versions": {
            "strict_latest": "htei-v3-multisource",
            "asof_release": "htei-v3-multisource-asof-2026",
        },
        "component_weights": HTEI_V3_WEIGHTS,
        "source_priority": HTEI_SOURCE_PRIORITY,
        "source_snapshot_ids": {key: snapshots[key].snapshot_id for key in ["ILOSTAT_HTEI", "OECD_HTEI", "EUROSTAT_HTEC", "UIS_HTEI", "WORLD_BANK_HTEI"]},
        "modes": {
            "HTEI_STRICT_LATEST": {
                "available_weight": HTEI_RELEASE_MIN_WEIGHT,
                "components": HTEI_RELEASE_MIN_COMPONENTS,
                "source_groups": HTEI_RELEASE_MIN_SOURCE_GROUPS,
                "same_year_only": True,
            },
            "HTEI_ASOF": {
                "available_weight": 0.30,
                "components": 2,
                "selection_rule": "latest selected official observation at or before release year",
                "quality_adjustment": True,
            },
        },
    }, "htei_v3_multisource_manifest.json")
    all_snapshots: list[Snapshot] = list(snapshots.values())

    gii_snapshots: list[Snapshot] = [snapshots["WIPO_GII"]]
    for year, url in GII_XLSX_URLS.items():
        if year == 2025:
            continue
        snap = ensure_optional_download("WIPO_GII", year, url, f"wipo-pub-2000-{year}-gii-tech1.xlsx")
        if snap:
            gii_snapshots.append(snap)
            all_snapshots.append(snap)
    gii_snapshots = sorted(gii_snapshots, key=lambda item: item.release_year)

    qs_snapshots = ensure_qs_snapshots(QS_YEARS)
    latest_qs = qs_snapshots[max(qs_snapshots)]
    snapshots["QS_ET"] = latest_qs
    all_snapshots.extend(qs_snapshots[year] for year in sorted(qs_snapshots))
    for year in sorted(qs_snapshots):
        all_snapshots.extend(qs_page_snapshots(qs_snapshots[year]))
    all_snapshots.extend(archived_snapshot_metadata())
    snapshot_by_id: dict[str, Snapshot] = {}
    for snapshot in sorted(all_snapshots, key=lambda item: item.retrieved_at):
        snapshot_by_id[snapshot.snapshot_id] = snapshot
    all_snapshots = list(snapshot_by_id.values())

    parsed_bundles: list[tuple[str, list[dict[str, Any]], list[dict[str, Any]], Snapshot, Transform, str, str, str, bool, bool]] = []
    hdi_scores, hdi_components, hdi_transform = parse_undp_hdi(snapshots["UNDP_HDR"])
    parsed_bundles.append(("HDI", hdi_scores, hdi_components, snapshots["UNDP_HDR"], hdi_transform, "official-hdi-timeseries-2025", "official_value", "recomputed_from_official_components", True, False))
    hci_scores, hci_components, hci_transform = parse_hci(snapshots["WORLD_BANK_HCI"])
    parsed_bundles.append(("HCI", hci_scores, hci_components, snapshots["WORLD_BANK_HCI"], hci_transform, "official-hci-series", "official_value", "recomputed_from_official_components", True, False))
    gtci_scores, gtci_components, gtci_transform = parse_gtci(snapshots["PORTULANS_GTCI"])
    parsed_bundles.append(("GTCI", gtci_scores, gtci_components, snapshots["PORTULANS_GTCI"], gtci_transform, "official-gtci-2025-pillars", "official_value", "recomputed_from_official_components", True, False))
    for gii_snapshot in gii_snapshots:
        try:
            gii_scores, gii_components, gii_transform = parse_gii(gii_snapshot)
        except Exception:
            continue
        if gii_scores:
            parsed_bundles.append(("GII", gii_scores, gii_components, gii_snapshot, gii_transform, "official-gii-2025", "official_value", "official_value", True, False))
    idi_scores, idi_components, idi_transform = parse_idi(snapshots["ITU_IDI"])
    parsed_bundles.append(("IDI", idi_scores, idi_components, snapshots["ITU_IDI"], idi_transform, "official-idi-series-2023-2025", "official_value", "official_value", True, False))
    htei_scores, htei_components, htei_source_observations, htei_transforms = parse_htei_multisource(snapshots, snapshots["HTEI_MULTI_SOURCE"])
    htei_transform = htei_transforms[0]
    htei_extra_transforms = htei_transforms[1:]
    parsed_bundles.append(("HTEI", htei_scores, htei_components, snapshots["HTEI_MULTI_SOURCE"], htei_transform, "htei-v3-multisource", "recomputed_from_official_components", "selected_official_observation", False, True))

    qs_institution_rows: list[dict[str, Any]] = []
    qs_unmatched: dict[int, list[str]] = {}
    for year in sorted(qs_snapshots):
        qs_scores, qs_components, qs_transform, institution_rows, unmatched = parse_qs_endpoint(qs_snapshots[year])
        parsed_bundles.append(("QS_ET", qs_scores, qs_components, qs_snapshots[year], qs_transform, "qs-country-aggregation-v1", "recomputed_from_official_components", "recomputed_from_official_components", False, True))
        qs_institution_rows.extend(institution_rows)
        if unmatched:
            qs_unmatched[year] = unmatched

    country_region_metadata, country_income_metadata, country_metadata_audit = parse_world_bank_country_metadata(snapshots["WORLD_BANK_COUNTRIES"])
    iso_to_name: dict[str, str] = {}
    iso_to_region: dict[str, str] = {}
    iso_to_income: dict[str, str] = {}
    all_scores: list[dict[str, Any]] = []
    all_components: list[dict[str, Any]] = []
    all_rankings: list[dict[str, Any]] = []
    all_diagnostics: list[dict[str, Any]] = []
    all_recommendations: list[dict[str, Any]] = []
    all_source_observations: list[dict[str, Any]] = list(htei_source_observations)
    transforms: list[Transform] = []
    for code, scores, comps, bundle_snapshot, transform, formula_version, score_quality, component_quality, score_official, score_recomputed in parsed_bundles:
        transforms.append(transform)
        for score in scores:
            iso3 = score["iso3"]
            iso_to_name.setdefault(iso3, country_name_en(iso3))
            iso_to_region.setdefault(iso3, country_region_metadata.get(iso3, COUNTRY_METADATA_FALLBACK_REGION))
            iso_to_income.setdefault(iso3, country_income_metadata.get(iso3, COUNTRY_METADATA_FALLBACK_INCOME))
        score_rows, comp_rows, ranking_rows, diag_rows, rec_rows = build_value_rows(
            code,
            scores,
            comps,
            bundle_snapshot,
            transform,
            formula_version,
            score_quality,
            component_quality,
            score_official,
            score_recomputed,
        )
        all_scores.extend(score_rows)
        all_components.extend(comp_rows)
        all_rankings.extend(ranking_rows)
        all_diagnostics.extend(diag_rows)
        all_recommendations.extend(rec_rows)
    transforms.extend(htei_extra_transforms)
    training_scores, training_components = build_training_model_rows(all_scores, all_components)

    init_db(db_path, reset=reset)
    with connect(db_path) as conn:
        upsert_many(conn, "source_registry", source_registry_rows(snapshots))
        upsert_many(conn, "raw_snapshots", [
            s.__dict__ | {
                "is_official": int(
                    (SOURCE_DEFS.get(s.source_id, {}).get("source_role") or source_role_for(s.source_id))
                    not in {"audit_register", "qualitative_evidence"}
                )
            }
            for s in all_snapshots
        ])
        upsert_many(conn, "transformation_runs", [t.__dict__ for t in transforms])
        upsert_many(conn, "countries", country_rows(iso_to_name, iso_to_region, iso_to_income))
        upsert_many(conn, "index_definitions", index_rows())
        upsert_many(conn, "indices", index_rows())
        upsert_many(conn, "index_aliases", alias_rows())
        comps = component_rows()
        upsert_many(conn, "index_components", comps)
        upsert_many(conn, "components", comps)
        upsert_many(conn, "index_formulas", formula_rows())
        upsert_many(conn, "quality_flags", quality_flag_rows())
        upsert_many(conn, "ui_translations", ui_translation_rows())
        upsert_many(conn, "index_scores", all_scores)
        upsert_many(conn, "component_values", all_components)
        upsert_many(conn, "source_observations", all_source_observations)
        upsert_many(conn, "rankings", all_rankings)
        upsert_many(conn, "qs_institution_rankings", qs_institution_rows)
        upsert_many(conn, "diagnostics", all_diagnostics)
        upsert_many(conn, "recommendations", all_recommendations)
        upsert_many(conn, "training_model_scores", training_scores)
        upsert_many(conn, "training_model_components", training_components)
        apply_rank_deltas(conn)
        ensure_htei_asof_release(conn)
        conn.commit()

    with connect(db_path) as conn:
        release_facts = {
            "sources": conn.execute("SELECT COUNT(*) AS n FROM source_registry").fetchone()["n"],
            "scores": conn.execute("SELECT COUNT(*) AS n FROM index_scores").fetchone()["n"],
            "components": conn.execute("SELECT COUNT(*) AS n FROM component_values").fetchone()["n"],
            "source_observations": conn.execute("SELECT COUNT(*) AS n FROM source_observations").fetchone()["n"],
            "training_model_scores": conn.execute("SELECT COUNT(*) AS n FROM training_model_scores").fetchone()["n"],
            "training_model_components": conn.execute("SELECT COUNT(*) AS n FROM training_model_components").fetchone()["n"],
            "countries": conn.execute("SELECT COUNT(*) AS n FROM countries").fetchone()["n"],
            "qs_institutions": conn.execute("SELECT COUNT(*) AS n FROM qs_institution_rankings").fetchone()["n"],
            "raw_snapshots": conn.execute("SELECT COUNT(*) AS n FROM raw_snapshots").fetchone()["n"],
        }

    return {
        "status": "ok",
        "db": str(db_path),
        **release_facts,
        "indices": len(INDEX_DEFINITIONS),
        "qs_unmatched_countries": qs_unmatched,
        "country_metadata": country_metadata_audit,
    }


def build_if_missing(db_path: Path | str | None = DB_PATH) -> None:
    db_path = Path(db_path) if db_path is not None else DB_PATH
    if not db_path.exists():
        build_production_database(db_path, reset=True)
        return
    try:
        with connect(db_path) as conn:
            existing_tables = {r["name"] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
            required_tables = {"source_observations", "training_model_scores", "training_model_components"}
            formula = conn.execute("SELECT formula_version FROM index_formulas WHERE index_code='HTEI' ORDER BY formula_version DESC LIMIT 1").fetchone()
            latest_htei = conn.execute("SELECT MAX(year) AS y FROM index_scores WHERE index_code='HTEI'").fetchone()
            country_meta = conn.execute("SELECT COUNT(*) AS n FROM source_registry WHERE source_id='WORLD_BANK_COUNTRIES'").fetchone()
            country_counts = conn.execute("SELECT COUNT(*) AS total, SUM(CASE WHEN region='Global' OR income_group='Not classified' THEN 1 ELSE 0 END) AS generic FROM countries").fetchone()
            needs_rebuild = (
                not required_tables.issubset(existing_tables)
                or not formula
                or not str(formula["formula_version"]).startswith("htei-v3-multisource")
                or latest_htei["y"] is None
                or int(latest_htei["y"]) < 2022
                or not country_meta
                or int(country_meta["n"] or 0) == 0
                or int(country_counts["generic"] or 0) > max(5, int(country_counts["total"] or 0) // 2)
            )
    except Exception:
        needs_rebuild = True
    if needs_rebuild:
        build_production_database(db_path, reset=True)
        return


if __name__ == "__main__":
    print(json.dumps(build_production_database(), ensure_ascii=False, indent=2))
